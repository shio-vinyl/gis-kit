#!/usr/bin/env python3
"""Inspect a DXF drawing before using it as GIS data: layers, entities, units, extent, candidate CRS and a check plot.

`cad.py inspect` reads model-space geometry through GDAL (block references expanded) and, when the optional
ezdxf is installed, inventories layers as CAD counts them: DXF entity types, closed and open polylines,
block names and attributes, raw and cleaned text, layer state, layouts and viewports. Without ezdxf the
inventory comes from GDAL and cad.json lists what is missing. It reads $ACADVER / $INSUNITS /
$DWGCODEPAGE from the header itself, and lists projected CRS
whose grid the coordinates fit, mainly zone-prefixed Chinese Gauss-Kruger grids. Coordinates without
the zone prefix are held: they may equally be a local grid, so no EPSG code is listed for them unless
--reference, a layer in a known projected CRS, lines up with the drawing's extent.
CAD files carry no CRS: every match is a candidate, the datum (CGCS2000, New Beijing, Xian 1980,
Beijing 1954) cannot be told apart from coordinates, and nothing is assigned to the data here.
Writes cad.json, check.png and record.json to a new directory. DWG is not read: `cad.py convert` runs an
installed ODA File Converter or LibreDWG on a copy and reads the DXF back; nothing is installed by it.

`cad.py extract` writes model-space entities of chosen layers to raw.gpkg (lines, closed_polygons, texts,
blocks) with layer and handle on every feature: curves flattened to --arc-tolerance, units scaled to metres
by an explicit --units, non-zero Z refused unless --drop-z, no CRS. It needs ezdxf.

`cad.py polygonize` merges endpoints within --snap (and moves free ends onto lines within it), nodes the
lines and polygonizes them into faces.gpkg: faces (small or invalid ones held, never dropped), gaps up to
--max-gap and dangles as QA layers, topo.py checks on the faces. Gaps are closed only with --close-gaps.

`cad.py extract` also writes each HATCH as the area it fills (hatches): boundary loops nested by containment and
filled by the hatch style, with pattern, colour and style kept. `cad.py compare-fills` matches those fills to faces
from polygonize, label or xiangyuan (each fill to the face with the highest IoU) and measures the area only one side covers.

`cad.py label` puts parcel fields on polygonized faces from CAD texts by YAML rules (one field per text, or several
by named groups), placing each text at its alignment point or leader arrow. Several distinct values in a face,
texts outside every face or on a shared edge, unparsed texts, duplicate parcel ids and unlinked leaders are
listed, never resolved by choice; road faces come only from the road rule.

`cad.py assign-crs` labels a CRS-less layer converted from an inspected drawing with a CRS someone has
confirmed: --confirmed-by and --basis are required, the CRS must be an inspection candidate as drawn
(or --outside-candidates), and geometry WKB and attributes are compared on readback before publishing.
A grid without the zone prefix, or a CRS outside the candidates, also needs a feature check against declared
thresholds: vertices on a reference layer's edges (--check-reference) or known points (--check-points).

`cad.py overlay` puts parcels (in a projected CRS) over reference layers in one check plot: hold parcels outlined,
QA layers of the same file (qa_*, label_conflicts) marked. References are read within the parcels' extent plus a
margin and reprojected; parcel x reference intersections go to overlay.gpkg with area and both shares. Invalid
polygons are listed and plotted, never repaired or measured.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from functools import lru_cache
import json
import math
from pathlib import Path
import re

from _sandbox import configure_writable_caches

configure_writable_caches()

from _budget import check_vector
from _delivery import bundle, digest, write_json

VERSION = '0.1.0'
VERSIONS = {'AC1009': 'R12', 'AC1012': 'R13', 'AC1014': 'R14', 'AC1015': 'R2000', 'AC1018': 'R2004',
            'AC1021': 'R2007', 'AC1024': 'R2010', 'AC1027': 'R2013', 'AC1032': 'R2018'}
# $INSUNITS code -> (name, metres per drawing unit); codes not listed are reported without a scale.
UNITS = {0: ('unitless', None), 1: ('inches', 0.0254), 2: ('feet', 0.3048), 4: ('millimeters', 0.001),
         5: ('centimeters', 0.01), 6: ('meters', 1.0), 7: ('kilometers', 1000.0), 10: ('yards', 0.9144),
         14: ('decimeters', 0.1)}
DATUMS = ('CGCS2000', 'New Beijing', 'Xian 1980', 'Beijing 1954')
TEXT_SAMPLES = 5
PLOT_LAYERS = 12
OUTLIER_SPREAD = 20
REFERENCE_OVERLAP = 0.5  # share of the smaller extent inside the larger one
OVERLAY_MARGIN = 0.1  # share of the parcels' longer side shown, and read from references, around them
# What the GDAL fallback cannot report; listed in cad.json when ezdxf is not installed.
EZDXF_ONLY = ['DXF entity types for R12', 'closed and open polylines', 'block names and attribute tags',
              'raw and cleaned text', 'layer off and frozen state', 'unused layers', 'layouts and viewports']


def read_header(path):
    """Header variables of an ASCII DXF; GDAL does not expose them."""
    with open(path, 'rb') as stream:
        if stream.read(18) == b'AutoCAD Binary DXF':
            return {'readable': False, 'reason': 'binary DXF header is not parsed'}
        stream.seek(0)
        values, name, in_header, code = {}, None, False, None
        for number, raw in enumerate(stream):
            text = raw.decode('latin-1').strip()
            if number % 2 == 0:
                code = text
                continue
            if code == '0' and text == 'ENDSEC' and in_header:
                break
            if code == '2' and text == 'HEADER':
                in_header = True
            elif code == '0' and text == 'SECTION':
                continue
            elif code == '2' and not in_header:
                break  # the first section is not HEADER: an entities-only R12 file
            elif in_header and code == '9':
                name = text
            elif in_header and name:
                values.setdefault(name, {})[code] = text
    version = values.get('$ACADVER', {}).get('1')
    units = values.get('$INSUNITS', {}).get('70')
    units = int(units) if units is not None and units.lstrip('-').isdigit() else None
    extent = [values.get(key, {}).get(axis) for key in ('$EXTMIN', '$EXTMAX') for axis in ('10', '20')]
    try:
        extent = [float(v) for v in extent]
    except (TypeError, ValueError):
        extent = None
    return {'readable': bool(values), 'acadver': version, 'release': VERSIONS.get(version),
            'insunits': units, 'units': UNITS.get(units, ('unknown', None))[0] if units is not None else None,
            'codepage': values.get('$DWGCODEPAGE', {}).get('3'), 'extent': extent,
            # R12 may carry no entity handles ($HANDLING 0 or missing); R13 and later always do.
            'handles': (version != 'AC1009' or values.get('$HANDLING', {}).get('70') == '1') if version else None}


@lru_cache(maxsize=None)
def gauss_kruger():
    """EPSG zone-prefixed Gauss-Kruger grids keyed by (width, zone), datums in DATUMS order."""
    from pyproj.database import query_crs_info
    pattern = re.compile(r'(%s) / (3-degree )?Gauss-Kruger zone (\d+)' % '|'.join(DATUMS))
    grids = {}
    for info in query_crs_info(auth_name='EPSG', pj_types='PROJECTED_CRS'):
        match = pattern.fullmatch(info.name)
        if match and not info.deprecated:
            key = (3 if match[2] else 6, int(match[3]))
            grids.setdefault(key, []).append((match[1], f'EPSG:{info.code}', info.name, info.area_of_use))
    return {key: sorted(rows, key=lambda row: DATUMS.index(row[0])) for key, rows in grids.items()}


def distance_km(a, b):
    lon1, lat1, lon2, lat2 = map(math.radians, (*a, *b))
    h = math.sin((lat2-lat1)/2)**2 + math.cos(lat1)*math.cos(lat2)*math.sin((lon2-lon1)/2)**2
    return 2*6371.0088*math.asin(math.sqrt(h))


def locate(code, x, y):
    from pyproj import Transformer
    lon, lat = Transformer.from_crs(code, 'EPSG:4326', always_xy=True).transform(x, y)
    return (lon, lat) if math.isfinite(lon) and math.isfinite(lat) else None


def check(code, name, datum, x, y, area, near, near_km):
    """A candidate row when the point lands in the CRS's area of use (0.5 deg slack) and near the hint."""
    place = locate(code, x, y)
    if place is None:
        return None, 'does not project'
    lon, lat = place
    if area is not None and not (area.west-0.5 <= lon <= area.east+0.5 and area.south-0.5 <= lat <= area.north+0.5):
        return None, 'outside the area of use'
    row = {'crs': code, 'name': name, 'datum': datum, 'lonlat': [round(lon, 5), round(lat, 5)]}
    if near is not None:
        row['distance_km'] = round(distance_km(near, place), 1)
        if row['distance_km'] > near_km:
            return None, f'{row["distance_km"]} km from --near'
    return row, None


def readings(x, y, scale, near, near_km):
    """Grid kinds the representative point fits at one drawing-unit scale, with candidate CRS."""
    from pyproj import CRS
    e, n = x*scale, y*scale
    found, rejected = [], []

    def consider(grid, rows):
        kept = []
        for datum, code, name, area in rows:
            row, reason = check(code, name, datum, e, n, area, near, near_km)
            if row:
                kept.append(row)
            else:
                rejected.append({'crs': code, 'reason': reason})
        if kept:
            found.append({'grid': grid, 'candidates': kept})

    if scale == 1 and -180 <= x <= 180 and -90 <= y <= 90:
        consider('geographic', [(CRS(code).name, code, CRS(code).name, CRS(code).area_of_use)
                                for code in ('EPSG:4490', 'EPSG:4326')])
    zone = e // 1e6
    if 0 < n < 6.5e6 and (13 <= zone <= 23 or 25 <= zone <= 45):
        width = 6 if zone <= 23 else 3
        consider(f'gauss-kruger {width}-degree zone {int(zone)} (zone prefix)', gauss_kruger().get((width, int(zone)), []))
    # Without the zone prefix a Gauss-Kruger grid, a UTM zone and a city or project local grid (Shenzhen's,
    # for one) look alike, so no EPSG code is listed for them, --near or not.
    bare = 0 < e < 1e6 and 0 < n < 6.5e6
    if not found and not bare and near is not None and abs(e) <= 20037508.35 and abs(n) <= 20037508.35:
        consider('web mercator', [('WGS 84', 'EPSG:3857', 'WGS 84 / Pseudo-Mercator', None)])
    return found, rejected, bare


def read_reference(path, layer):
    """Extent and CRS of a reference layer; the CRS must be projected in metres to compare plane coordinates."""
    import pyogrio
    from pyproj import CRS
    info = pyogrio.read_info(path, layer=layer, force_total_bounds=True)
    if info.get('crs') is None:
        raise ValueError('--reference has no CRS')
    crs = CRS.from_user_input(info['crs'])
    unit = crs.axis_info[0].unit_name if crs.axis_info else None
    if not crs.is_projected or unit != 'metre':
        raise ValueError(f'--reference must be in a projected CRS in metres; got {crs.name}')
    bounds = info.get('total_bounds')
    if bounds is None or not all(math.isfinite(float(v)) for v in bounds):
        raise ValueError('--reference has no extent')
    authority = crs.to_authority(min_confidence=70)
    return {'name': Path(path).name, 'layer': info.get('layer_name') or layer,
            'crs': ':'.join(authority) if authority else crs.to_wkt(), 'crs_name': crs.name,
            'datum': crs.datum.name if crs.datum else None, 'bounds': [round(float(v), 3) for v in bounds]}


def share(a, b):
    """Fraction of the smaller box inside the larger one, per axis multiplied; a degenerate axis counts as inside."""
    total = 1.0
    for low, high in ((0, 2), (1, 3)):
        start, end = max(a[low], b[low]), min(a[high], b[high])
        if end < start:
            return 0.0
        small = min(a[high]-a[low], b[high]-b[low])
        total *= 1.0 if small <= 0 else (end-start)/small
    return total


@lru_cache(maxsize=None)
def epsg_by_name():
    from pyproj.database import query_crs_info
    return {info.name: f'EPSG:{info.code}' for info in query_crs_info(auth_name='EPSG', pj_types='PROJECTED_CRS')
            if not info.deprecated}


def prefix_twin(name, zone):
    """The same datum's grid with the zone prefix added (zone > 0: reference has it) or removed (zone < 0)."""
    if zone > 0:
        match = re.fullmatch(r'(.+ / (3-degree )?Gauss-Kruger) zone (\d+)', name)
        if not match or int(match[3]) != zone:
            return None
        twin = f'{match[1]} CM {3*zone if match[2] else 6*zone-3}E'
    else:
        match = re.fullmatch(r'(.+ / (3-degree )?Gauss-Kruger) CM (\d+)E', name)
        if not match:
            return None
        cm = int(match[3])
        if (cm if match[2] else cm+3) != (3 if match[2] else 6)*-zone:
            return None
        twin = f'{match[1]} zone {-zone}'
    code = epsg_by_name().get(twin)
    return (code, twin) if code else None


def compare(core, scale, reference):
    """How the drawing's main extent, at one unit scale, sits against the reference extent."""
    box = [v*scale for v in core]
    ref = reference['bounds']
    centre = ((box[0]+box[2])/2, (box[1]+box[3])/2)
    ref_centre = ((ref[0]+ref[2])/2, (ref[1]+ref[3])/2)
    offset = [round(ref_centre[0]-centre[0], 3), round(ref_centre[1]-centre[1], 3)]
    span, ref_span = math.hypot(box[2]-box[0], box[3]-box[1]), math.hypot(ref[2]-ref[0], ref[3]-ref[1])
    row = {'overlap': round(share(box, ref), 3), 'offset': offset,
           'size_ratio': round(span/ref_span, 4) if ref_span > 0 else None}
    if row['overlap'] >= REFERENCE_OVERLAP:
        row['verdict'] = 'extents-overlap'
        return row
    zone = round(offset[0]/1e6)
    shifted = [box[0]+zone*1e6, box[1], box[2]+zone*1e6, box[3]]
    if (13 <= abs(zone) <= 23 or 25 <= abs(zone) <= 45) and share(shifted, ref) >= REFERENCE_OVERLAP:
        row.update(verdict='zone-prefix', zone=abs(zone), drawing_has_prefix=zone < 0)
        twin = prefix_twin(reference['crs_name'], zone)
        if twin:
            row['twin'] = {'crs': twin[0], 'name': twin[1]}
        return row

    def order(v):
        return math.floor(math.log10(abs(v))) if abs(v) >= 1 else 0
    same = all(order(a) == order(b) for a, b in zip(centre, ref_centre))
    row['verdict'] = 'shift-or-rotation' if same else 'different-magnitude'
    return row


def read_gdal(source, encoding=None):
    """GDAL's DXF features with no per-block feature limit; geometries shapely cannot read (polyhedral surfaces
    from polyface meshes) become empty and are counted. Returns (frame, unreadable count, GDAL text warnings)."""
    import warnings as python_warnings
    import geopandas as gpd
    import pyogrio
    import shapely
    from pyogrio.raw import read
    options = {'DXF_FEATURE_LIMIT_PER_BLOCK': '-1', 'DXF_ENCODING': encoding}
    pyogrio.set_gdal_config_options(options)
    try:
        with python_warnings.catch_warnings(record=True) as caught:
            python_warnings.simplefilter('always')
            meta, _, wkb, fields = read(source)
    finally:
        pyogrio.set_gdal_config_options({key: None for key in options})
    geometry = shapely.from_wkb(wkb, on_invalid='ignore')
    unreadable = sum(1 for g, raw in zip(geometry, wkb) if g is None and raw is not None)
    frame = gpd.GeoDataFrame(dict(zip(meta['fields'], fields)), geometry=gpd.GeoSeries(geometry))
    notes = sorted({str(c.message).split('.')[0] for c in caught if 'converted' in str(c.message)})
    return frame, unreadable, notes


def undecodable(doc):
    """Layer names and model-space texts holding bytes the encoding could not decode (ezdxf keeps them as surrogates)."""
    strings = [layer.dxf.name for layer in doc.layers]
    for entity in doc.modelspace():
        kind = entity.dxftype()
        if kind == 'TEXT':
            strings.append(entity.dxf.text)
        elif kind == 'MTEXT':
            strings.append(entity.text)
        elif kind == 'INSERT':
            strings += [a.dxf.text for a in entity.attribs]
    return sum(1 for s in strings if any(0xDC80 <= ord(c) <= 0xDCFF for c in s or ''))


def utf8_dxf(header):
    """DXF R2007 (AC1021) and later is UTF-8 whatever $DWGCODEPAGE says; GDAL follows the code page unless told."""
    return (header.get('acadver') or '') >= 'AC1021'


def choose_encoding(source, names, header, option):
    """The encoding to read with and how it was chosen: --encoding, the header code page, or an automatic retry
    with CP936 then GB18030 when names look like GBK read as Latin-1 (fewest undecodable strings wins)."""
    if option:
        return option, {'declared': header.get('codepage'), 'used': option, 'source': 'option'}
    if utf8_dxf(header):
        return None, {'declared': header.get('codepage'), 'used': 'UTF-8', 'source': 'dxf-version'}
    if not looks_misdecoded(names):
        return None, {'declared': header.get('codepage'), 'used': 'header', 'source': 'header'}
    tried = {}
    try:
        import ezdxf
    except ImportError:
        ezdxf = None
    for candidate in ('CP936', 'GB18030'):
        if ezdxf is None:
            tried[candidate] = None
            break
        tried[candidate] = undecodable(ezdxf.readfile(source, encoding=candidate))
    counted = {k: v for k, v in tried.items() if v is not None}
    used = min(counted, key=lambda k: (counted[k], list(tried).index(k))) if counted else 'CP936'
    return used, {'declared': header.get('codepage'), 'used': used, 'source': 'auto', 'tried': tried}


def looks_misdecoded(values):
    """Names that are GBK bytes read as Latin-1, as when $DWGCODEPAGE is missing or says ANSI_1252."""
    for value in values:
        if not isinstance(value, str) or value.isascii():
            continue
        try:
            raw = value.encode('latin-1')
            if raw.decode('gbk') != value:
                return True
        except UnicodeError:
            continue
    return False


def entity_type(subclasses):
    if isinstance(subclasses, str) and subclasses:
        return subclasses.split(':')[-1]
    return None


def inventory(frame):
    """Per-layer counts, entity and geometry types, text samples and bounds of model-space entities."""
    import numpy as np
    types = frame['SubClasses'].map(entity_type) if 'SubClasses' in frame else None
    text = frame['Text'] if 'Text' in frame else None
    is_text = np.zeros(len(frame), bool)
    if text is not None:
        has_text = text.notna() & (text.astype(str).str.strip() != '')
        if types is not None and types.notna().any():
            is_text = (has_text & types.fillna('').str.contains('Text')).to_numpy()
        else:  # R12 has no subclass markers: TEXT entities come through as points with text
            is_text = (has_text & (frame.geom_type == 'Point')).to_numpy()
    layers = []
    for name, rows in frame.groupby(frame['Layer'].fillna(''), sort=False):
        index = rows.index
        row = {'layer': name, 'entities': len(rows),
               'geometry_types': {k: int(v) for k, v in rows.geom_type.value_counts().items()}}
        if types is not None and types.loc[index].notna().any():
            row['entity_types'] = {k: int(v) for k, v in types.loc[index].value_counts().items()}
        mask = is_text[frame.index.get_indexer(index)]
        if mask.any():
            values = text.loc[index][mask].astype(str).str.strip()
            row['texts'] = int(mask.sum())
            row['text_samples'] = [{'raw': v[:80]} for v in dict.fromkeys(values)][:TEXT_SAMPLES]
        row['bounds'] = [round(float(v), 3) for v in rows.total_bounds]
        layers.append(row)
    return sorted(layers, key=lambda row: -row['entities'])


def clean_text(text):
    """Text with full-width ASCII folded to half-width and whitespace (MTEXT \\P included) collapsed."""
    folded = ''.join(chr(ord(c) - 0xFEE0) if 0xFF01 <= ord(c) <= 0xFF5E else ' ' if c == '　' else c for c in text)
    return ' '.join(folded.split())


def strip_bad_xdata(source, target):
    """Copy a DXF dropping tags below group code 1000 inside XDATA (after a 1001 application tag, before the next
    entity), which LibreDWG writes into some XRECORDs and ezdxf refuses. Returns the number of tags dropped."""
    dropped, in_xdata = 0, False
    with open(source, 'rb') as reader, open(target, 'wb') as writer:
        while True:
            code_line, value_line = reader.readline(), reader.readline()
            if not code_line:
                break
            code = code_line.strip()
            if code == b'0':
                in_xdata = False
            elif code == b'1001':
                in_xdata = True
            elif in_xdata and code.lstrip(b'-').isdigit() and int(code) < 1000:
                dropped += 1
                continue
            writer.write(code_line + value_line)
    return dropped


def read_drawing(path, encoding):
    """(ezdxf document, audit counts); (None, None) when ezdxf is not installed. A drawing with malformed XDATA is
    read from a copy without the offending tags, and the count is reported as audit['xdata_tags_dropped']."""
    import tempfile
    try:
        import ezdxf
        from ezdxf import recover
        from ezdxf.lldxf.const import DXFXDataError
    except ImportError:
        return None, None

    def load(file):
        if encoding:  # recover mode takes the encoding from the header only
            return ezdxf.readfile(file, encoding=encoding), {'errors': 0, 'fixes': 0}
        doc, auditor = recover.readfile(file)
        return doc, {'errors': len(auditor.errors), 'fixes': len(auditor.fixes)}
    try:
        return load(path)
    except DXFXDataError:
        with tempfile.TemporaryDirectory() as folder:
            copy = Path(folder) / 'clean.dxf'
            dropped = strip_bad_xdata(path, copy)
            doc, audit = load(copy)
        return doc, {**audit, 'xdata_tags_dropped': dropped}


def drawing_inventory(doc, bounds):
    """Per-layer DXF entity types counted as CAD counts them (one INSERT per block reference), block names,
    attribute tags, raw and cleaned text, and layer state; bounds come from the expanded GDAL geometry."""
    from collections import Counter
    layers = {}
    for entity in doc.modelspace():
        kind = entity.dxftype()
        if kind == 'LWPOLYLINE':
            kind += ' closed' if entity.closed else ' open'
        elif kind == 'POLYLINE':
            kind += (' polyface' if entity.is_poly_face_mesh else ' mesh' if entity.is_polygon_mesh
                     else ' closed' if entity.is_closed else ' open')
        row = layers.setdefault(entity.dxf.layer, {'layer': entity.dxf.layer, 'entities': 0, 'dxf_types': Counter(),
                                                    'blocks': Counter(), 'attributes': {}, 'text': []})
        row['entities'] += 1
        row['dxf_types'][kind] += 1
        if kind == 'INSERT':
            row['blocks'][entity.dxf.name] += 1
            for attrib in entity.attribs:
                samples = row['attributes'].setdefault(attrib.dxf.tag, [])
                value = clean_text(attrib.dxf.text)
                if value not in samples and len(samples) < TEXT_SAMPLES:
                    samples.append(value)
        elif kind == 'TEXT':
            row['text'].append((entity.dxf.text, clean_text(entity.plain_text())))
        elif kind == 'MTEXT':
            row['text'].append((entity.text, clean_text(entity.plain_text(split=False))))
    table = {layer.dxf.name: layer for layer in doc.layers}
    result = []
    for name, row in layers.items():
        out = {'layer': name, 'entities': row['entities'], 'dxf_types': dict(row['dxf_types'].most_common())}
        if row['blocks']:
            out['blocks'] = dict(row['blocks'].most_common())
        if row['attributes']:
            out['attributes'] = row['attributes']
        if row['text']:
            out['texts'] = len(row['text'])
            out['text_samples'] = [{'raw': raw[:80], 'clean': clean[:80]}
                                   for raw, clean in dict(row['text']).items()][:TEXT_SAMPLES]
        layer = table.get(name)
        state = [flag for flag, on in (('off', layer and layer.is_off()), ('frozen', layer and layer.is_frozen())) if on]
        if state:
            out['state'] = state
        out['bounds'] = bounds.get(name)
        result.append(out)
    unused = sorted(set(table) - set(layers) - {'0', 'Defpoints'})
    layouts = [{'name': layout.name, 'entities': len(layout),
                'viewports': sum(1 for v in layout.query('VIEWPORT') if v.dxf.get('id', 0) != 1)}
               for layout in doc.layouts if layout.name != 'Model']
    return sorted(result, key=lambda row: -row['entities']), unused, layouts


def split_core(frame):
    """Entities more than OUTLIER_SPREAD median distances from the median centre (title blocks at the origin,
    strays) are outliers; returns the core mask and the core's median centre."""
    import numpy as np
    bounds = frame.geometry.bounds
    cx, cy = ((bounds.minx+bounds.maxx)/2).to_numpy(), ((bounds.miny+bounds.maxy)/2).to_numpy()
    distance = np.hypot(cx-np.median(cx), cy-np.median(cy))
    core = distance <= OUTLIER_SPREAD*max(float(np.median(distance)), 1e-9)
    return core, float(np.median(cx[core])), float(np.median(cy[core]))


def plot(frame, core, layers, path, title, font):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    outliers = not core.all()
    figure, axes = plt.subplots(1, 2 if outliers else 1, figsize=(15 if outliers else 9, 8), squeeze=False)
    axes = axes[0]
    shown = [row['layer'] for row in layers[:PLOT_LAYERS]]
    colours = matplotlib.colormaps['tab20']
    names = frame['Layer'].fillna('')
    panels = [(axes[-1], frame[core], 'core extent')] + ([(axes[0], frame, 'full extent')] if outliers else [])
    for axis, part, label in panels:
        other = part[~names.loc[part.index].isin(shown)]
        if len(other):
            other.plot(ax=axis, color='#b0b0b0', linewidth=0.5, markersize=3)
        for i, layer in enumerate(shown):
            rows = part[names.loc[part.index] == layer]
            if len(rows):
                rows.plot(ax=axis, color=colours(i % 20), linewidth=0.7, markersize=4)
        axis.set_title(label)
        axis.set_xlabel('x (drawing units)')
        axis.set_ylabel('y (drawing units)')
        axis.set_aspect('equal', adjustable='datalim')
        axis.ticklabel_format(useOffset=False, style='plain')
        axis.tick_params(labelsize=7)
    if outliers:
        minx, miny, maxx, maxy = frame[core].total_bounds
        axes[0].add_patch(Rectangle((minx, miny), maxx-minx, maxy-miny, fill=False, linestyle='--', edgecolor='black'))
        stray = frame[~core].geometry.centroid
        axes[0].scatter(stray.x, stray.y, marker='x', color='red', s=40, label='outlier', zorder=5)
    handles = [plt.Line2D([], [], color=colours(i % 20), label=layer or '(no layer)') for i, layer in enumerate(shown)]
    if len(layers) > PLOT_LAYERS:
        handles.append(plt.Line2D([], [], color='#b0b0b0', label=f'{len(layers)-PLOT_LAYERS} other layers'))
    figure.legend(handles=handles, loc='center right', fontsize=8, frameon=False)
    figure.suptitle(title, fontsize=11)
    figure.tight_layout(rect=(0, 0, 0.85, 0.95))
    figure.savefig(path, dpi=150)
    plt.close(figure)


def parse_near(text):
    try:
        lon, lat = (float(part) for part in text.split(','))
    except ValueError as error:
        raise ValueError('--near must be LON,LAT in degrees') from error
    if not (-180 <= lon <= 180 and -90 <= lat <= 90):
        raise ValueError('--near must be LON,LAT in degrees')
    return lon, lat


def inspect(args):
    import pyogrio
    source = Path(args.source)
    if source.suffix.lower() == '.dwg':
        raise ValueError('DWG is not read; convert it with cad.py convert, or save it as DXF in AutoCAD / ZWCAD')
    if source.suffix.lower() != '.dxf' or not source.is_file():
        raise ValueError('Source must be an existing .dxf file')
    near = parse_near(args.near) if args.near else None
    header = read_header(source)
    check_vector(source)
    frame, unreadable, gdal_notes = read_gdal(source, args.encoding or ('UTF-8' if utf8_dxf(header) else None))
    warnings = []
    names = [*frame['Layer'].dropna().unique()[:200], *(frame['Text'].dropna().unique()[:200] if 'Text' in frame else [])]
    encoding, encoding_report = choose_encoding(source, names, header, args.encoding)
    if encoding_report['source'] == 'auto':
        frame, unreadable, gdal_notes = read_gdal(source, encoding)
        warnings.append(f'Layer or text names looked like GBK read as Latin-1 ($DWGCODEPAGE = {header.get("codepage")}); '
                        f'read again as {encoding} (pass --encoding to choose)')
    warnings += [f'GDAL: {note}' for note in gdal_notes]
    if unreadable:
        warnings.append(f'{unreadable} polyface meshes (GDAL returns them as polyhedral surfaces) are left out of the geometry')
    empty = frame.geometry.isna() | frame.geometry.is_empty
    paper = frame['PaperSpace'].eq(True) if 'PaperSpace' in frame else empty & False
    if paper.any():
        warnings.append(f'{int(paper.sum())} paper-space entities left out of extent and CRS checks')
    model = frame[~empty & ~paper]
    if model.empty:
        raise ValueError('No model-space geometry in the drawing')
    core, x, y = split_core(model)
    try:
        doc, audit = read_drawing(source, encoding)
    except Exception as error:  # ezdxf's own errors vary by damage; GDAL already read the drawing
        doc, audit = None, None
        warnings.append(f'ezdxf could not read the drawing ({type(error).__name__}: {error}); layer inventory falls back to GDAL')
    if doc is not None:
        import ezdxf
        bounds = {name: [round(float(v), 3) for v in rows.total_bounds] for name, rows in model.groupby(model['Layer'].fillna(''))}
        layers, unused, layouts = drawing_inventory(doc, bounds)
        reader = {'inventory': f'ezdxf {ezdxf.__version__}', 'missing': []}
        encoding_report['undecodable_strings'] = undecodable(doc)
        if encoding_report['undecodable_strings']:
            warnings.append(f'{encoding_report["undecodable_strings"]} layer names or texts hold bytes that do not decode as '
                            f'{encoding_report["used"] if encoding_report["used"] != "header" else header.get("codepage")}')
        if audit and audit.get('xdata_tags_dropped'):
            warnings.append(f'Read without {audit["xdata_tags_dropped"]} malformed XDATA tags (group codes below 1000 after an '
                            'application name, as LibreDWG writes them); entities are unaffected')
        if audit and audit['errors']:
            warnings.append(f'ezdxf recovered the drawing with {audit["errors"]} structure errors and {audit["fixes"]} fixes')
    else:
        layers, unused, layouts = inventory(model), None, None
        reader = {'inventory': 'gdal', 'missing': EZDXF_ONLY}
    reader['geometry'] = f'gdal {pyogrio.__gdal_version_string__}'
    stray = model[~core]
    if len(stray):
        counts = stray['Layer'].fillna('').value_counts()
        warnings.append(f'{len(stray)} entities lie far outside the main drawing (layers: '
                        + ', '.join(f'{k or "(no layer)"} {v}' for k, v in counts.head(5).items()) + ')')

    scales = [(1.0, 'coordinates as drawn')]
    declared = UNITS.get(header.get('insunits'), (None, None))[1]
    if declared and declared != 1.0:
        scales.append((declared, f'header $INSUNITS = {header["units"]}'))
    tried, rejected, unprefixed = [], [], False
    for scale, basis in scales:
        found, dropped, bare = readings(x, y, scale, near, args.near_km)
        tried.append({'scale_to_metre': scale, 'basis': basis, 'matches': found})
        rejected += dropped
        unprefixed = unprefixed or bare
    reference = None
    if args.reference:
        reference = read_reference(args.reference, args.reference_layer)
        core_box = [float(v) for v in model[core].total_bounds]
        reference['comparisons'] = []
        for reading in tried:
            row = compare(core_box, reading['scale_to_metre'], reference)
            reference['comparisons'].append({'scale_to_metre': reading['scale_to_metre'], **row})
            if row['verdict'] == 'extents-overlap':
                found = (reference['crs'], reference['crs_name'])
            elif row['verdict'] == 'zone-prefix' and 'twin' in row:
                found = (row['twin']['crs'], row['twin']['name'])
            else:
                found = None
            if found and found[0] not in {c['crs'] for m in reading['matches'] for c in m['candidates']}:
                reading['matches'].append({'grid': f'reference layer ({row["verdict"]})', 'candidates': [
                    {'crs': found[0], 'name': found[1], 'datum': reference['datum'], 'basis': 'reference'}]})
    candidates = [row for reading in tried for match in reading['matches'] for row in match['candidates']]
    notes = ['Coordinates cannot distinguish CGCS2000, New Beijing, Xian 1980 and Beijing 1954 grids; '
             'confirm the datum with the data provider or known control points']
    if unprefixed:
        notes.append('Coordinates look like a Gauss-Kruger grid without zone prefix, a UTM zone or a city or project '
                     'local grid, which coordinates alone cannot tell apart; no EPSG code is listed. Compare with a '
                     'layer in a known CRS or fit control points')
    if not candidates and not unprefixed:
        notes.append('No national grid fits; treat as a local or engineering coordinate system that needs control points')
    verdicts = {row['verdict']: row for row in (reference or {}).get('comparisons', [])}
    if 'extents-overlap' in verdicts or 'zone-prefix' in verdicts:
        notes.append('Extents that line up with the reference are consistent with its grid but do not prove it; '
                     'overlay shared features (roads, parcel corners) before assigning a CRS')
    elif 'shift-or-rotation' in verdicts:
        dx, dy = verdicts['shift-or-rotation']['offset']
        notes.append(f'Same magnitude as the reference but offset by about ({dx:.0f}, {dy:.0f}) m: likely a local grid or '
                     'a moved origin; fit control points instead of assigning a CRS')
    elif reference:
        notes.append('Coordinate magnitudes differ from the reference at every scale tried; not the same grid')
    if declared and declared != 1.0 and any(r['matches'] for r in tried[:1]):
        warnings.append(f'Header says {header["units"]}, but coordinates as drawn fit a metre grid')

    crs = {'status': 'candidate' if candidates else 'hold', 'representative_point': [round(x, 3), round(y, 3)],
           'near': list(near) if near else None, 'near_km': args.near_km if near else None,
           'readings': tried, 'rejected': rejected, 'reference': reference, 'notes': notes, 'assigned': None}
    extent = {'full': [round(float(v), 3) for v in model.total_bounds],
              'core': [round(float(v), 3) for v in model[core].total_bounds], 'outliers': int(len(stray))}
    summary = (', '.join(dict.fromkeys(c['crs'] for c in candidates[:4])) + (' …' if len(candidates) > 4 else '')
               if candidates else 'CRS on hold')
    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    stat = source.stat()
    total = sum(row['entities'] for row in layers)
    with bundle(args.output) as stage:
        plot(model, core, layers, stage / 'check.png',
             f'{source.name} · {total} entities · {len(layers)} layers · {summary}', font)
        report = {'report': 'cad-inventory', 'schema_version': 1,
                  'source': {'name': source.name, 'bytes': stat.st_size, 'sha256': digest(source),
                             'modified': datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(timespec='seconds')},
                  'header': header, 'reader': reader, 'encoding': encoding_report,
                  'entities': {'model_space': total, 'geometries': int(len(model)), 'paper_space': int(paper.sum()),
                               'empty': int(empty.sum()),
                               'geometry_types': {k: int(v) for k, v in model.geom_type.value_counts().items()}},
                  'blocks': ('each reference counts once per layer; geometry, extent and the plot expand references'
                             if doc is not None else 'expanded into their geometry; block names are not reported'),
                  'layers': layers, 'unused_layers': unused, 'layouts': layouts,
                  'extent': extent, 'crs': crs, 'warnings': warnings,
                  'files': {'plot': 'check.png', 'record': 'record.json'}}
        write_json(stage / 'cad.json', report)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'source': {'path': str(source.absolute()), 'sha256': report['source']['sha256']},
                                            'reference': ({'path': str(Path(args.reference).absolute()), 'layer': reference['layer'],
                                                           'sha256': digest(args.reference) if Path(args.reference).is_file() else None}
                                                          if reference else None),
                                            'files': {'cad.json': digest(stage / 'cad.json'), 'check.png': digest(stage / 'check.png')},
                                            'warnings': warnings,
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['Geometry comes from GDAL with block references expanded',
                                                            'CRS candidates come from coordinate magnitude and EPSG areas of use only',
                                                            'No CRS is assigned and no geometry is written']})
    return report


CURVES = {'LINE', 'LWPOLYLINE', 'POLYLINE', 'ARC', 'CIRCLE', 'ELLIPSE', 'SPLINE'}
MAX_BLOCK_DEPTH = 16
# Output layers of raw.gpkg: name -> (geometry type, fields and dtypes)
EXTRACT_LAYERS = {
    'lines': ('LineString', {'layer': str, 'dxf_type': str, 'handle': str, 'closed': bool, 'block': str, 'source_handle': str}),
    'closed_polygons': ('Polygon', {'layer': str, 'dxf_type': str, 'handle': str, 'valid': bool, 'block': str, 'source_handle': str}),
    'texts': ('Point', {'text_raw': str, 'text_clean': str, 'layer': str, 'dxf_type': str, 'tag': str, 'height': float,
                        'rotation': float, 'anchor': str, 'arrows': int, 'leader_handle': str, 'handle': str, 'block': str,
                        'source_handle': str}),
    'blocks': ('Point', {'layer': str, 'block': str, 'handle': str, 'rotation': float, 'xscale': float, 'yscale': float,
                         'attributes': str}),
    # One row per leader arrow; text_handle is the text it labels, linked tells whether that text was extracted.
    'leaders': ('Point', {'layer': str, 'dxf_type': str, 'handle': str, 'text_handle': str, 'linked': bool}),
    # The area each HATCH fills; color is the ACI number (256 = BYLAYER, 0 = BYBLOCK), true_color #rrggbb when set.
    'hatches': ('MultiPolygon', {'layer': str, 'dxf_type': str, 'handle': str, 'pattern': str, 'solid': bool, 'color': int,
                                 'true_color': str, 'hatch_style': int, 'loops': int, 'area_m2': float, 'valid': bool,
                                 'block': str, 'source_handle': str}),
}
HATCH_STYLES = {0: 'normal', 1: 'outer', 2: 'ignore'}


def fill_area(loops, style):
    """(area a HATCH fills, valid) from its boundary loops, nested by containment: the normal style fills even-odd
    (islands are holes, islands in them filled again), outer stops at the first island level, ignore keeps only the
    outermost loops. Loops that are not simple polygons are returned as drawn, flagged invalid, never repaired."""
    import shapely
    rings = [shapely.Polygon(loop) for loop in loops]
    if not all(ring.is_valid and ring.area > 0 for ring in rings):
        return shapely.MultiPolygon([ring for ring in rings if not ring.is_empty]), False
    depth = [sum(1 for j, other in enumerate(rings) if j != i and other.covers(ring) and (other.area > ring.area or j < i))
             for i, ring in enumerate(rings)]
    limit = {1: 1, 2: 0}.get(style, len(rings))
    area = shapely.GeometryCollection()
    for ring, level in zip(rings, depth):
        if level <= limit:
            area = shapely.symmetric_difference(area, ring) if style != 2 else shapely.union(area, ring)
    parts = [part for part in shapely.get_parts(area) if part.geom_type == 'Polygon' and not part.is_empty]
    return shapely.MultiPolygon(parts), True


class Extraction:
    """Rows per output layer, walked from model-space entities; block references are kept or exploded."""

    def __init__(self, tolerance, explode):
        from collections import Counter
        self.tolerance, self.explode = tolerance, explode
        self.rows = {name: [] for name in EXTRACT_LAYERS}
        self.read, self.skipped = {}, Counter()

    def text(self, entity, layer, block, source):
        kind = entity.dxftype()
        if kind == 'MTEXT':
            raw, clean = entity.text, clean_text(entity.plain_text(split=False))
            point, height, rotation, anchor = entity.dxf.insert, entity.dxf.char_height, entity.get_rotation(), 'insert'
        else:  # TEXT and ATTRIB: aligned text is placed by its alignment point
            raw, clean = entity.dxf.text, clean_text(entity.plain_text())
            aligned = (entity.dxf.get('halign', 0) or entity.dxf.get('valign', 0)) and entity.dxf.hasattr('align_point')
            point = entity.dxf.align_point if aligned else entity.dxf.insert
            height, rotation, anchor = entity.dxf.height, entity.dxf.get('rotation', 0.0), 'align' if aligned else 'insert'
            point = entity.ocs().to_wcs(point)
        self.rows['texts'].append({'text_raw': raw, 'text_clean': clean, 'layer': layer, 'dxf_type': kind,
                                   'tag': entity.dxf.tag if kind == 'ATTRIB' else None, 'height': float(height),
                                   'rotation': float(rotation), 'anchor': anchor, 'arrows': 0, 'leader_handle': None,
                                   'handle': entity.dxf.get('handle'), 'block': block, 'source_handle': source,
                                   'coords': [tuple(point)]})

    def multileader(self, entity, layer, block, source):
        """An MLEADER's text placed at its arrow when it has exactly one; every arrow also goes to leaders."""
        from ezdxf.tools.text import plain_mtext
        context, handle = entity.context, entity.dxf.get('handle')
        arrows = [tuple(line.vertices[0]) for leader in context.leaders for line in leader.lines if line.vertices]
        content = context.mtext
        for arrow in arrows:
            self.rows['leaders'].append({'layer': layer, 'dxf_type': 'MULTILEADER', 'handle': handle,
                                         'text_handle': handle if content is not None else None,
                                         'linked': content is not None, 'coords': [arrow]})
        if content is None:  # block content: nothing to read as text
            self.skipped['MULTILEADER (block content)'] += 1
            return
        single = len(arrows) == 1
        self.rows['texts'].append({'text_raw': content.default_content, 'text_clean': clean_text(plain_mtext(content.default_content)),
                                   'layer': layer, 'dxf_type': 'MULTILEADER', 'tag': None,
                                   'height': float(getattr(context, 'char_height', 0.0) or 0.0), 'rotation': 0.0,
                                   'anchor': 'leader' if single else 'insert', 'arrows': len(arrows),
                                   'leader_handle': handle, 'handle': handle, 'block': block, 'source_handle': source,
                                   'coords': [arrows[0] if single else tuple(content.insert)]})

    def link_leaders(self):
        """Move each text labelled by exactly one LEADER to that leader's arrow."""
        texts = {row['handle']: row for row in self.rows['texts'] if row['handle'] and row['dxf_type'] != 'MULTILEADER'}
        pointing = {}
        for row in self.rows['leaders']:
            if row['dxf_type'] == 'LEADER' and row['text_handle'] in texts:
                row['linked'] = True
                pointing.setdefault(row['text_handle'], []).append(row)
        for text_handle, leaders in pointing.items():
            text = texts[text_handle]
            text['arrows'] = len(leaders)
            if len(leaders) == 1:
                text.update(anchor='leader', leader_handle=leaders[0]['handle'], coords=leaders[0]['coords'])

    def walk(self, entity, layer, block=None, source=None, depth=0):
        from ezdxf import path
        kind = entity.dxftype()
        handle = entity.dxf.get('handle')
        if kind in CURVES:
            try:
                points = [tuple(p) for p in path.make_path(entity).flattening(distance=self.tolerance)]
            except (TypeError, ValueError):  # polyface meshes and degenerate curves have no 2D path
                self.skipped[f'{kind} (no path)'] += 1
                return
            if len(points) < 2:
                self.skipped[f'{kind} (degenerate)'] += 1
                return
            closed = (entity.closed if kind == 'LWPOLYLINE' else entity.is_closed if kind == 'POLYLINE' else False)
            self.rows['lines'].append({'layer': layer, 'dxf_type': kind, 'handle': handle, 'closed': bool(closed),
                                       'block': block, 'source_handle': source, 'coords': points})
            if closed:
                if len({p[:2] for p in points}) < 3:
                    self.skipped[f'{kind} closed (fewer than 3 vertices)'] += 1
                else:
                    self.rows['closed_polygons'].append({'layer': layer, 'dxf_type': kind, 'handle': handle,
                                                         'block': block, 'source_handle': source, 'coords': points})
        elif kind in ('TEXT', 'MTEXT'):
            self.text(entity, layer, block, source)
        elif kind == 'LEADER':
            vertices = list(entity.vertices)
            if not vertices:
                self.skipped['LEADER (no vertices)'] += 1
                return
            self.rows['leaders'].append({'layer': layer, 'dxf_type': kind, 'handle': handle,
                                         'text_handle': entity.dxf.get('annotation_handle'), 'linked': False,
                                         'coords': [tuple(vertices[0])]})
        elif kind == 'MULTILEADER':
            self.multileader(entity, layer, block, source)
        elif kind == 'HATCH':
            loops = []
            try:
                for item in path.from_hatch(entity):
                    points = [tuple(p) for p in item.flattening(distance=self.tolerance)]
                    if len({p[:2] for p in points}) >= 3:
                        loops.append(points)
            except (TypeError, ValueError):
                loops = []
            if not loops:
                self.skipped['HATCH (no boundary)'] += 1
                return
            rgb = entity.rgb
            self.rows['hatches'].append({'layer': layer, 'dxf_type': kind, 'handle': handle,
                                         'pattern': entity.dxf.get('pattern_name'), 'solid': bool(entity.dxf.get('solid_fill', 0)),
                                         'color': int(entity.dxf.get('color', 256)),
                                         'true_color': '#%02x%02x%02x' % tuple(rgb) if rgb else None,
                                         'hatch_style': int(entity.dxf.get('hatch_style', 0)), 'loops': len(loops),
                                         'block': block, 'source_handle': source, 'loop_points': loops,
                                         'coords': [p for loop in loops for p in loop]})
        elif kind == 'INSERT':
            name = entity.dxf.name
            for attrib in entity.attribs:
                self.text(attrib, layer, name, source or handle)
            if not self.explode:
                self.rows['blocks'].append({'layer': layer, 'block': name, 'handle': handle,
                                            'rotation': float(entity.dxf.get('rotation', 0.0)),
                                            'xscale': float(entity.dxf.get('xscale', 1.0)), 'yscale': float(entity.dxf.get('yscale', 1.0)),
                                            'attributes': json.dumps({a.dxf.tag: a.dxf.text for a in entity.attribs}, ensure_ascii=False),
                                            'coords': [tuple(entity.ocs().to_wcs(entity.dxf.insert))]})
                return
            if depth >= MAX_BLOCK_DEPTH:
                raise ValueError(f'Block nesting deeper than {MAX_BLOCK_DEPTH} at {name}')
            for part in entity.virtual_entities():
                # Entities on layer 0 inside a block take the reference's layer, as in CAD.
                own = part.dxf.get('layer', '0')
                self.walk(part, layer if own == '0' else own, f'{block}/{name}' if block else name, source or handle, depth+1)
        else:
            self.skipped[kind] += 1


def extract(args):
    """Model-space entities of chosen layers to raw.gpkg (lines, closed_polygons, texts, blocks), in metres, no CRS."""
    import warnings as python_warnings
    import geopandas as gpd
    import numpy as np
    import pandas as pd
    import pyogrio
    import shapely
    from collections import Counter
    source = Path(args.source)
    if source.suffix.lower() != '.dxf' or not source.is_file():
        raise ValueError('Source must be an existing .dxf file')
    if args.arc_tolerance <= 0:
        raise ValueError('--arc-tolerance must be positive')
    header = read_header(source)
    doc, audit = read_drawing(source, args.encoding)
    if doc is None:
        raise ValueError('cad.py extract needs ezdxf (pip install ezdxf); GDAL keeps neither handles nor block attributes')
    import ezdxf
    warnings = []
    names = [layer.dxf.name for layer in doc.layers]
    encoding, encoding_report = choose_encoding(source, names, header, args.encoding)
    if encoding_report['source'] == 'auto':
        doc, audit = read_drawing(source, encoding)
        warnings.append(f'Layer names looked like GBK read as Latin-1 ($DWGCODEPAGE = {header.get("codepage")}); '
                        f'read again as {encoding} (pass --encoding to choose)')
    encoding_report['undecodable_strings'] = undecodable(doc)
    if audit and audit.get('xdata_tags_dropped'):
        warnings.append(f'Read without {audit["xdata_tags_dropped"]} malformed XDATA tags; entities are unaffected')
    if header.get('handles') is not True:
        warnings.append('The drawing may carry no entity handles (R12 or no readable header): handle fields can be '
                        'assigned by ezdxf on reading and then cannot be looked up in CAD')
    if audit and audit['errors']:
        warnings.append(f'ezdxf recovered the drawing with {audit["errors"]} structure errors and {audit["fixes"]} fixes')

    if args.units == 'header':
        name, scale = UNITS.get(header.get('insunits'), (header.get('units') or 'missing', None))
        if not scale:
            raise ValueError(f'Header $INSUNITS is {name}; pass --units m or --units mm')
        basis = f'header $INSUNITS = {name}'
    else:
        scale, basis = {'m': 1.0, 'mm': 0.001}[args.units], f'--units {args.units}'

    model = doc.modelspace()
    present = Counter(entity.dxf.layer for entity in model)
    wanted = [name.strip() for name in args.layers.split(',')] if args.layers else sorted(present)
    unknown = [name for name in wanted if name not in present]
    if unknown:
        raise ValueError(f'No model-space entities on layers {", ".join(unknown)}; layers with entities: {", ".join(sorted(present))}')
    run = Extraction(args.arc_tolerance/scale, args.include_blocks)
    for entity in model:
        layer = entity.dxf.layer
        if layer in wanted:
            run.read.setdefault(layer, Counter())[entity.dxftype()] += 1
            run.walk(entity, layer)

    run.link_leaders()
    coords = [p for rows in run.rows.values() for row in rows for p in row['coords']]
    if not coords:
        raise ValueError('Nothing to extract from the chosen layers')
    xs, ys, zs = (np.array(axis, dtype=float) for axis in zip(*coords))
    if scale != 1.0 and args.units == 'header':
        x, y = float(np.median(xs)), float(np.median(ys))
        if readings(x, y, 1.0, None, 50.0)[0] and not readings(x, y, scale, None, 50.0)[0]:
            raise ValueError(f'Header says {header["units"]}, but coordinates as drawn fit a metre grid; pass --units m')
    lifted = {name: sum(1 for row in rows if any(abs(p[2]) > 1e-9 for p in row['coords'])) for name, rows in run.rows.items()}
    z = {'features': sum(lifted.values()), 'by_layer': {k: v for k, v in lifted.items() if v},
         'min': float(zs.min())*scale, 'max': float(zs.max())*scale, 'dropped': bool(args.drop_z and any(lifted.values()))}
    if z['features'] and not args.drop_z:
        raise ValueError(f'{z["features"]} features have non-zero Z ({z["min"]:g} to {z["max"]:g} m); pass --drop-z to flatten them')

    frames = {}
    for name, (geometry_type, fields) in EXTRACT_LAYERS.items():
        rows = run.rows[name]
        points = [[(p[0]*scale, p[1]*scale) for p in row['coords']] for row in rows]
        if name == 'hatches':
            filled = [fill_area([[(x*scale, y*scale) for x, y, *_ in loop] for loop in row['loop_points']], row['hatch_style'])
                      for row in rows]
            for row, (shape, valid) in zip(rows, filled):
                row.update(area_m2=round(float(shape.area), 3) if valid else None, valid=valid)
            geometry = [shape for shape, _ in filled]
        else:
            maker = {'LineString': shapely.linestrings, 'Polygon': shapely.polygons, 'Point': shapely.points}[geometry_type]
            geometry = [maker(p[0] if geometry_type == 'Point' else p) for p in points]
        table = pd.DataFrame({field: pd.Series([row.get(field) for row in rows], dtype=object if kind is str else kind)
                              for field, kind in fields.items() if field != 'valid' or name == 'hatches'})
        if name == 'closed_polygons':
            table['valid'] = pd.Series(shapely.is_valid(np.array(geometry, dtype=object)) if rows else [], dtype=bool)
        if name == 'texts':
            table['height'] = table['height']*scale
        frames[name] = gpd.GeoDataFrame(table, geometry=gpd.GeoSeries(geometry, dtype='geometry'))
    paper = sum(len(layout) for layout in doc.layouts if layout.name != 'Model')
    if paper:
        warnings.append(f'{paper} paper-space entities not extracted')
    invalid = int((~frames['closed_polygons']['valid']).sum()) if len(frames['closed_polygons']) else 0
    if invalid:
        warnings.append(f'{invalid} closed polygons are invalid (valid = false); kept as drawn, not repaired')
    hatches = frames['hatches']
    invalid = int((~hatches['valid'].astype(bool)).sum()) if len(hatches) else 0
    if invalid:
        warnings.append(f'{invalid} hatches have boundary loops that are not simple polygons (valid = false); '
                        'their loops are kept as drawn, not filled or repaired')
    empty = int((hatches['valid'].astype(bool) & (hatches['area_m2'] == 0)).sum()) if len(hatches) else 0
    if empty:
        warnings.append(f'{empty} hatches fill no area: their loops cancel out under the hatch style (e.g. a loop drawn twice)')

    stat = source.stat()
    with bundle(args.output) as stage:
        target = stage / 'raw.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message="'crs' was not provided")
            for name, frame in frames.items():
                pyogrio.write_dataframe(frame, target, layer=name, driver='GPKG', geometry_type=EXTRACT_LAYERS[name][0])
        counts = {name: len(frame) for name, frame in frames.items()}
        back = {name: int(pyogrio.read_info(target, layer=name, force_feature_count=True)['features']) for name in frames}
        if back != counts:
            raise ValueError(f'Readback feature counts {back} differ from {counts}; nothing was published')
        manifest = {'manifest': 'cad-extract', 'schema_version': 1,
                    'source': {'name': source.name, 'bytes': stat.st_size, 'sha256': digest(source)},
                    'header': header, 'reader': f'ezdxf {ezdxf.__version__}', 'encoding': encoding_report,
                    'parameters': {'layers': wanted, 'units': args.units, 'scale_to_metre': scale, 'units_basis': basis,
                                   'arc_tolerance_m': args.arc_tolerance, 'include_blocks': args.include_blocks,
                                   'drop_z': args.drop_z, 'encoding': args.encoding},
                    'read': {layer: dict(types.most_common()) for layer, types in sorted(run.read.items())},
                    'layers': counts, 'skipped': dict(run.skipped.most_common()), 'paper_space': paper, 'z': z,
                    'crs': None, 'warnings': warnings,
                    'output': {'path': target.name, 'sha256': digest(target)}}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad extract', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'source': {'path': str(source.absolute()), 'sha256': manifest['source']['sha256']},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256']},
                                            'warnings': warnings,
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['Coordinates are only scaled to metres; no CRS is set and nothing is moved',
                                                            'Curves are flattened so no chord is farther than --arc-tolerance from the curve',
                                                            'Closed polygons are taken as drawn; invalid ones are flagged, not repaired',
                                                            'A hatch is the area its boundary loops fill under its hatch style; '
                                                            'its pattern and colour are kept, not rendered']})
    return manifest


# A plane in metres with no location, so topo.py can measure faces of a drawing that has no CRS yet.
DRAWING_PLANE = '+proj=tmerc +lat_0=0 +lon_0=0 +k=1 +x_0=0 +y_0=0 +ellps=GRS80 +units=m +no_defs'
FACE_LAYERS = {
    'faces': ('Polygon', ('face_id', 'area_m2', 'status', 'reason', 'valid', 'closed_gap', 'source_handles', 'source_layers')),
    'qa_gaps': ('LineString', ('kind', 'length_m', 'handles_a', 'handles_b', 'closed')),
    'qa_dangles': ('LineString', ('kind', 'length_m', 'handles')),
    'qa_topology': ('Unknown', ('check', 'detail', 'area_m2', 'face_a', 'face_b')),
    'added_lines': ('LineString', ('length_m', 'handles_a', 'handles_b')),
    'excluded': ('LineString', ('layer', 'handle', 'block', 'rule')),
}
EXCLUDE_KEYS = ('layers', 'blocks', 'boxes')


def load_exclusions(path):
    """Exclusion rules from YAML: layers and blocks are fnmatch patterns on CAD layer and block names, boxes are
    [minx, miny, maxx, maxy] in metres."""
    import yaml
    text = Path(path).read_text(encoding='utf-8')
    rules = yaml.safe_load(text) or {}
    if not isinstance(rules, dict) or set(rules) - set(EXCLUDE_KEYS):
        raise ValueError(f'--exclude {path} must be a mapping with only {", ".join(EXCLUDE_KEYS)}')
    for key in ('layers', 'blocks'):
        if not isinstance(rules.get(key, []), list) or not all(isinstance(v, str) and v for v in rules.get(key, [])):
            raise ValueError(f'--exclude {path}: {key} must be a list of name patterns')
    boxes = rules.get('boxes', [])
    if not isinstance(boxes, list) or not all(isinstance(b, list) and len(b) == 4 and all(isinstance(v, (int, float)) for v in b)
                                              and b[0] < b[2] and b[1] < b[3] for b in boxes):
        raise ValueError(f'--exclude {path}: boxes must be [minx, miny, maxx, maxy] lists with min < max')
    return {key: rules.get(key, []) for key in EXCLUDE_KEYS}, text


def exclusions(rows, args, crs):
    """(rule matched by each line or None, report): the first of layer patterns, block patterns (any name in an exploded
    block path), boxes and --exclude-area polygons. A line counts as inside an area only when wholly within it; lines
    crossing an area's edge are kept and counted, never clipped."""
    from fnmatch import fnmatchcase
    import geopandas as gpd
    import shapely
    from pyproj import CRS
    rules, text = load_exclusions(args.exclude) if args.exclude else ({key: [] for key in EXCLUDE_KEYS}, None)
    if args.boundary and any(fnmatchcase(args.boundary, pattern) for pattern in rules['layers']):
        raise ValueError(f'--exclude matches the boundary layer {args.boundary}; the planning boundary cannot be excluded')
    areas = [(f'box:{i}', shapely.box(*b)) for i, b in enumerate(rules['boxes'], 1)]
    area_report = None
    if args.exclude_area:
        path, layer = (args.exclude_area.rsplit(':', 1) if ':' in args.exclude_area and not Path(args.exclude_area).exists()
                       else (args.exclude_area, None))
        check_vector(path, layer)
        polygons = gpd.read_file(path, layer=layer, engine='pyogrio')
        if (polygons.crs is None) != (crs is None) or (crs is not None and not CRS.from_user_input(crs).equals(polygons.crs)):
            raise ValueError(f'--exclude-area must share the lines\' CRS (or both have none); got {polygons.crs} and {crs}')
        polygons = polygons[polygons.geometry.notna() & polygons.geom_type.isin(['Polygon', 'MultiPolygon'])]
        if not len(polygons):
            raise ValueError(f'--exclude-area {args.exclude_area} has no polygons')
        areas += [(f'area:{i}', g) for i, g in zip(polygons.index, polygons.geometry)]
        area_report = {'path': str(Path(path).absolute()), 'layer': layer, 'features': int(len(polygons)),
                       'sha256': digest(path) if Path(path).is_file() else None}
    blocks = rows['block'] if 'block' in rows else [None]*len(rows)
    matched, crossing, used = [], 0, set()
    for layer, block, line in zip(rows['layer'], blocks, rows.geometry):
        rule = next((f'layer:{p}' for p in rules['layers'] if fnmatchcase(str(layer), p)), None)
        names = block.split('/') if isinstance(block, str) and block else []
        rule = rule or next((f'block:{p}' for p in rules['blocks'] if any(fnmatchcase(n, p) for n in names)), None)
        if rule is None:
            for label, area in areas:
                if area.covers(line):
                    rule = label
                    break
            else:
                crossing += any(area.intersects(line) for _, area in areas)
        used.add(rule)
        matched.append(rule)
    unmatched = [f'layer:{p}' for p in rules['layers'] if f'layer:{p}' not in used] + \
                [f'block:{p}' for p in rules['blocks'] if f'block:{p}' not in used] + \
                [label for label, _ in areas if label not in used and label.startswith('box:')]
    if rules['blocks'] and 'block' not in rows:
        unmatched.append('blocks: the line layer has no block field (cad.py extract --include-blocks writes one)')
    report = {'rules': {'path': str(Path(args.exclude).absolute()), 'sha256': digest(args.exclude), **rules} if text else None,
              'areas': area_report, 'crossing_kept': crossing, 'unmatched': unmatched}
    return matched, report


def plot_faces(frames, path, title, font):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    figure, axis = plt.subplots(figsize=(12, 9))
    faces = frames['faces']
    for status, colour in (('pass', '#cfe3f5'), ('hold', '#f6b26b')):
        part = faces[faces['status'] == status]
        if len(part):
            part.plot(ax=axis, facecolor=colour, edgecolor='#4a4a4a', linewidth=0.4)
    if len(frames.get('excluded', [])):
        frames['excluded'].plot(ax=axis, color='#9e9e9e', linewidth=0.5, linestyle='--')
    for name, colour, width in (('qa_dangles', '#8e44ad', 1.0), ('added_lines', '#27ae60', 1.6), ('qa_gaps', '#d62728', 1.6)):
        if len(frames[name]):
            frames[name].plot(ax=axis, color=colour, linewidth=width)
            ends = frames[name].geometry.boundary.explode(index_parts=False)
            axis.scatter(ends.x, ends.y, color=colour, s=14, zorder=5)
    handles = [plt.Rectangle((0, 0), 1, 1, facecolor=c, edgecolor='#4a4a4a', label=l)
               for c, l in (('#cfe3f5', 'face pass'), ('#f6b26b', 'face hold'))]
    handles += [plt.Line2D([], [], color=c, linewidth=2, label=l)
                for c, l in (('#d62728', 'gap'), ('#8e44ad', 'dangle / cut edge'), ('#27ae60', 'added by --close-gaps'))]
    if len(frames.get('excluded', [])):
        handles.append(plt.Line2D([], [], color='#9e9e9e', linestyle='--', label=f'excluded ({len(frames["excluded"])})'))
    axis.legend(handles=handles, loc='center left', bbox_to_anchor=(1.01, 0.5), fontsize=8, frameon=False)
    axis.set_aspect('equal', adjustable='datalim')
    axis.ticklabel_format(useOffset=False, style='plain')
    axis.tick_params(labelsize=7)
    axis.set_title(title, fontsize=10)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def polygonize(args):
    """Snap, node and polygonize extracted CAD lines into candidate faces with gap, dangle and topology checks."""
    import warnings as python_warnings
    import geopandas as gpd
    import numpy as np
    import pandas as pd
    import pyogrio
    import shapely
    from pyproj import CRS
    import _cad_faces as F
    from topo import check_topology
    if not (args.snap > 0 and args.max_gap > args.snap and args.min_area >= 0):
        raise ValueError('Need --snap > 0, --max-gap > --snap and --min-area >= 0')
    source = Path(args.source)
    check_vector(source, args.lines_layer)
    frame = gpd.read_file(source, layer=args.lines_layer, engine='pyogrio')
    if 'layer' not in frame:
        raise ValueError(f'{source}:{args.lines_layer} has no "layer" field; use cad.py extract output')
    if frame.crs is not None:
        crs = CRS.from_user_input(frame.crs)
        unit = crs.axis_info[0].unit_name if crs.axis_info else None
        if not crs.is_projected or unit != 'metre':
            raise ValueError(f'Input CRS must be projected in metres; got {crs.name}')
    frame = frame[frame.geometry.notna() & ~frame.geometry.is_empty].explode(index_parts=False)
    frame = frame[frame.geom_type == 'LineString']
    present = sorted(frame['layer'].dropna().unique())
    wanted = [n.strip() for n in args.cad_layers.split(',')] if args.cad_layers else list(present)
    if args.boundary and args.boundary not in wanted:
        wanted.append(args.boundary)
    unknown = [n for n in wanted if n not in present]
    if unknown:
        raise ValueError(f'No lines on CAD layers {", ".join(unknown)}; layers present: {", ".join(present)}')
    rows = frame[frame['layer'].isin(wanted)].reset_index(drop=True)
    from collections import Counter
    matched, exclusion = exclusions(rows, args, frame.crs)
    dropped = rows[[rule is not None for rule in matched]].assign(rule=[r for r in matched if r is not None])
    rows = rows[[rule is None for rule in matched]].reset_index(drop=True)
    exclusion.update(excluded=int(len(dropped)), by_rule=dict(Counter(dropped['rule']).most_common()),
                     by_layer=dict(Counter(dropped['layer']).most_common()))
    if not len(rows):
        raise ValueError('Every line was excluded; nothing is left to polygonize')
    labels = [h if isinstance(h, str) and h else (s if isinstance(s, str) and s else None)
              for h, s in zip(rows.get('handle', [None]*len(rows)), rows.get('source_handle', [None]*len(rows)))]

    snapped, moves = F.snap([shapely.get_coordinates(g) for g in rows.geometry], args.snap)
    lines = F.as_lines(snapped)
    kept = shapely.length(lines) > 0
    collapsed = int((~kept).sum())
    lines, labels, layers = lines[kept], [l for l, k in zip(labels, kept) if k], rows['layer'].to_numpy()[kept]
    edges, polys, cuts, dangles, invalid_rings = F.faces(lines)
    found = F.gaps(edges, args.max_gap)
    edge_labels = F.sources(edges, lines, labels)
    edge_layers = F.sources(edges, lines, list(layers))
    added = np.array([g[0] for g in found], dtype=object) if args.close_gaps and found else np.array([], dtype=object)
    if len(added):
        edges, polys, cuts, dangles, invalid_rings = F.faces(np.concatenate([lines, added]))
        edge_labels = F.sources(edges, lines, labels)
        edge_layers = F.sources(edges, lines, list(layers))

    notes = []
    boundary = None
    if args.boundary:
        _, ring_faces, *_ = F.faces(lines[layers == args.boundary])
        if not len(ring_faces):
            raise ValueError(f'Lines on boundary layer {args.boundary} do not close into an area at --snap {args.snap}')
        boundary = shapely.union_all(ring_faces)
        inside = shapely.within(shapely.point_on_surface(polys), boundary) if len(polys) else np.array([], bool)
        outside = int((~inside).sum())
        polys = polys[inside]
    else:
        outside = None
        notes.append('No --boundary: faces are not limited to a planning area, so their total is not a planning-area total')

    edge_tree = shapely.STRtree(edges)
    added_tree = shapely.STRtree(added) if len(added) else None
    face_rows = []
    for polygon in polys:
        valid = bool(shapely.is_valid(polygon))
        area = float(shapely.area(polygon))
        reasons = ([] if valid else ['invalid geometry']) + (['sliver: area below --min-area'] if area < args.min_area else [])
        closed_gap = bool(added_tree is not None and any(
            shapely.distance(shapely.boundary(polygon), shapely.line_interpolate_point(added[k], 0.5, normalized=True)) <= F.ON_EDGE
            for k in added_tree.query(polygon, predicate='intersects')))
        face_rows.append({'area_m2': round(area, 4), 'status': 'hold' if reasons else 'pass', 'reason': '; '.join(reasons) or None,
                          'valid': valid, 'closed_gap': closed_gap,
                          'source_handles': ','.join(F.boundary_labels(polygon, edges, edge_labels, edge_tree)) or None,
                          'source_layers': ','.join(F.boundary_labels(polygon, edges, edge_layers, edge_tree)) or None,
                          'geometry': polygon})
    centre = lambda row: shapely.get_coordinates(shapely.point_on_surface(row['geometry']))[0]
    face_rows.sort(key=lambda row: (-round(centre(row)[1], 3), round(centre(row)[0], 3)))
    for number, row in enumerate(face_rows, 1):
        row['face_id'] = number

    def labelled(index):
        return ','.join(edge_labels[index]) or None
    gap_rows = [{'kind': kind, 'length_m': round(float(shapely.length(line)), 4), 'handles_a': labelled(a), 'handles_b': labelled(b),
                 'closed': bool(len(added)), 'geometry': line} for line, kind, a, b in found]
    loose = [(kind, g) for kind, parts in (('dangle', dangles), ('cut-edge', cuts), ('invalid-ring', invalid_rings)) for g in parts]
    loose_labels = F.sources([g for _, g in loose], lines, labels) if loose else []
    dangle_rows = [{'kind': kind, 'length_m': round(float(shapely.length(g)), 4), 'handles': ','.join(h) or None, 'geometry': g}
                   for (kind, g), h in zip(loose, loose_labels)]
    added_rows = [{'length_m': r['length_m'], 'handles_a': r['handles_a'], 'handles_b': r['handles_b'], 'geometry': r['geometry']}
                  for r in gap_rows] if len(added) else []

    out_crs = frame.crs
    face_frame = gpd.GeoDataFrame(face_rows, columns=[*FACE_LAYERS['faces'][1], 'geometry'], geometry='geometry', crs=out_crs)
    topology = []
    if len(face_frame):
        report = check_topology(face_frame.set_crs(out_crs or DRAWING_PLANE, allow_override=True), id_field='face_id',
                                checks={'gaps', 'overlaps', 'slivers', 'invalid-geometries', 'duplicates'},
                                sliver_threshold_m2=args.min_area)
        topology = [{'check': e.check, 'detail': e.detail, 'area_m2': e.area_m2, 'face_a': e.feature_a, 'face_b': e.feature_b,
                     'geometry': e.geometry} for e in report.errors]
        if out_crs is None:
            notes.append('Topology was measured in the drawing plane (metres, no CRS)')
    frames = {'faces': face_frame}
    for name, rows_ in (('qa_gaps', gap_rows), ('qa_dangles', dangle_rows), ('qa_topology', topology), ('added_lines', added_rows)):
        frames[name] = gpd.GeoDataFrame(rows_, columns=[*FACE_LAYERS[name][1], 'geometry'], geometry='geometry', crs=out_crs)
    frames['excluded'] = gpd.GeoDataFrame(dropped.reindex(columns=[*FACE_LAYERS['excluded'][1], 'geometry']), geometry='geometry',
                                          crs=out_crs)
    if exclusion['unmatched']:
        notes.append(f'Exclusion rules that matched no line: {", ".join(exclusion["unmatched"])}')
    if exclusion['crossing_kept']:
        notes.append(f'{exclusion["crossing_kept"]} lines cross an exclusion area\'s edge and were kept, not clipped')
    counts = {name: len(f) for name, f in frames.items()}
    hold = int((face_frame['status'] == 'hold').sum()) if len(face_frame) else 0
    total = float(face_frame['area_m2'].sum()) if len(face_frame) else 0.0
    coverage = {'faces_area_m2': round(total, 3), 'boundary_area_m2': round(float(shapely.area(boundary)), 3) if boundary is not None else None}
    if boundary is not None and coverage['boundary_area_m2']:
        coverage['share_of_boundary'] = round(total/coverage['boundary_area_m2'], 6)
    moved = [m[2] for m in moves]

    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    stat = source.stat()
    with bundle(args.output) as stage:
        target = stage / 'faces.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message="'crs' was not provided")
            for name, f in frames.items():
                pyogrio.write_dataframe(f, target, layer=name, driver='GPKG', geometry_type=FACE_LAYERS[name][0])
        back = {name: int(pyogrio.read_info(target, layer=name, force_feature_count=True)['features']) for name in frames}
        if back != counts:
            raise ValueError(f'Readback feature counts {back} differ from {counts}; nothing was published')
        plot_faces(frames, stage / 'check.png',
                   f'{source.name} · {counts["faces"]} faces ({hold} hold) · {counts["qa_gaps"]} gaps · {counts["qa_dangles"]} dangles', font)
        manifest = {'manifest': 'cad-polygonize', 'schema_version': 1,
                    'source': {'name': source.name, 'layer': args.lines_layer, 'bytes': stat.st_size, 'sha256': digest(source),
                               'crs': CRS.from_user_input(out_crs).to_string() if out_crs is not None else None},
                    'parameters': {'cad_layers': wanted, 'boundary': args.boundary, 'snap_m': args.snap, 'max_gap_m': args.max_gap,
                                   'min_area_m2': args.min_area, 'close_gaps': args.close_gaps},
                    'lines': {'used': int(len(rows)), 'collapsed_by_snap': collapsed, 'edges_after_noding': int(len(edges))},
                    'snap': {'endpoints_merged': sum(1 for m in moves if m[0] == 'endpoint'),
                             'endpoints_moved_to_line': sum(1 for m in moves if m[0] == 'to-line'),
                             'largest_move_m': round(max(moved), 6) if moved else 0.0},
                    'layers': counts, 'faces': {'pass': counts['faces']-hold, 'hold': hold, 'outside_boundary': outside},
                    'exclusion': exclusion, 'area': coverage, 'notes': notes,
                    'output': {'path': target.name, 'sha256': digest(target)}, 'plot': 'check.png'}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad polygonize', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'source': {'path': str(source.absolute()), 'sha256': manifest['source']['sha256']},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256'],
                                                      'check.png': digest(stage / 'check.png')},
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['Coordinates are in metres (cad.py extract output or a projected CRS in metres)',
                                                            'Endpoints within --snap are merged and free ends within --snap are moved onto the nearest line',
                                                            'Gaps up to --max-gap are reported; they are closed only with --close-gaps, and marked',
                                                            'Faces are candidates: no face is dropped for size, small and invalid ones are held',
                                                            'Excluded lines (--exclude, --exclude-area) are left out before snapping and listed in excluded']})
    return manifest


FACE_KEYS = ('face_id', 'parcel_id', 'landuse', 'control_type', 'status', 'layer', 'handle')


def plot_fill_compare(faces, fills, fill_only, face_only, path, title, font):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    figure, axis = plt.subplots(figsize=(12, 9))
    handles = []
    for frame, style, label in ((face_only, {'facecolor': '#f6b26b', 'edgecolor': 'none'}, 'face only'),
                                (fills, {'facecolor': '#9fc5e8', 'edgecolor': '#3d85c6', 'linewidth': 0.3, 'alpha': 0.6}, 'fill'),
                                (fill_only, {'facecolor': '#d62728', 'edgecolor': 'none'}, 'fill only')):
        if len(frame):
            frame.plot(ax=axis, **style)
        handles.append(plt.Rectangle((0, 0), 1, 1, facecolor=style['facecolor'], label=f'{label} ({len(frame)})'))
    if len(faces):
        faces.boundary.plot(ax=axis, color='#4a4a4a', linewidth=0.4)
        handles.append(plt.Line2D([], [], color='#4a4a4a', label=f'face ({len(faces)})'))
    axis.legend(handles=handles, loc='center left', bbox_to_anchor=(1.01, 0.5), fontsize=8, frameon=False)
    axis.set_aspect('equal', adjustable='datalim')
    axis.ticklabel_format(useOffset=False, style='plain')
    axis.tick_params(labelsize=7)
    axis.set_title(title, fontsize=10)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def compare_fills(args):
    """Hatch areas from cad.py extract against faces (polygonize faces, label or xiangyuan parcels): which face each
    fill covers, and the area only one side covers. Nothing is snapped, merged or chosen."""
    import warnings as python_warnings
    from collections import Counter
    import geopandas as gpd
    import pandas as pd
    import pyogrio
    import shapely
    from pyproj import CRS
    faces_path, fills_path = Path(args.faces), Path(args.fills)
    check_vector(faces_path, args.faces_layer)
    check_vector(fills_path, args.fills_layer)
    faces = gpd.read_file(faces_path, layer=args.faces_layer, engine='pyogrio')
    fills = gpd.read_file(fills_path, layer=args.fills_layer, engine='pyogrio')
    if not {'pattern', 'hatch_style', 'valid'} <= set(fills.columns):
        raise ValueError(f'{fills_path}:{args.fills_layer} is not a hatches layer from cad.py extract')
    if (faces.crs is None) != (fills.crs is None) or (faces.crs is not None and not CRS.from_user_input(faces.crs).equals(fills.crs)):
        raise ValueError(f'Faces and fills must share a CRS (or both have none); got {faces.crs} and {fills.crs}')
    if faces.crs is not None:
        crs = CRS.from_user_input(faces.crs)
        if not crs.is_projected or (crs.axis_info[0].unit_name if crs.axis_info else None) != 'metre':
            raise ValueError(f'Input CRS must be projected in metres; got {crs.name}')
    if args.cad_layers:
        wanted = [name.strip() for name in args.cad_layers.split(',')]
        present = set(fills['layer'].dropna())
        unknown = [name for name in wanted if name not in present]
        if unknown:
            raise ValueError(f'No hatches on CAD layers {", ".join(unknown)}; layers present: {", ".join(sorted(present))}')
        fills = fills[fills['layer'].isin(wanted)]
    faces = faces[faces.geometry.notna() & faces.geom_type.isin(['Polygon', 'MultiPolygon'])].copy()
    held = 0
    if args.exclude_hold:
        if 'status' not in faces:
            raise ValueError(f'--exclude-hold needs a status field in {faces_path}:{args.faces_layer}')
        held = int((faces['status'] == 'hold').sum())
        faces = faces[faces['status'] != 'hold']
    if not len(faces):
        raise ValueError(f'{faces_path}:{args.faces_layer} has no polygons to compare')
    faces['face_index'] = faces.index
    fills = fills.reset_index(drop=True)
    fills['fill_index'] = fills.index
    usable_faces = faces[faces.geometry.is_valid]
    usable_fills = fills[fills['valid'].astype(bool) & fills.geometry.notna() & ~fills.geometry.is_empty]
    keys = [k for k in FACE_KEYS if k in faces.columns]

    tree = shapely.STRtree(usable_faces.geometry.to_numpy())
    pairs = tree.query(usable_fills.geometry.to_numpy(), predicate='intersects')
    pieces = shapely.intersection(usable_fills.geometry.to_numpy()[pairs[0]], usable_faces.geometry.to_numpy()[pairs[1]])
    pair_frame = pd.DataFrame({'fill': usable_fills['fill_index'].to_numpy()[pairs[0]],
                               'face': usable_faces['face_index'].to_numpy()[pairs[1]], 'overlap': shapely.area(pieces)})
    pair_frame = pair_frame[pair_frame['overlap'] > 0]
    face_area = faces.geometry.area
    # The best face has the highest IoU: a frame or boundary face around everything overlaps every fill but fits none.
    pair_frame['iou'] = pair_frame['overlap']/(fills.geometry.area.to_numpy()[pair_frame['fill'].to_numpy()]
                                               + face_area.loc[pair_frame['face']].to_numpy() - pair_frame['overlap'])
    by_fill = pair_frame.sort_values('iou', ascending=False).groupby('fill')
    rows = []
    for _, fill in fills.iterrows():
        row = {k: fill[k] for k in fills.columns if k not in ('geometry', 'fill_index')}
        hits = by_fill.get_group(fill['fill_index']) if fill['fill_index'] in by_fill.groups else pair_frame.iloc[:0]
        usable = fill['fill_index'] in usable_fills['fill_index'].values
        if not usable:
            row.update(match='invalid fill' if not fill['valid'] else 'empty fill')
        elif not len(hits):
            row.update(match='no face', fill_only_m2=round(float(fill.geometry.area), 3))
        else:
            best = hits.iloc[0]
            face = faces.loc[best['face']]
            covered = shapely.intersection(fill.geometry, shapely.union_all(faces.loc[hits['face']].geometry.to_numpy())).area
            overlap = float(best['overlap'])
            row.update(match='one face' if len(hits) == 1 else 'several faces', faces=len(hits), face_index=int(best['face']),
                       **{f'face_{k}': face[k] for k in keys}, overlap_m2=round(overlap, 3),
                       iou=round(float(best['iou']), 6),
                       fill_only_m2=round(float(fill.geometry.area - covered), 3))
        rows.append({**row, 'geometry': fill.geometry})
    fill_frame = gpd.GeoDataFrame(pd.DataFrame(rows), geometry='geometry', crs=fills.crs)

    by_face = pair_frame.groupby('face')
    matched = Counter(int(i) for i in fill_frame.get('face_index', pd.Series(dtype=float)).dropna())
    face_rows = []
    for index, face in faces.iterrows():
        hits = by_face.get_group(index) if index in by_face.groups else pair_frame.iloc[:0]
        covered = (shapely.intersection(face.geometry, shapely.union_all(fills.loc[hits['fill']].geometry.to_numpy())).area
                   if len(hits) and face.geometry.is_valid else 0.0)
        face_rows.append({'face_index': index, **{k: face[k] for k in keys}, 'valid': bool(face.geometry.is_valid),
                          'area_m2': round(float(face_area[index]), 3), 'fills': matched.get(index, 0), 'fills_overlapping': len(hits),
                          'fill_cover_m2': round(float(covered), 3), 'face_only_m2': round(float(face_area[index] - covered), 3),
                          'geometry': face.geometry})
    face_frame = gpd.GeoDataFrame(pd.DataFrame(face_rows), geometry='geometry', crs=faces.crs)

    fill_union = shapely.union_all(usable_fills.geometry.to_numpy())
    face_union = shapely.union_all(usable_faces.geometry.to_numpy())

    def parts(shape, side):
        found = [p for p in shapely.get_parts(shape) if p.geom_type == 'Polygon' and p.area > 0]
        return gpd.GeoDataFrame({'side': [side]*len(found), 'area_m2': [round(float(p.area), 6) for p in found]},
                                geometry=gpd.GeoSeries(found, dtype='geometry'), crs=faces.crs)
    fill_only = parts(shapely.difference(fill_union, face_union), 'fill only')
    face_only = parts(shapely.difference(face_union, fill_union), 'face only')
    matches = Counter(fill_frame['match'])
    summary = {'fills': int(len(fills)), 'faces': int(len(faces)), 'matches': dict(matches.most_common()),
               'faces_held_out': held, 'faces_without_fill': int((face_frame['fills'] == 0).sum()),
               'faces_with_several_fills': int((face_frame['fills'] > 1).sum()),
               'invalid_faces': int((~face_frame['valid']).sum()),
               'iou_median': round(float(fill_frame['iou'].median()), 6) if 'iou' in fill_frame and fill_frame['iou'].notna().any() else None,
               'iou_min': round(float(fill_frame['iou'].min()), 6) if 'iou' in fill_frame and fill_frame['iou'].notna().any() else None}
    area = {'fills_m2': round(float(fill_union.area), 3), 'faces_m2': round(float(face_union.area), 3),
            'both_m2': round(float(shapely.intersection(fill_union, face_union).area), 3),
            'fill_only_m2': round(float(fill_only['area_m2'].sum()), 3), 'face_only_m2': round(float(face_only['area_m2'].sum()), 3)}
    warnings = []
    if matches.get('invalid fill'):
        warnings.append(f'{matches["invalid fill"]} invalid hatches are listed but not compared')
    if summary['invalid_faces']:
        warnings.append(f'{summary["invalid_faces"]} invalid faces are listed but not compared')
    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    layers = {'fills': (fill_frame, 'MultiPolygon'), 'faces': (face_frame, 'MultiPolygon'),
              'fill_only': (fill_only, 'Polygon'), 'face_only': (face_only, 'Polygon')}
    with bundle(args.output) as stage:
        target = stage / 'fill_compare.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message="'crs' was not provided")
            for name, (frame, kind) in layers.items():
                pyogrio.write_dataframe(frame, target, layer=name, driver='GPKG', geometry_type=kind,
                                        promote_to_multi=kind == 'MultiPolygon')
        counts = {name: len(frame) for name, (frame, _) in layers.items()}
        back = {name: int(pyogrio.read_info(target, layer=name, force_feature_count=True)['features']) for name in layers}
        if back != counts:
            raise ValueError(f'Readback feature counts {back} differ from {counts}; nothing was published')
        plot_fill_compare(face_frame, fill_frame[fill_frame['match'] != 'invalid fill'], fill_only, face_only, stage / 'check.png',
                          f'{fills_path.name} vs {faces_path.name} · {len(fills)} fills, {len(faces)} faces · '
                          f'fill only {area["fill_only_m2"]:.0f} m², face only {area["face_only_m2"]:.0f} m²', font)
        manifest = {'manifest': 'cad-fill-compare', 'schema_version': 1,
                    'faces': {'name': faces_path.name, 'layer': args.faces_layer, 'sha256': digest(faces_path) if faces_path.is_file() else None},
                    'fills': {'name': fills_path.name, 'layer': args.fills_layer, 'sha256': digest(fills_path) if fills_path.is_file() else None,
                              'cad_layers': args.cad_layers.split(',') if args.cad_layers else None},
                    'crs': str(faces.crs) if faces.crs is not None else None, 'face_keys': keys,
                    'layers': counts, 'summary': summary, 'area': area, 'warnings': warnings,
                    'output': {'path': target.name, 'sha256': digest(target)}, 'plot': 'check.png'}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad compare-fills', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256'],
                                                      'check.png': digest(stage / 'check.png')},
                                            'warnings': warnings,
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['Each fill is matched to the face with the highest IoU; other overlaps are counted, not chosen',
                                                            'fill_only and face_only are measured on the unions, as drawn; nothing is snapped or repaired']})
    return manifest


LABEL_FIXED = ('face_id', 'parcel_id', 'landuse', 'landuse_major', 'road', 'status', 'reason', 'area_m2',
               'source_handles', 'source_layers', 'label_handles')
CONFLICT_FIELDS = ('kind', 'field', 'value', 'text_raw', 'text_clean', 'layer', 'handle', 'face_id')
UNPARSED_FIELDS = ('text_raw', 'text_clean', 'layer', 'tag', 'handle')


def plot_parcels(parcels, conflicts, path, title, font):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    figure, axis = plt.subplots(figsize=(12, 9))
    colours = matplotlib.colormaps['tab10']
    classes = sorted({c for c in parcels['landuse_major'].dropna()}) if len(parcels) else []
    handles = []
    for i, name in enumerate(classes):
        part = parcels[(parcels['landuse_major'] == name) & ~parcels['road']]
        part.plot(ax=axis, facecolor=colours(i % 10), edgecolor='#4a4a4a', linewidth=0.4, alpha=0.6)
        handles.append(plt.Rectangle((0, 0), 1, 1, facecolor=colours(i % 10), alpha=0.6, label=name))
    for mask, colour, label in ((parcels['road'] if len(parcels) else [], '#dddddd', 'road'),
                                ((parcels['landuse_major'].isna() & ~parcels['road']) if len(parcels) else [], '#ffffff', 'no class')):
        part = parcels[mask] if len(parcels) else parcels
        if len(part):
            part.plot(ax=axis, facecolor=colour, edgecolor='#4a4a4a', linewidth=0.4)
            handles.append(plt.Rectangle((0, 0), 1, 1, facecolor=colour, edgecolor='#4a4a4a', label=label))
    held = parcels[parcels['status'] == 'hold'] if len(parcels) else parcels
    if len(held):
        held.boundary.plot(ax=axis, color='#e67e22', linewidth=1.6)
        handles.append(plt.Line2D([], [], color='#e67e22', linewidth=2, label='hold'))
    for _, row in parcels[parcels['parcel_id'].notna()].iterrows() if len(parcels) else []:
        point = row.geometry.representative_point()
        axis.annotate(row['parcel_id'], (point.x, point.y), fontsize=6, ha='center', va='center')
    if len(conflicts):
        points = conflicts[conflicts.geometry.notna()]
        axis.scatter(points.geometry.x, points.geometry.y, marker='x', color='#d62728', s=30, zorder=5)
        handles.append(plt.Line2D([], [], color='#d62728', marker='x', linestyle='', label='label conflict'))
    if handles:
        axis.legend(handles=handles, loc='center left', bbox_to_anchor=(1.01, 0.5), fontsize=8, frameon=False)
    axis.set_aspect('equal', adjustable='datalim')
    axis.ticklabel_format(useOffset=False, style='plain')
    axis.tick_params(labelsize=7)
    axis.set_title(title, fontsize=10)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def label(args):
    """Attach parcel fields from CAD texts to faces by point-in-face, listing every conflict instead of choosing."""
    import warnings as python_warnings
    from collections import Counter
    import geopandas as gpd
    import pandas as pd
    import pyogrio
    import shapely
    from pyproj import CRS
    import _cad_label as L
    rules, rules_text = L.load(args.rules)
    faces_path, texts_path = Path(args.faces), Path(args.texts)
    check_vector(faces_path, args.faces_layer)
    faces = gpd.read_file(faces_path, layer=args.faces_layer, engine='pyogrio')
    texts = gpd.read_file(texts_path, layer=args.texts_layer, engine='pyogrio')
    missing = {'face_id', 'status'} - set(faces.columns)
    if missing:
        raise ValueError(f'{faces_path}:{args.faces_layer} lacks {", ".join(sorted(missing))}; use cad.py polygonize output')
    if {'text_clean', 'layer'} - set(texts.columns):
        raise ValueError(f'{texts_path}:{args.texts_layer} lacks text_clean / layer; use cad.py extract output')
    if faces.crs is not None and texts.crs is not None and not CRS.from_user_input(faces.crs).equals(CRS.from_user_input(texts.crs)):
        raise ValueError('Faces and texts are in different CRS')
    layers = [name for name, _ in pyogrio.list_layers(texts_path)]
    leaders = gpd.read_file(texts_path, layer='leaders', engine='pyogrio') if 'leaders' in layers else None

    conflicts, unparsed, ignored = [], [], Counter()
    found = {}  # face position -> field -> [(value, text row)]
    tree = shapely.STRtree(faces.geometry.to_numpy())

    def conflict(kind, row=None, field=None, value=None, face_id=None, geometry=None):
        conflicts.append({'kind': kind, 'field': field, 'value': value,
                          'text_raw': None if row is None else row['text_raw'], 'text_clean': None if row is None else row['text_clean'],
                          'layer': None if row is None else row['layer'], 'handle': None if row is None else row.get('handle'),
                          'face_id': face_id, 'geometry': geometry if geometry is not None else (None if row is None else row.geometry)})

    for _, row in texts.iterrows():
        relevant, values = rules.classify(row['layer'], row.get('tag'), row['text_clean'])
        if not relevant:
            ignored[row['layer']] += 1
            continue
        if not values:
            unparsed.append({k: row.get(k) for k in UNPARSED_FIELDS} | {'geometry': row.geometry})
            continue
        if int(row.get('arrows') or 0) > 1:
            conflict('leader-ambiguous', row, ', '.join(values), ', '.join(values.values()))
            continue
        hits = tree.query(row.geometry, predicate='intersects')
        if len(hits) == 0:
            for field, value in values.items():
                conflict('unassigned', row, field, value)
            continue
        if len(hits) > 1:
            for field, value in values.items():
                conflict('on-boundary', row, field, value)
            continue
        for field, value in values.items():
            found.setdefault(int(hits[0]), {}).setdefault(field, []).append((value, row))
    if leaders is not None and len(leaders):
        for _, row in leaders[~leaders['linked'].astype(bool)].iterrows():
            conflicts.append({'kind': 'leader-unlinked', 'field': None, 'value': None, 'text_raw': None, 'text_clean': None,
                              'layer': row['layer'], 'handle': row['handle'], 'face_id': None, 'geometry': row.geometry})

    records = []
    for position, face in faces.reset_index(drop=True).iterrows():
        record = {'face_id': int(face['face_id']), 'area_m2': face.get('area_m2'), 'source_handles': face.get('source_handles'),
                  'source_layers': face.get('source_layers'), 'geometry': face.geometry}
        reasons = [face['reason']] if face['status'] == 'hold' and isinstance(face.get('reason'), str) else []
        used = []
        for field in rules.fields:
            hits = found.get(position, {}).get(field, [])
            distinct = list(dict.fromkeys(value for value, _ in hits))
            used += [r.get('handle') for _, r in hits if r.get('handle')]
            if len(distinct) > 1:
                reasons.append(f'multi-value {field}')
                for value, r in hits:
                    conflict('multi-value', r, field, value, record['face_id'])
            record[field] = distinct[0] if len(distinct) == 1 else None
        record['label_handles'] = ','.join(sorted(set(used))) or None
        landuse = record.get('landuse')
        record['landuse_major'] = rules.major(landuse)
        labelled = any(record.get(f) for f in rules.fields)
        record['road'] = rules.is_road(landuse, record['source_layers'], labelled)
        if not record['road'] and 'landuse' in rules.fields and not landuse and not any(r.startswith('multi-value landuse') for r in reasons):
            reasons.append('unknown use')
        record['reason'] = '; '.join(reasons) or None
        record['status'] = 'hold' if reasons else 'pass'
        records.append(record)
    if 'parcel_id' in rules.fields:
        counts = Counter(r['parcel_id'] for r in records if r.get('parcel_id'))
        for record in records:
            if record.get('parcel_id') and counts[record['parcel_id']] > 1:
                record['status'] = 'hold'
                record['reason'] = '; '.join(filter(None, [record['reason'], 'duplicate parcel_id']))
                conflicts.append({'kind': 'duplicate-id', 'field': 'parcel_id', 'value': record['parcel_id'], 'text_raw': None,
                                  'text_clean': None, 'layer': None, 'handle': None, 'face_id': record['face_id'],
                                  'geometry': shapely.point_on_surface(record['geometry'])})

    extra = [f for f in rules.fields if f not in LABEL_FIXED]
    columns = [*LABEL_FIXED, *extra, 'geometry']
    for record in records:
        record.setdefault('parcel_id', None)
        record.setdefault('landuse', None)
    parcels = gpd.GeoDataFrame(pd.DataFrame(records, columns=columns), geometry='geometry', crs=faces.crs)
    conflict_frame = gpd.GeoDataFrame(pd.DataFrame(conflicts, columns=[*CONFLICT_FIELDS, 'geometry']), geometry='geometry', crs=faces.crs)
    unparsed_frame = gpd.GeoDataFrame(pd.DataFrame(unparsed, columns=[*UNPARSED_FIELDS, 'geometry']), geometry='geometry', crs=faces.crs)

    land = parcels[~parcels['road']] if len(parcels) else parcels
    coverage = {'faces': len(parcels), 'road_faces': int(parcels['road'].sum()) if len(parcels) else 0,
                'unknown_use': int(land['reason'].fillna('').str.contains('unknown use').sum()) if len(land) else 0}
    if 'parcel_id' in rules.fields:
        known = land[~land['reason'].fillna('').str.contains('unknown use')] if len(land) else land
        with_id = int(known['parcel_id'].notna().sum()) if len(known) else 0
        coverage.update(non_road_known_use=len(known), with_parcel_id=with_id,
                        parcel_id_share=round(with_id/len(known), 4) if len(known) else None)
    counts = {'parcels': len(parcels), 'label_conflicts': len(conflict_frame), 'unparsed_text': len(unparsed_frame)}
    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    with bundle(args.output) as stage:
        target = stage / 'parcels.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message="'crs' was not provided")
            for name, frame, kind in (('parcels', parcels, 'Polygon'), ('label_conflicts', conflict_frame, 'Point'),
                                      ('unparsed_text', unparsed_frame, 'Point')):
                pyogrio.write_dataframe(frame, target, layer=name, driver='GPKG', geometry_type=kind)
        back = {name: int(pyogrio.read_info(target, layer=name, force_feature_count=True)['features']) for name in counts}
        if back != counts:
            raise ValueError(f'Readback feature counts {back} differ from {counts}; nothing was published')
        hold = int((parcels['status'] == 'hold').sum()) if len(parcels) else 0
        plot_parcels(parcels, conflict_frame, stage / 'check.png',
                     f'{faces_path.name} · {len(parcels)} faces ({hold} hold) · {len(conflict_frame)} conflicts', font)
        (stage / 'rules.yaml').write_text(rules_text, encoding='utf-8')
        manifest = {'manifest': 'cad-label', 'schema_version': 1,
                    'faces': {'name': faces_path.name, 'layer': args.faces_layer, 'sha256': digest(faces_path)},
                    'texts': {'name': texts_path.name, 'layer': args.texts_layer, 'sha256': digest(texts_path),
                              'leaders': None if leaders is None else len(leaders)},
                    'rules': {'path': 'rules.yaml', 'sha256': digest(stage / 'rules.yaml'), 'fields': rules.fields},
                    'layers': counts, 'parcels': {'pass': len(parcels)-hold, 'hold': hold},
                    'conflicts': dict(Counter(c['kind'] for c in conflicts).most_common()),
                    'coverage': coverage, 'ignored_text_layers': dict(ignored.most_common()),
                    'output': {'path': target.name, 'sha256': digest(target)}, 'plot': 'check.png'}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad label', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256'],
                                                      'check.png': digest(stage / 'check.png'), 'rules.yaml': manifest['rules']['sha256']},
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['A text belongs to the face its anchor (alignment point or leader arrow) falls in',
                                                            'Several distinct values for one field in a face are a conflict, never a choice',
                                                            'Road faces come only from the road rule; unlabelled non-road faces are held']})
    return manifest


# DXF versions the converters write, as cad.py names them -> (LibreDWG --as, ODA output version)
DXF_VERSIONS = {'r2000': ('r2000', 'ACAD2000'), 'r2004': ('r2004', 'ACAD2004'), 'r2007': ('r2007', 'ACAD2007'),
                'r2010': ('r2010', 'ACAD2010'), 'r2013': ('r2013', 'ACAD2013'), 'r2018': ('r2018', 'ACAD2018')}
ODA_PATHS = ('/Applications/ODAFileConverter.app/Contents/MacOS/ODAFileConverter',
             'C:/Program Files/ODA/ODAFileConverter*/ODAFileConverter.exe', '/usr/bin/ODAFileConverter')


def find_converters():
    """Installed DWG-to-DXF converters: {'oda': path, 'libredwg': path}; nothing is installed here."""
    import glob
    import shutil
    found = {}
    oda = shutil.which('ODAFileConverter') or next((hit for pattern in ODA_PATHS for hit in sorted(glob.glob(pattern))), None)
    if oda:
        found['oda'] = oda
    if shutil.which('dwg2dxf'):
        found['libredwg'] = shutil.which('dwg2dxf')
    return found


def dwg_version(path):
    with open(path, 'rb') as stream:
        magic = stream.read(6).decode('ascii', 'replace')
    if not magic.startswith('AC'):
        raise ValueError(f'{path} does not start like a DWG file (got {magic!r})')
    return magic


XIANGYUAN_KEY = '用地代码'   # 湘源控规 writes each parcel attribute as its own XDATA application; this one marks a parcel
XIANGYUAN_ALIASES = {'parcel_id': '地块编号', 'landuse': '用地代码', 'landuse_name': '用地性质'}
XIANGYUAN_FIXED = ('parcel_id', 'landuse', 'landuse_name', 'control_type', 'unit_id', 'status', 'reason', 'area_m2', 'area_declared',
                   'area_diff', 'number_block', 'filled_from_block', 'fills', 'valid', 'layer', 'handle', 'other_xdata')
# Parcel attribute names seen on 湘源控规 drawings: the land-attribute table of 湘源 7.0 and the Shenzhen submission
# fields some drawings carry instead. Applications outside this list (other vendors' plug-ins) go to other_xdata.
XIANGYUAN_FIELDS = frozenset((
    '地块编号 用地代码 用地名称 用地类型 用地性质 第一位比例 第二位比例 第三位比例 第四位比例 用地主代码 第二位代码 第三位代码 '
    '第四位代码 混合代码 单位名称 单位类别 单位规模 一类兼容 二类兼容 非机车位 机动车位 配套公建 地块标高 总面积 用地面积 道路面积 '
    '绿地面积 建筑面积 容积率 建筑密度 绿地率 建筑限高 建筑限低 最小容积率 最小建筑密度 最大绿地率 主出入口 东面退后 南面退后 '
    '西面退后 北面退后 范围参数 等级参数 地区类型 控制信息 用电容量 开闭所号 电信容量 用水容量 总户数 总人数 人口密度 地价参数 '
    '总地价 备注 用水指标 用电指标 用电系数 电信指标 非机车位指标 机动车位指标 户均面积 户均人数 现状用地 用气指标 用气容量 '
    '交通指标 交通流量 公建配套要求 道路交通要求 景观色彩要求 其它特殊要求 消防防灾要求 绿化控制要求 给水排水要求 电力电信要求 '
    '燃气热力要求 文物保护要求 '
    '规划项目编号 图则名称 图则编号 标准单元编号 子单元编号 地块容积 地块容积描述 配套设施名称 规划要求 管控类型 主要功能代码 区位 '
    '图则编号-地块编号 图则编号&地块编号 汇交阶段 删除标记 要素代码').split())
# Attributes that hold numbers (as strings, like every 湘源 XDATA value); --schema can name more.
XIANGYUAN_NUMERIC = frozenset((
    '用地面积 总面积 道路面积 绿地面积 建筑面积 地块容积 容积率 建筑密度 绿地率 建筑限高 建筑限低 最小容积率 最小建筑密度 最大绿地率 '
    '机动车位 非机车位 总户数 总人数 人口密度 户均面积 户均人数 地块标高 东面退后 南面退后 西面退后 北面退后 '
    '第一位比例 第二位比例 第三位比例 第四位比例').split())
UNIT_FIELDS = ('unit_id', 'labels', 'status', 'reason', 'area_m2', 'duplicate_of', 'layer', 'handle')


def xdata_values(entity, multi=None):
    """{application: 1000-group string} of an entity's XDATA, the way 湘源 stores attributes. An empty string is a
    value, not a missing field. With several 1000 groups the last one wins, as 湘源 itself reads them; `multi`
    (a Counter) counts the applications where that happened."""
    values = {}
    for app, tags in (entity.xdata.data.items() if entity.xdata else ()):
        strings = [tag.value for tag in tags if tag.code == 1000]
        if strings:
            values[app] = strings[-1]
            if multi is not None and len(strings) > 1:
                multi[app] += 1
    return values


def load_schema(path):
    """Extra field knowledge from a user's JSON: numeric_like and field lists extend the built-in ones."""
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    if not isinstance(data, dict):
        raise ValueError(f'--schema {path} must be a JSON object')
    names = set(data.get('land_attribute_fields_order') or []) | set(data.get('gis_field_mapping') or {})
    names |= set((data.get('observed_on_mawan') or {}).get('entity_apps') or [])
    numeric = set(data.get('numeric_like') or [])
    if not names and not numeric:
        raise ValueError(f'--schema {path} has no land_attribute_fields_order, gis_field_mapping or numeric_like')
    return names, numeric


def hatch_point(hatch):
    """A point inside a HATCH's largest boundary path, or None."""
    import shapely
    from ezdxf import path as dxfpath
    best = None
    for item in dxfpath.from_hatch(hatch):
        points = [(p.x, p.y) for p in item.flattening(distance=0.01)]
        if len(points) >= 4:
            shape = shapely.Polygon(points)
            if shape.is_valid and (best is None or shape.area > best.area):
                best = shape
    return shapely.point_on_surface(best) if best is not None and not best.is_empty else None


# OLE ProgID -> extension of the Office file held in its Package stream (Office 2007 and later).
OLE_PACKAGES = {'Excel.Sheet.12': '.xlsx', 'Excel.SheetMacroEnabled.12': '.xlsm', 'Word.Document.12': '.docx',
                'PowerPoint.Show.12': '.pptx'}
CFB_MAGIC = bytes.fromhex('D0CF11E0A1B11AE1')
OLE_FRAME_FIELDS = ('handle', 'layer', 'owner', 'prog_id', 'user_type', 'file', 'bytes', 'sha256', 'shown_sheet', 'shown_range',
                    'sheets', 'corners_from')
TABLE_CELL_FIELDS = ('row', 'col', 'text', 'x', 'y', 'layer', 'handle')


def comp_obj(data, encoding):
    """(user type, ProgID) from an OLE CompObj stream: a 28-byte header, the user type, a clipboard format
    (a length-prefixed name, or a marker and a format id) and the ProgID."""
    import struct

    def string(at):
        size = struct.unpack_from('<I', data, at)[0]
        return data[at+4:at+4+size].rstrip(b'\0'), at+4+size
    try:
        user, at = string(28)
        marker = struct.unpack_from('<I', data, at)[0]
        at += 8 if marker in (0xFFFFFFFF, 0xFFFFFFFE) else 4 + marker
        prog, _ = string(at)
    except struct.error:
        return None, None
    return user.decode(encoding, 'replace') or None, prog.decode('ascii', 'replace') or None


def ole_corners(entity, data):
    """Frame corners as drawn and where they came from. OLE2FRAME binary data starts with an 0x80 0x55 marker and the
    four corners as 3D doubles; LibreDWG writes placeholders in group codes 10/11, so the header is preferred."""
    import struct
    if data[:2] == b'\x80\x55' and len(data) >= 98:
        values = struct.unpack_from('<12d', data, 2)
        corners = [values[i:i+2] for i in range(0, 12, 3)]
        if all(math.isfinite(v) for corner in corners for v in corner):
            return corners, 'binary header'
    tags = entity.acdb_ole2frame
    v10, v11 = (tags.get_first_value(code, None) for code in (10, 11)) if tags is not None else (None, None)
    if v10 is not None and v11 is not None:
        return [(v10[0], v10[1]), (v11[0], v10[1]), (v11[0], v11[1]), (v10[0], v11[1])], 'group codes 10/11'
    return None, None


def ole_payload(data, encoding):
    """(ProgID, user type, extension, file bytes) of the compound file inside OLE2FRAME data; needs olefile.
    Office 2007+ objects keep their file in a Package stream; others are saved as the compound file itself."""
    import io
    import olefile
    start = data.find(CFB_MAGIC)
    if start < 0:
        return None, None, None, None
    ole = olefile.OleFileIO(io.BytesIO(data[start:]))
    try:
        streams = {'/'.join(path).lstrip('\x01\x02\x03'): path for path in ole.listdir()}
        user, prog = comp_obj(ole.openstream(streams['CompObj']).read(), encoding) if 'CompObj' in streams else (None, None)
        if 'Package' in streams:
            payload = ole.openstream(streams['Package']).read()
            return prog, user, OLE_PACKAGES.get(prog, '.zip' if payload[:2] == b'PK' else '.bin'), payload
    finally:
        ole.close()
    extension = '.xls' if 'Workbook' in streams or 'Book' in streams else '.doc' if 'WordDocument' in streams else '.cfb'
    return prog, user, extension, data[start:]


def workbook_sheets(payload):
    """(active sheet, [{name, range, digest of its values}]) of an xlsx / xlsm; needs openpyxl."""
    import hashlib
    import io
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    sheets = [{'name': sheet.title, 'range': sheet.dimensions,
               'digest': hashlib.sha256(json.dumps(list(sheet.iter_rows(values_only=True)), default=str,
                                                   ensure_ascii=False).encode('utf-8')).hexdigest()[:16]}
              for sheet in book.worksheets]
    return book.active.title if book.active is not None else None, sheets


def count_entities(path, name):
    """How many entities of a DXF type the file holds, counted on its group 0 tags (ezdxf drops empty ones)."""
    count, wanted = 0, name.encode('ascii')
    with open(path, 'rb') as reader:
        while True:
            code_line, value_line = reader.readline(), reader.readline()
            if not code_line:
                break
            count += code_line.strip() == b'0' and value_line.strip() == wanted
    return count


def rank(values, tolerance):
    """1-based rank of each value among groups of sorted values no more than tolerance apart."""
    lookup, level, last = {}, 0, None
    for value in sorted(set(values)):
        if last is None or value - last > tolerance:
            level += 1
        lookup[value], last = level, value
    return [lookup[value] for value in values]


def embedded_objects(doc, source, scale, warnings):
    """OLE objects and table blocks of a drawing: frame records (geometry for model space only), the embedded files,
    sheets that differ between workbooks, and table cells as placed in each *T block. Nothing is interpreted."""
    import hashlib
    import re
    from collections import Counter
    import shapely
    owners = {block.block_record_handle: block.name for block in doc.blocks}
    frames, files = [], {}
    oles = [e for e in doc.entitydb.values() if e.dxftype() == 'OLE2FRAME']
    try:
        import olefile  # noqa: F401
    except ImportError:
        olefile = None
        if oles:
            warnings.append(f'{len(oles)} OLE objects are listed without their contents: olefile is not installed')
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        openpyxl = None
    sheet_copies = {}  # sheet name -> {digest: [handles]}
    for entity in oles:
        data = entity.binary_data()
        handle, owner = entity.dxf.handle, owners.get(entity.dxf.owner, entity.dxf.owner)
        corners, corners_from = ole_corners(entity, data)
        shape = None
        if corners and owner == '*Model_Space':
            shape = shapely.Polygon([(x*scale, y*scale) for x, y in corners])
            shape = shape if shape.is_valid and shape.area > 0 else None
        record = {'handle': handle, 'layer': entity.dxf.layer, 'owner': owner, 'corners_from': corners_from, 'geometry': shape}
        payload = None
        if olefile is not None:
            try:
                prog, user, extension, payload = ole_payload(data, doc.encoding or 'cp1252')
            except OSError as error:  # olefile raises OSError subclasses on malformed compound files
                warnings.append(f'OLE object {handle} could not be read: {error}')
            if payload:
                name = f'ole/{handle}{extension}'
                files[name] = payload
                record.update(prog_id=prog, user_type=user, file=name, bytes=len(payload),
                              sha256=hashlib.sha256(payload).hexdigest())
                if extension in ('.xlsx', '.xlsm'):
                    if openpyxl is None:
                        warnings.append(f'Sheets of {name} are not listed: openpyxl is not installed')
                    else:
                        active, sheets = workbook_sheets(payload)
                        shown = next((s for s in sheets if s['name'] == active), None)
                        record.update(shown_sheet=active, shown_range=shown['range'] if shown else None,
                                      sheets=', '.join(f'{s["name"]} {s["range"]}' for s in sheets))
                        for sheet in sheets:
                            sheet_copies.setdefault(sheet['name'], {}).setdefault(sheet['digest'], []).append(
                                (handle, sheet['name'] == active))
        frames.append(record)
    versions = [{'sheet': name, 'versions': [{'frames': [h for h, _ in copies], 'shown_in': [h for h, shown in copies if shown]}
                                             for copies in by_digest.values()]}
                for name, by_digest in sorted(sheet_copies.items()) if len(by_digest) > 1]
    if versions:
        warnings.append(f'Embedded workbooks hold different versions of {len(versions)} sheets '
                        f'({", ".join(v["sheet"] for v in versions)}); see embedded.sheet_versions')

    # Tables: AutoCAD draws an ACAD_TABLE as an anonymous *T block of texts and lines, which is where cells are read.
    referenced = {e.get_block_name(): e for e in doc.entitydb.values()
                  if e.dxftype() == 'ACAD_TABLE' and hasattr(e, 'get_block_name')}
    inserted = Counter(e.dxf.name for e in doc.entitydb.values() if e.dxftype() == 'INSERT')
    empty = count_entities(source, 'ACAD_TABLE') - sum(1 for e in doc.entitydb.values() if e.dxftype() == 'ACAD_TABLE')
    if empty > 0:
        warnings.append(f'{empty} ACAD_TABLE entities in the file carry no data (LibreDWG does not convert their contents) and '
                        'were not read; their cells are read from the *T blocks as drawn, but which table owns which block is '
                        'lost: convert with ODA File Converter to keep it')
    tables = []
    for block in doc.blocks:
        if not re.fullmatch(r'\*T\d+', block.name):
            continue
        cells = []
        for entity in block:
            if entity.dxftype() not in ('TEXT', 'MTEXT'):
                continue
            mtext = entity.dxftype() == 'MTEXT'
            text = clean_text(entity.plain_text(split=False) if mtext else entity.plain_text())
            if text:
                insert = entity.dxf.insert
                cells.append({'text': text, 'x': round(insert.x, 3), 'y': round(insert.y, 3), 'layer': entity.dxf.layer,
                              'handle': entity.dxf.handle,
                              'height': entity.dxf.get('char_height' if mtext else 'height', 0) or 0})
        if not cells:
            continue
        tolerance = 0.5*sorted(c['height'] for c in cells)[len(cells)//2]
        for cell, row, col in zip(cells, rank([-c['y'] for c in cells], tolerance), rank([c['x'] for c in cells], tolerance)):
            cell.update(row=row, col=col)
        cells.sort(key=lambda c: (c['row'], c['col']))
        table = referenced.get(block.name)
        name = f'tables/{block.name.lstrip("*")}.csv'
        files[name] = ('﻿' + ','.join(TABLE_CELL_FIELDS) + '\n' + ''.join(
            ','.join('"' + str(c[f]).replace('"', '""') + '"' if f == 'text' else str(c[f]) for f in TABLE_CELL_FIELDS) + '\n'
            for c in cells)).encode('utf-8')
        tables.append({'block': block.name, 'file': name, 'cells': len(cells), 'rows': max(c['row'] for c in cells),
                       'cols': max(c['col'] for c in cells), 'first_row': ' | '.join(c['text'] for c in cells if c['row'] == 1),
                       'table_entity': table.dxf.handle if table is not None else None,
                       'table_owner': owners.get(table.dxf.owner, table.dxf.owner) if table is not None else None,
                       'block_inserts': inserted.get(block.name, 0)})
    return frames, files, versions, tables


def plot_xiangyuan(parcels, boundary, path, title, font, ole_frames=None):
    import re
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    figure, axis = plt.subplots(figsize=(12, 9))
    colours = matplotlib.colormaps['tab10']
    land = parcels[parcels['control_type'] == '地块']
    groups = land['landuse'].fillna('').map(lambda code: (re.match(r'[A-Z]+', code) or [''])[0])
    handles = []
    for i, group in enumerate(sorted(set(groups))):
        land[groups == group].plot(ax=axis, facecolor=colours(i % 10), edgecolor='#4a4a4a', linewidth=0.4, alpha=0.6)
        handles.append(plt.Rectangle((0, 0), 1, 1, facecolor=colours(i % 10), alpha=0.6, label=group or '(no code)'))
    units = parcels[parcels['control_type'] == '规划控制单元']
    if len(units):
        units.boundary.plot(ax=axis, color='#555555', linewidth=0.6, linestyle='--')
        handles.append(plt.Line2D([], [], color='#555555', linestyle='--', label='规划控制单元'))
    if boundary is not None:
        axis.plot(*boundary.exterior.xy, color='#000000', linewidth=1.2)
        handles.append(plt.Line2D([], [], color='#000000', label='规划边界'))
    held = parcels[parcels['status'] == 'hold']
    if len(held):
        held.boundary.plot(ax=axis, color='#e67e22', linewidth=1.6)
        handles.append(plt.Line2D([], [], color='#e67e22', linewidth=2, label='hold'))
    shown = ole_frames[ole_frames.geometry.notna()] if ole_frames is not None and len(ole_frames) else []
    if len(shown):
        shown.boundary.plot(ax=axis, color='#1f77b4', linewidth=0.8, linestyle=':')
        for _, row in shown.iterrows():
            corner = row.geometry.bounds
            axis.annotate(f'OLE {row["handle"]} {row["shown_sheet"] or ""}', (corner[0], corner[3]), fontsize=5,
                          color='#1f77b4', va='top')
        handles.append(plt.Line2D([], [], color='#1f77b4', linestyle=':', label='OLE object'))
    for _, row in land[land['parcel_id'].fillna('') != ''].iterrows():
        point = row.geometry.representative_point()
        axis.annotate(row['parcel_id'], (point.x, point.y), fontsize=5, ha='center', va='center')
    axis.legend(handles=handles, loc='center left', bbox_to_anchor=(1.01, 0.5), fontsize=8, frameon=False)
    axis.set_aspect('equal', adjustable='datalim')
    axis.ticklabel_format(useOffset=False, style='plain')
    axis.tick_params(labelsize=7)
    axis.set_title(title, fontsize=10)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def xiangyuan(args):
    """Parcels of a 湘源控规 drawing straight from its closed polylines and their XDATA attributes."""
    import warnings as python_warnings
    from collections import Counter
    import geopandas as gpd
    import pandas as pd
    import pyogrio
    import shapely
    from ezdxf import path as dxfpath
    source = Path(args.source)
    if source.suffix.lower() != '.dxf' or not source.is_file():
        raise ValueError('Source must be an existing .dxf file (convert DWG first with cad.py convert)')
    header = read_header(source)
    doc, audit = read_drawing(source, args.encoding)
    if doc is None:
        raise ValueError('cad.py xiangyuan needs ezdxf: the parcel attributes are XDATA, which GDAL does not read')
    names = [layer.dxf.name for layer in doc.layers]
    encoding, encoding_report = choose_encoding(source, names, header, args.encoding)
    if encoding_report['source'] == 'auto':
        doc, audit = read_drawing(source, encoding)
    if args.units == 'header':
        scale = UNITS.get(header.get('insunits'), (None, None))[1]
        if not scale:
            raise ValueError(f'Header $INSUNITS is {header.get("units") or "missing"}; pass --units m or --units mm')
        units_basis = f'header $INSUNITS = {header.get("units")}'
    else:
        scale = {'m': 1.0, 'mm': 0.001}[args.units]
        units_basis = f'--units {args.units} (header $INSUNITS = {header.get("units")})'
    known, numeric = set(XIANGYUAN_FIELDS), set(XIANGYUAN_NUMERIC)
    schema = None
    if args.schema:
        extra_names, extra_numeric = load_schema(args.schema)
        known |= extra_names | extra_numeric
        numeric |= extra_numeric
        schema = {'path': str(Path(args.schema).absolute()), 'sha256': digest(args.schema),
                  'fields': len(extra_names), 'numeric': len(extra_numeric)}
    warnings = []
    model = doc.modelspace()
    # An unexploded 湘源 parcel is a custom LZX_LAND entity whose DXF form keeps only its class marker.
    registered = any(str(key[0]).upper() == 'LZX_LAND' for key in getattr(doc.classes, 'classes', {}))
    unexploded = sum(1 for e in model if e.dxftype() == 'LZX_LAND')
    proxies = sum(1 for e in model if e.dxftype() == 'ACAD_PROXY_ENTITY')
    if unexploded:
        raise ValueError(f'{unexploded} 湘源 parcels are unexploded LZX_LAND custom entities, whose attributes are not in DXF; '
                         'explode or export them in 湘源控规 and convert the drawing again')
    if proxies:
        warnings.append(f'{proxies} proxy entities from a plug-in that was not loaded; their contents are not read')

    def polygon(entity):
        points = [(p.x*scale, p.y*scale) for p in dxfpath.make_path(entity).flattening(distance=args.arc_tolerance/scale)]
        return shapely.Polygon(points) if len({pt for pt in points}) >= 3 else None

    records, skipped, open_marked, multi = [], Counter(), 0, Counter()
    for entity in model:
        several = Counter()
        values = xdata_values(entity, several)
        if XIANGYUAN_KEY not in values:
            continue
        multi.update(several)  # only parcels count: AutoCAD's own ACAD application often repeats 1000 groups
        kind = entity.dxftype()
        if kind not in ('LWPOLYLINE', 'POLYLINE'):
            skipped[kind] += 1
            continue
        closed = entity.closed if kind == 'LWPOLYLINE' else entity.is_closed
        shape = polygon(entity) if closed else None
        if shape is None:
            open_marked += 1
            continue
        layer = entity.dxf.layer
        code = values.get(XIANGYUAN_KEY, '')
        record = {field: values.get(source_field) for field, source_field in XIANGYUAN_ALIASES.items()}
        other = {app: value for app, value in values.items() if app not in known}
        record.update({app: value for app, value in values.items() if app in known})
        record['other_xdata'] = json.dumps(other, ensure_ascii=False) if other else None
        reasons = []
        if layer not in (code, f'{code}_K'):  # 湘源 puts parcel outlines on <code>_K; some drawings drop the suffix
            reasons.append(f'layer {layer} does not match 用地代码 {code}')
        if (values.get('删除标记') or '').strip():
            reasons.append('marked deleted')
        if not (record['parcel_id'] or '').strip():
            reasons.append('no 地块编号')
        valid = bool(shapely.is_valid(shape))
        if not valid:
            reasons.append('invalid geometry')
        area = float(shape.area)
        try:
            declared = float(values.get('用地面积') or 'nan')
        except ValueError:
            declared = float('nan')
        diff = (area - declared)/declared if declared == declared and declared else None
        if diff is not None and args.max_area_diff is not None and abs(diff) > args.max_area_diff:
            reasons.append(f'area differs from 用地面积 by {diff:+.2%}')
        record.update(control_type='规划控制单元' if code == 'DY' else '地块', area_m2=round(area, 3),
                      area_declared=declared if declared == declared else None,
                      area_diff=round(diff, 6) if diff is not None else None, valid=valid, layer=layer,
                      handle=entity.dxf.handle, reasons=reasons, geometry=shape)
        records.append(record)
    if not records:
        raise ValueError(f'No closed polylines carry a {XIANGYUAN_KEY} XDATA attribute; this does not look like a 湘源控规 drawing')
    if args.units == 'header' and scale != 1.0:
        # 湘源 records each parcel's area: if it fits the drawing as metres but not at the header unit, the header is wrong.
        pairs = [(r['area_m2'], r['area_declared']) for r in records if r['area_declared']]
        if pairs:
            import statistics
            at_header = statistics.median(abs(a - d)/d for a, d in pairs)
            as_metres = statistics.median(abs(a/scale**2 - d)/d for a, d in pairs)
            if as_metres < 0.01 <= at_header:
                raise ValueError(f'Header says {header.get("units")}, but 用地面积 matches the drawing in metres '
                                 f'(median difference {as_metres:.4%} vs {at_header:.2%}); pass --units m')
    counts = Counter((r['control_type'], (r['parcel_id'] or '').strip()) for r in records if (r['parcel_id'] or '').strip())
    for record in records:
        if counts[(record['control_type'], (record['parcel_id'] or '').strip())] > 1:
            record['reasons'].append('duplicate 地块编号')

    # Number blocks: their attributes repeat the parcel's; check each sits in the parcel it names.
    tree = shapely.STRtree([r['geometry'] for r in records])
    checks, numbered = [], set()
    for entity in model.query('INSERT'):
        attributes = {a.dxf.tag: a.dxf.text for a in entity.attribs}
        if '地块编号' not in attributes or entity.dxf.layer == args.facility_layer:
            continue
        insert = entity.dxf.insert
        point = shapely.Point(insert.x*scale, insert.y*scale)
        hits = [int(i) for i in tree.query(point, predicate='intersects')]
        same = [i for i in hits if (records[i]['parcel_id'] or '') == attributes['地块编号'] and records[i]['control_type'] == '地块']
        same = same or [i for i in hits if (records[i]['parcel_id'] or '') == attributes['地块编号']]
        if same:
            numbered.update(same)
            if len(same) == 1:  # the outline is authoritative: its block only fills fields the outline leaves empty
                record = records[same[0]]
                for tag, value in attributes.items():
                    if (tag in known and tag not in XIANGYUAN_ALIASES.values() and (value or '').strip()
                            and not (record.get(tag) or '').strip()):
                        record[tag] = value
                        record['filled_from_block'] = ','.join(filter(None, [record.get('filled_from_block'), tag]))
            continue
        checks.append({'kind': 'number-block-mismatch' if hits else 'number-block-outside', 'parcel_id': attributes['地块编号'],
                       'found': ','.join(sorted({records[i]['parcel_id'] or '' for i in hits})) or None,
                       'layer': entity.dxf.layer, 'handle': entity.dxf.handle, 'geometry': point})
    # Units: closed outlines on the unit layer, named by the unit labels (TEXT / MTEXT) inside them.
    units = []
    if args.unit_layer:
        labels = []
        for entity in model:
            if entity.dxf.layer == args.unit_label_layer and entity.dxftype() in ('TEXT', 'MTEXT'):
                text = clean_text(entity.plain_text(split=False) if entity.dxftype() == 'MTEXT' else entity.plain_text())
                insert = entity.dxf.insert
                labels.append((text, shapely.Point(insert.x*scale, insert.y*scale)))
        kept = []  # (shape, handle) of outlines seen so far
        for entity in model:
            if entity.dxf.layer != args.unit_layer or entity.dxftype() not in ('LWPOLYLINE', 'POLYLINE'):
                continue
            if not (entity.closed if entity.dxftype() == 'LWPOLYLINE' else entity.is_closed):
                continue
            shape = polygon(entity)
            if shape is None:
                continue
            inside = sorted({text for text, point in labels if shape.covers(point)})
            # The same outline drawn twice, whatever its start vertex or direction.
            twin = next((handle for other, handle in kept
                         if shapely.area(shapely.symmetric_difference(shape, other)) <= 1e-6*shape.area), None)
            reasons = (['no unit label inside'] if not inside else ['several unit labels inside'] if len(inside) > 1 else [])
            if twin:
                reasons.append('duplicate outline')
            units.append({'unit_id': inside[0] if len(inside) == 1 else None, 'labels': ','.join(inside) or None,
                          'reason': '; '.join(reasons) or None, 'status': 'hold' if reasons else 'pass',
                          'area_m2': round(float(shape.area), 3), 'duplicate_of': twin, 'layer': entity.dxf.layer,
                          'handle': entity.dxf.handle, 'geometry': shape})
            if not twin:
                kept.append((shape, entity.dxf.handle))
    # Fills: 湘源 puts one HATCH per parcel on <code>_T. They carry no attributes; pair them with the outlines.
    fills, stray_fills, fill_mismatch = Counter(), [], []
    for hatch in model.query('HATCH'):
        layer = hatch.dxf.layer
        if not layer.endswith('_T'):
            continue
        try:
            point = hatch_point(hatch)
        except (TypeError, ValueError):
            point = None
        if point is None:
            continue
        point = shapely.Point(point.x*scale, point.y*scale)
        hits = [int(i) for i in tree.query(point, predicate='intersects')]
        hits = [i for i in hits if records[i]['control_type'] == '地块'] or hits
        if not hits:
            stray_fills.append(hatch.dxf.handle)
            continue
        for i in hits:
            fills[i] += 1
            if layer[:-2] != (records[i]['landuse'] or ''):
                fill_mismatch.append({'parcel_id': records[i]['parcel_id'], 'fill_layer': layer, 'landuse': records[i]['landuse']})
    named = [u for u in units if u['unit_id'] and not u['duplicate_of']]
    for i, record in enumerate(records):
        record['number_block'] = i in numbered
        record['fills'] = fills[i]
        if named and record['control_type'] == '地块':
            point = shapely.point_on_surface(record['geometry'])
            ids = sorted({u['unit_id'] for u in named if u['geometry'].covers(point)})
            record['unit_id'] = ids[0] if len(ids) == 1 else None
            if len(ids) > 1:
                record['reasons'].append(f'in several units {",".join(ids)}')
        record['reason'] = '; '.join(record.pop('reasons')) or None
        record['status'] = 'hold' if record['reason'] else 'pass'

    facilities = []
    if args.facility_layer:
        for entity in model.query('INSERT'):
            if entity.dxf.layer != args.facility_layer:
                continue
            values = {**{a.dxf.tag: a.dxf.text for a in entity.attribs}, **xdata_values(entity, multi)}
            insert = entity.dxf.insert
            point = shapely.Point(insert.x*scale, insert.y*scale)
            number = (values.get('地块编号') or '').strip()
            named_parcels = [r for r in records if (r['parcel_id'] or '').strip() == number] if number else []
            match = ('no id' if not number else 'unknown id' if not named_parcels
                     else 'inside' if any(r['geometry'].covers(point) for r in named_parcels) else 'elsewhere')
            distance = (round(min(float(r['geometry'].distance(point)) for r in named_parcels), 1)
                        if match == 'elsewhere' else None)
            facilities.append({'block': entity.dxf.name, 'layer': entity.dxf.layer, 'handle': entity.dxf.handle, **values,
                               'parcel_match': match, 'parcel_distance_m': distance, 'geometry': point})
    boundary = None
    if args.boundary_layer:
        rings = [polygon(e) for e in model if e.dxf.layer == args.boundary_layer and e.dxftype() in ('LWPOLYLINE', 'POLYLINE')
                 and (e.closed if e.dxftype() == 'LWPOLYLINE' else e.is_closed)]
        rings = [r for r in rings if r is not None]
        boundary = shapely.union_all(rings) if rings else None

    # Parcel attributes copied into block definitions (road red lines, inset maps): reported, never taken as parcels.
    inserted = Counter(e.dxf.name for e in model.query('INSERT'))
    copies = []
    for block in doc.blocks:
        if block.name.startswith('*'):
            continue
        carrying = sum(1 for e in block if XIANGYUAN_KEY in xdata_values(e))
        if carrying:
            copies.append({'block': block.name, 'entities': carrying, 'inserts_in_model_space': inserted.get(block.name, 0)})
    if copies:
        warnings.append(f'{sum(c["entities"] for c in copies)} entities inside {len(copies)} block definitions carry parcel '
                        'attributes (copies on road lines and insets); they are not read as parcels')
    if multi:
        warnings.append(f'{sum(multi.values())} XDATA applications held several values; the last was kept, as 湘源 reads them')

    extra = sorted({k for r in records for k in r} - set(XIANGYUAN_FIXED) - {'geometry'})
    parcels = gpd.GeoDataFrame(pd.DataFrame(records, columns=[*XIANGYUAN_FIXED, *extra, 'geometry']), geometry='geometry')
    converted = []
    for field in sorted(numeric & set(extra)):
        text = parcels[field].astype(object)
        present = text.notna() & (text.astype(str).str.strip() != '')
        numbers = pd.to_numeric(text.where(present), errors='coerce')
        bad = text[present & numbers.isna()]
        if len(bad):
            samples = ', '.join(repr(v) for v in list(dict.fromkeys(bad))[:3])
            warnings.append(f'{field} kept as text: {len(bad)} values are not numbers ({samples})')
            continue
        parcels[field] = numbers.astype(float)
        converted.append(field)
    facility_fields = sorted({k for f in facilities for k in f} - {'geometry'})
    frames = {'parcels': (parcels, 'Polygon'),
              'facilities': (gpd.GeoDataFrame(pd.DataFrame(facilities, columns=[*facility_fields, 'geometry']), geometry='geometry'), 'Point'),
              'qa_number_blocks': (gpd.GeoDataFrame(pd.DataFrame(checks, columns=['kind', 'parcel_id', 'found', 'layer', 'handle', 'geometry']),
                                                    geometry='geometry'), 'Point')}
    if args.unit_layer:
        frames['units'] = (gpd.GeoDataFrame(pd.DataFrame(units, columns=[*UNIT_FIELDS, 'geometry']), geometry='geometry'), 'Polygon')
    if boundary is not None:
        frames['boundary'] = (gpd.GeoDataFrame({'layer': [args.boundary_layer]}, geometry=[boundary]), 'MultiPolygon'
                              if boundary.geom_type == 'MultiPolygon' else 'Polygon')
    ole_frames, embedded_files, sheet_versions, tables = embedded_objects(doc, source, scale, warnings)
    frames['ole_frames'] = (gpd.GeoDataFrame(pd.DataFrame(ole_frames, columns=[*OLE_FRAME_FIELDS, 'geometry']), geometry='geometry'),
                            'Polygon')
    land = parcels[parcels['control_type'] == '地块']
    union = shapely.union_all(land.geometry.to_numpy()) if len(land) else None
    area = {'parcels_m2': round(float(land['area_m2'].sum()), 3), 'parcels_union_m2': round(float(union.area), 3) if union else 0.0,
            'units_m2': round(float(parcels.loc[parcels['control_type'] == '规划控制单元', 'area_m2'].sum()), 3)}
    if named:
        totals = {}
        for unit in named:
            totals[unit['unit_id']] = totals.get(unit['unit_id'], 0.0) + unit['area_m2']
        area['units_by_id_m2'] = {k: round(v, 3) for k, v in sorted(totals.items())}
    if boundary is not None:
        area.update(boundary_m2=round(float(boundary.area), 3),
                    parcels_inside_boundary_m2=round(float(union.intersection(boundary).area), 3) if union else 0.0)
    diffs = parcels['area_diff'].dropna().abs()
    summary = {'parcels': int(len(land)), 'control_units': int(len(parcels) - len(land)),
               'pass': int((parcels['status'] == 'pass').sum()), 'hold': int((parcels['status'] == 'hold').sum()),
               'reasons': dict(Counter(r for value in parcels['reason'].dropna() for r in value.split('; ')).most_common()),
               'number_blocks_checked': len(numbered) + len(checks), 'number_block_problems': len(checks),
               'area_diff_median': round(float(diffs.median()), 8) if len(diffs) else None,
               'area_diff_max': round(float(diffs.max()), 6) if len(diffs) else None,
               'skipped_non_polylines': dict(skipped), 'open_polylines_with_code': open_marked,
               'fills': {'parcels_without_fill': int(sum(1 for i, r in enumerate(records) if r['control_type'] == '地块' and not fills[i])),
                         'fills_without_parcel': len(stray_fills), 'fill_code_mismatches': fill_mismatch[:20]},
               'facility_matches': dict(Counter(f['parcel_match'] for f in facilities)),
               'filled_from_block': dict(Counter(t for value in parcels['filled_from_block'].dropna() for t in value.split(','))),
               'other_xdata_parcels': int(parcels['other_xdata'].notna().sum()), 'numeric_fields': converted}
    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    with bundle(args.output) as stage:
        target = stage / 'parcels.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message="'crs' was not provided")
            for name, (frame, kind) in frames.items():
                pyogrio.write_dataframe(frame, target, layer=name, driver='GPKG', geometry_type=kind)
        layer_counts = {name: len(frame) for name, (frame, _) in frames.items()}
        back = {name: int(pyogrio.read_info(target, layer=name, force_feature_count=True)['features']) for name in frames}
        if back != layer_counts:
            raise ValueError(f'Readback feature counts {back} differ from {layer_counts}; nothing was published')
        for name, content in embedded_files.items():
            (stage / name).parent.mkdir(exist_ok=True)
            (stage / name).write_bytes(content)
        plot_xiangyuan(parcels, boundary if boundary is not None and boundary.geom_type == 'Polygon' else None, stage / 'check.png',
                       f'{source.name} · {summary["parcels"]} parcels, {summary["control_units"]} control units · {summary["hold"]} hold', font,
                       frames['ole_frames'][0])
        manifest = {'manifest': 'cad-xiangyuan', 'schema_version': 1,
                    'source': {'name': source.name, 'bytes': source.stat().st_size, 'sha256': digest(source)},
                    'header': header, 'encoding': encoding_report, 'audit': audit, 'schema': schema,
                    'custom_entities': {'lzx_land_class_registered': registered, 'lzx_land_instances': unexploded,
                                        'proxy_entities': proxies},
                    'block_copies': copies, 'warnings': warnings,
                    'embedded': {'ole_frames': [{k: v for k, v in f.items() if k != 'geometry'} for f in ole_frames],
                                 'sheet_versions': sheet_versions, 'tables': tables},
                    'parameters': {'units': args.units, 'units_basis': units_basis, 'scale_to_metre': scale,
                                   'arc_tolerance_m': args.arc_tolerance,
                                   'max_area_diff': args.max_area_diff, 'boundary_layer': args.boundary_layer,
                                   'unit_layer': args.unit_layer, 'unit_label_layer': args.unit_label_layer,
                                   'facility_layer': args.facility_layer},
                    'fields': extra, 'layers': layer_counts, 'summary': summary, 'area': area, 'crs': None,
                    'output': {'path': target.name, 'sha256': digest(target)}, 'plot': 'check.png'}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad xiangyuan', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'source': {'path': str(source.absolute()), 'sha256': manifest['source']['sha256']},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256'],
                                                      'check.png': digest(stage / 'check.png'),
                                                      **{name: digest(stage / name) for name in embedded_files}},
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['A parcel is a closed polyline with a 用地代码 XDATA application, as 湘源控规 writes them',
                                                            'Attribute values are the XDATA strings as stored; nothing is recalculated except area_m2',
                                                            'A number block only fills fields its outline leaves empty (filled_from_block)',
                                                            'An embedded workbook shows its active sheet as last saved (shown_sheet); '
                                                            'embedded files are saved as stored, table cells as placed in their *T block',
                                                            'No CRS is set; coordinates are only scaled to metres']})
    return manifest


def convert(args):
    """Convert a DWG to DXF with an installed converter and read the result back."""
    import shutil
    import subprocess
    import tempfile
    import time
    source = Path(args.source)
    if source.suffix.lower() != '.dwg' or not source.is_file():
        raise ValueError('Source must be an existing .dwg file')
    version = dwg_version(source)
    found = find_converters()
    if args.tool == 'auto':
        tool = 'oda' if 'oda' in found else 'libredwg' if 'libredwg' in found else None
    else:
        tool = args.tool if args.tool in found else None
    if tool is None:
        wanted = 'ODA File Converter or LibreDWG (dwg2dxf)' if args.tool == 'auto' else args.tool
        raise ValueError(f'No {wanted} found. Install one (ODA File Converter needs its licence accepted; '
                         'LibreDWG: brew install libredwg), or save the drawing as DXF 2013/2018 in AutoCAD / ZWCAD')
    as_libredwg, as_oda = DXF_VERSIONS[args.version]
    stat = source.stat()
    with tempfile.TemporaryDirectory() as folder:
        work = Path(folder)
        given = work / 'in'
        given.mkdir()
        copy = given / source.name  # converters see a copy; the original is never opened for writing
        shutil.copy2(source, copy)
        out = work / 'out'
        out.mkdir()
        target = out / f'{source.stem}.dxf'
        if tool == 'libredwg':
            command = [found[tool], '-y', '--as', as_libredwg, '-o', str(target), str(copy)]
        else:
            command = [found[tool], str(given), str(out), as_oda, 'DXF', '0', '1', source.name]
        started = time.monotonic()
        run = subprocess.run(command, capture_output=True, text=True, errors='replace', timeout=args.timeout)
        seconds = round(time.monotonic() - started, 2)
        messages = [line for line in (run.stdout + run.stderr).splitlines() if line.strip()]
        if run.returncode != 0 or not target.is_file() or target.stat().st_size == 0:
            raise ValueError(f'{tool} failed (exit {run.returncode}): ' + ' | '.join(messages[-5:]))
        header = read_header(target)
        doc, audit = read_drawing(target, None)
        readback = {'header': header}
        if doc is not None:
            model = doc.modelspace()
            layers = {entity.dxf.layer for entity in model}
            readback.update(model_space_entities=len(model), layers_with_entities=len(layers), layer_table=len(doc.layers),
                            paper_space_entities=sum(len(l) for l in doc.layouts if l.name != 'Model'),
                            extent=header.get('extent'), audit=audit)
        else:
            readback['note'] = 'ezdxf not installed: only the header was read back'
        with bundle(args.output) as stage:
            shutil.copy2(target, stage / target.name)
            published = stage / target.name
            warnings = sorted(set(m for m in messages if 'warning' in m.lower()))
            manifest = {'manifest': 'cad-convert', 'schema_version': 1,
                        'source': {'name': source.name, 'bytes': stat.st_size, 'sha256': digest(source), 'dwg_version': version,
                                   'release': VERSIONS.get(version)},
                        'converter': {'tool': tool, 'path': found[tool], 'command': [Path(command[0]).name, *command[1:]],
                                      'exit_code': run.returncode, 'seconds': seconds, 'messages': messages[-20:],
                                      'available': sorted(found)},
                        'output': {'path': published.name, 'bytes': published.stat().st_size, 'sha256': digest(published),
                                   'requested_version': args.version},
                        'readback': readback, 'warnings': warnings[:20]}
            write_json(stage / 'manifest.json', manifest)
            write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad convert', 'version': VERSION,
                                                'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                                'source': {'path': str(source.absolute()), 'sha256': manifest['source']['sha256']},
                                                'files': {'manifest.json': digest(stage / 'manifest.json'), published.name: manifest['output']['sha256']},
                                                'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                                'assumptions': ['The converter reads a copy of the DWG; the original is not modified',
                                                                'Readback counts come from the DXF; the DWG itself is not parsed here']})
    return manifest


def load_inspection(path):
    path = Path(path)
    path = path / 'cad.json' if path.is_dir() else path
    report = json.loads(path.read_text(encoding='utf-8'))
    if report.get('report') != 'cad-inventory' or report.get('schema_version') != 1:
        raise ValueError(f'--inspection is not a cad.py inspect report (cad.json v1): {path}')
    return path, report


PREFIX_EASTING = 1_000_000  # a false easting this large carries the zone number, as in 39500000


def false_easting(crs):
    operation = crs.coordinate_operation
    return next((float(p.value) for p in operation.params if p.name == 'False easting'), None) if operation else None


def check_points(path, limit):
    """Distances between each point's drawing coordinates (src_x, src_y) and its known coordinates in the assigned
    CRS (dst_x, dst_y), with nothing transformed: the vector-georef control-point CSV, every row used as a check."""
    import csv
    with open(path, encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream))
    missing = {'id', 'src_x', 'src_y', 'dst_x', 'dst_y'} - set(rows[0] if rows else [])
    if not rows or missing:
        raise ValueError(f'--check-points {path} needs rows with id, src_x, src_y, dst_x, dst_y ({", ".join(sorted(missing)) or "no rows"})')
    try:
        errors = [{'id': r['id'], 'error_m': round(math.hypot(float(r['dst_x'])-float(r['src_x']), float(r['dst_y'])-float(r['src_y'])), 4)}
                  for r in rows]
    except ValueError as error:
        raise ValueError(f'--check-points {path}: coordinates must be numbers ({error})') from error
    worst = max(e['error_m'] for e in errors)
    return {'path': str(Path(path).absolute()), 'sha256': digest(path), 'points': len(errors), 'max_check_error_m': limit,
            'max_error_m': worst, 'rmse_m': round(math.sqrt(sum(e['error_m']**2 for e in errors)/len(errors)), 4),
            'errors': errors, 'passed': worst <= limit}


def check_reference(frames, crs, spec, tolerance, min_share, layers):
    """Share of the drawing's vertices (line and polygon layers) within tolerance of a reference layer's edges, read
    around the drawing and reprojected to the assigned CRS when it has another. A right CRS puts shared boundaries
    on top of each other; a shift of a few metres, which an extent comparison cannot see, does not."""
    import geopandas as gpd
    import numpy as np
    import pyogrio
    import shapely
    from pyproj import CRS, Transformer
    path, layer = spec.rsplit(':', 1) if ':' in spec and not Path(spec).exists() else (spec, None)
    info = pyogrio.read_info(path, layer=layer)
    if not info.get('crs'):
        raise ValueError(f'--check-reference {spec} has no CRS')
    ref_crs = CRS.from_user_input(info['crs'])
    names = layers or [n for n, f in frames.items() if hasattr(f, 'geometry') and
                       f.geometry.geom_type.isin(['LineString', 'MultiLineString', 'Polygon', 'MultiPolygon']).any()]
    if not names:
        raise ValueError('The source has no line or polygon layer to check against --check-reference')
    vertices, per_layer = [], {}
    for name in names:
        if name not in frames:
            raise ValueError(f'--check-layer {name} is not a layer of the source ({", ".join(frames)})')
        coords = np.unique(np.round(shapely.get_coordinates(frames[name].geometry.dropna().to_numpy()), 3), axis=0)
        per_layer[name] = coords
        vertices.append(coords)
    vertices = np.unique(np.concatenate(vertices), axis=0)
    if not len(vertices):
        raise ValueError('The checked layers have no vertices')
    bounds = [*vertices.min(axis=0), *vertices.max(axis=0)]
    pad = max(tolerance*10, 1.0)
    box = (bounds[0]-pad, bounds[1]-pad, bounds[2]+pad, bounds[3]+pad)
    if not ref_crs.equals(crs):
        box = Transformer.from_crs(crs, ref_crs, always_xy=True).transform_bounds(*box, densify_pts=21)
    check_vector(path, layer)
    reference = gpd.read_file(path, layer=layer, bbox=box, engine='pyogrio')
    reference = reference[reference.geometry.notna() & reference.geom_type.str.contains('LineString|Polygon')]
    if not ref_crs.equals(crs):
        reference = reference.to_crs(crs)
    edges = shapely.force_2d(reference.geometry.boundary.where(reference.geom_type.str.contains('Polygon'), reference.geometry).to_numpy())
    edges = edges[~shapely.is_empty(edges)]
    if not len(edges):
        raise ValueError(f'--check-reference {spec} has no lines or polygons around the drawing')
    tree = shapely.STRtree(edges)

    def distances(coords):
        _, found = tree.query_nearest(shapely.points(coords), return_distance=True, all_matches=False)
        return found
    near = distances(vertices)
    share = float(np.mean(near <= tolerance))
    return {'path': str(Path(path).absolute()), 'layer': layer, 'sha256': digest(path) if Path(path).is_file() else None,
            'crs': ref_crs.to_string(), 'reprojected': not ref_crs.equals(crs), 'features_read': int(len(reference)),
            'layers_checked': names, 'vertices': int(len(vertices)), 'tolerance_m': tolerance, 'min_share': min_share,
            'share_within_tolerance': round(share, 4), 'median_distance_m': round(float(np.median(near)), 4),
            'p10_distance_m': round(float(np.percentile(near, 10)), 4),
            'share_by_layer': {n: round(float(np.mean(distances(c) <= tolerance)), 4) for n, c in per_layer.items() if len(c)},
            'passed': share >= min_share}


def assign_crs(args):
    """Label CRS-less layers converted from an inspected drawing with a confirmed CRS; coordinates are not touched.
    Every layer of the source is labelled unless --layer names one; layer names are kept."""
    import warnings as python_warnings
    import pandas as pd
    import pyogrio
    import shapely
    from pyproj import CRS
    for flag, value in (('--confirmed-by', args.confirmed_by), ('--basis', args.basis)):
        if not value or not value.strip():
            raise ValueError(f'{flag} is required and must not be blank')
    crs = CRS.from_user_input(args.crs)
    unit = crs.axis_info[0].unit_name if crs.axis_info else None
    if not crs.is_projected or unit != 'metre':
        raise ValueError(f'--crs must be a projected CRS in metres; got {crs.name}')
    authority = crs.to_authority(min_confidence=70)
    code = ':'.join(authority) if authority else None
    inspection_path, inspection = load_inspection(args.inspection)

    source = Path(args.source)
    available = [name for name, _ in pyogrio.list_layers(source)]
    if args.layer and args.layer not in available:
        raise ValueError(f'{source} has no layer {args.layer} ({", ".join(available)})')
    names = [args.layer] if args.layer else available
    full = inspection['extent']['full']
    slack = 1e-9*max(abs(v) for v in full) + 1e-6
    layers = {}
    for name in names:
        info = pyogrio.read_info(source, layer=name, force_total_bounds=True)
        if info.get('crs') is not None:
            raise ValueError(f'{source}:{name} already has a CRS ({info["crs"]}); assigning would relabel it, reproject instead')
        bounds = info.get('total_bounds')
        bounds = [float(v) for v in bounds] if bounds is not None and all(v == v for v in bounds) else None
        if bounds and not (full[0]-slack <= bounds[0] and full[1]-slack <= bounds[1]
                           and bounds[2] <= full[2]+slack and bounds[3] <= full[3]+slack):
            raise ValueError(f'{source}:{name} extent {bounds} lies outside the inspected drawing\'s model-space extent {full}; '
                             'it is not this drawing as drawn (paper-space features, scaled or moved coordinates?)')
        layers[name] = {'bounds': bounds, 'geometry_type': info.get('geometry_type')}

    readings = inspection['crs']['readings']
    as_drawn = {c['crs'] for r in readings if r['scale_to_metre'] == 1.0 for m in r['matches'] for c in m['candidates']}
    scaled = {c['crs'] for r in readings if r['scale_to_metre'] != 1.0 for m in r['matches'] for c in m['candidates']}
    if code in scaled and code not in as_drawn:
        raise ValueError(f'{code} fits the drawing only after unit scaling; convert the units before assigning')
    outside = code not in as_drawn
    if outside and not args.outside_candidates:
        raise ValueError(f'{code or crs.name} is not among the inspection candidates as drawn '
                         f'({", ".join(sorted(as_drawn)) or "none"}); pass --outside-candidates to assign it anyway')
    if args.check_reference and (args.tolerance is None or args.min_share is None):
        raise ValueError('--check-reference needs --tolerance (metres) and --min-share (0-1); there are no defaults')
    if args.check_reference and not (args.tolerance > 0 and 0 < args.min_share <= 1):
        raise ValueError('--tolerance must be positive and --min-share in (0, 1]')
    if args.check_points and (args.max_check_error is None or args.max_check_error <= 0):
        raise ValueError('--check-points needs a positive --max-check-error (metres); there is no default')
    easting = false_easting(crs)
    prefixed = easting is not None and easting >= PREFIX_EASTING
    needed = not prefixed or outside
    if needed and not (args.check_reference or args.check_points):
        why = ('it is outside the inspection candidates' if outside else
               f'it has no zone prefix (false easting {easting:g}), the grids a city or project grid is most easily mistaken for'
               if easting is not None else 'its false easting is unknown')
        raise ValueError(f'Assigning {code or crs.name} needs a feature check, since {why}: --check-reference with --tolerance '
                         'and --min-share (shared boundaries), or --check-points with --max-check-error (known points)')

    frames = {}
    for name in names:
        check_vector(source, name)
        frames[name] = pyogrio.read_dataframe(source, layer=name)
    evidence = {'required': needed, 'reason': None if not needed else 'outside candidates' if outside else 'no zone prefix'}
    if args.check_reference:
        evidence['reference'] = check_reference(frames, crs, args.check_reference, args.tolerance, args.min_share,
                                                [n.strip() for n in args.check_layer.split(',')] if args.check_layer else None)
    if args.check_points:
        evidence['points'] = check_points(args.check_points, args.max_check_error)
    failed = [kind for kind in ('reference', 'points') if kind in evidence and not evidence[kind]['passed']]
    if failed:
        detail = {'reference': lambda e: f'{e["share_within_tolerance"]:.1%} of {e["vertices"]} vertices within {e["tolerance_m"]} m '
                                         f'of the reference edges, below --min-share {e["min_share"]:.1%}',
                  'points': lambda e: f'largest check-point error {e["max_error_m"]} m exceeds --max-check-error {e["max_check_error_m"]}'}
        raise ValueError('The feature check failed, nothing was published: ' + '; '.join(detail[k](evidence[k]) for k in failed))
    stat = source.stat()
    with bundle(args.output) as stage:
        target = stage / f'{source.stem}.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message='.*geometry column.*')
            for name, frame in frames.items():
                kind = layers[name]['geometry_type']
                pyogrio.write_dataframe(frame.set_crs(crs), target, layer=name, driver='GPKG',
                                        **({'geometry_type': kind} if kind and kind != 'Unknown' else {}))
        checks = {}
        for name, frame in frames.items():
            back = pyogrio.read_dataframe(target, layer=name)
            if len(back) != len(frame) or not (shapely.to_wkb(back.geometry.to_numpy()) == shapely.to_wkb(frame.geometry.to_numpy())).all():
                raise ValueError(f'Readback geometry of {name} differs from the input; nothing was published')
            try:
                pd.testing.assert_frame_equal(back.drop(columns=back.geometry.name).reset_index(drop=True),
                                              frame.drop(columns=frame.geometry.name).reset_index(drop=True), check_dtype=False)
            except AssertionError as error:
                raise ValueError(f'Readback attributes of {name} differ from the input; nothing was published ({error})') from error
            if not CRS.from_user_input(back.crs).equals(crs):
                raise ValueError(f'Readback CRS of {name} differs from --crs; nothing was published')
            checks[name] = {'features': len(back), 'bounds': [round(v, 3) for v in layers[name]['bounds']] if layers[name]['bounds'] else None,
                            'geometry_wkb_identical': True, 'attributes_identical': True}
        manifest = {'manifest': 'crs-assignment', 'schema_version': 1,
                    'input': {'name': source.name, 'sha256': digest(source) if source.is_file() else None, 'bytes': stat.st_size,
                              'layers': names},
                    'crs': {'id': code, 'name': crs.name, 'datum': crs.datum.name if crs.datum else None},
                    'confirmed_by': args.confirmed_by.strip(), 'basis': args.basis.strip(),
                    'confirmed_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                    'inspection': {'path': str(inspection_path.absolute()), 'sha256': digest(inspection_path),
                                   'drawing': inspection['source']['name'], 'drawing_sha256': inspection['source']['sha256'],
                                   'crs_status': inspection['crs']['status'], 'candidates_as_drawn': sorted(as_drawn)},
                    'outside_candidates': outside, 'zone_prefix': prefixed, 'evidence': evidence,
                    'checks': checks, 'crs_readback': code or crs.name,
                    'output': {'path': target.name, 'sha256': digest(target)}}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad assign-crs', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256']},
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['Only the CRS label changes; geometry WKB and attributes are compared on readback',
                                                            'The confirmation and its basis are the user\'s; the tool does not verify them',
                                                            'A feature check (evidence) is measured against thresholds the user declared; '
                                                            'it is required for grids without a zone prefix and outside candidates']})
    return manifest


OVERLAY_FIELDS = ('parcel_id', 'landuse', 'control_type', 'status')
REFERENCE_COLOURS = ('#1f77b4', '#2ca02c', '#9467bd', '#8c564b', '#17becf', '#e377c2', '#7f7f7f', '#bcbd22')


def layer_spec(text):
    """PATH[:LAYER]; split only when the whole text is not an existing path."""
    if ':' in text and not Path(text).exists():
        path, layer = text.rsplit(':', 1)
        return Path(path), layer or None
    return Path(text), None


def projected_metres(crs):
    unit = crs.axis_info[0].unit_name if crs.axis_info else None
    return crs.is_projected and unit == 'metre'


def polygon_parts(geometry):
    import shapely
    parts = [part for part in shapely.get_parts(geometry) if part.geom_type == 'Polygon' and not part.is_empty]
    return shapely.MultiPolygon(parts) if parts else None


def plot_overlay(parcels, references, qa, view, path, title, font):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = [font, 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    figure, axis = plt.subplots(figsize=(12, 9))
    held = parcels['status'] == 'hold' if 'status' in parcels else parcels.geometry.isna()
    parcels[~held].plot(ax=axis, facecolor='#eeeeee', edgecolor='none')
    handles = [plt.Rectangle((0, 0), 1, 1, facecolor='#eeeeee', edgecolor='#222222', label=f'parcels ({len(parcels)})')]
    for i, (name, frame) in enumerate(references):
        colour = REFERENCE_COLOURS[i % len(REFERENCE_COLOURS)]
        kinds = frame.geometry.geom_type.fillna('')
        areas, lines, points = (frame[kinds.str.contains(k)] for k in ('Polygon', 'LineString', 'Point'))
        if len(areas):
            areas.plot(ax=axis, facecolor=colour, edgecolor=colour, linewidth=1.0, alpha=0.18)
            areas.boundary.plot(ax=axis, color=colour, linewidth=1.0)
        if len(lines):
            lines.plot(ax=axis, color=colour, linewidth=1.0)
        if len(points):
            points.plot(ax=axis, color=colour, markersize=10)
        handles.append(plt.Rectangle((0, 0), 1, 1, facecolor=colour, edgecolor=colour, alpha=0.5, label=f'{name} ({len(frame)})'))
    parcels[~held].boundary.plot(ax=axis, color='#222222', linewidth=0.6)
    if held.any():
        parcels[held].boundary.plot(ax=axis, color='#e67e22', linewidth=1.6)
        handles.append(plt.Line2D([], [], color='#e67e22', linewidth=2, label=f'hold ({int(held.sum())})'))
    shown = 0
    for frame in qa.values():
        frame = frame[frame.geometry.notna() & ~frame.geometry.is_empty]
        kinds = frame.geometry.geom_type
        points = frame[kinds.str.contains('Point')].geometry.centroid
        if len(points):
            axis.scatter(points.x, points.y, marker='x', color='#d62728', s=30, zorder=5)
        lines = frame[kinds.str.contains('LineString')]
        if len(lines):
            lines.plot(ax=axis, color='#d62728', linewidth=1.4)
        areas = frame[kinds.str.contains('Polygon')]
        if len(areas):
            areas.boundary.plot(ax=axis, color='#d62728', linewidth=1.2)
        shown += len(frame)
    if shown:
        handles.append(plt.Line2D([], [], color='#d62728', marker='x', linestyle='', label=f'QA ({shown})'))
    if 'parcel_id' in parcels:
        for _, row in parcels[parcels['parcel_id'].fillna('') != ''].iterrows():
            point = row.geometry.representative_point()
            axis.annotate(row['parcel_id'], (point.x, point.y), fontsize=5, ha='center', va='center')
    axis.set_xlim(view[0], view[2])
    axis.set_ylim(view[1], view[3])
    axis.legend(handles=handles, loc='center left', bbox_to_anchor=(1.01, 0.5), fontsize=8, frameon=False)
    axis.set_aspect('equal', adjustable='datalim')
    axis.ticklabel_format(useOffset=False, style='plain')
    axis.tick_params(labelsize=7)
    axis.set_title(title, fontsize=10)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def overlay(args):
    """Parcels over reference layers in one check plot, with their intersections measured; nothing is repaired."""
    import warnings as python_warnings
    import geopandas as gpd
    import numpy as np
    import pyogrio
    import shapely
    from pyproj import CRS
    from pyproj.transformer import Transformer

    source = Path(args.source)
    available = [name for name, _ in pyogrio.list_layers(source)]
    if args.parcels_layer not in available:
        raise ValueError(f'{source} has no layer {args.parcels_layer} ({", ".join(available)})')
    check_vector(source, args.parcels_layer)
    parcels = pyogrio.read_dataframe(source, layer=args.parcels_layer, fid_as_index=True)
    if parcels.crs is None:
        raise ValueError(f'{source}:{args.parcels_layer} has no CRS; run cad.py assign-crs (or vector-georef.py) first')
    crs = CRS.from_user_input(parcels.crs)
    if not projected_metres(crs):
        raise ValueError(f'{source}:{args.parcels_layer} must be in a projected CRS in metres; got {crs.name}')
    parcels = parcels[parcels.geometry.notna() & ~parcels.geometry.is_empty]
    if not len(parcels):
        raise ValueError(f'{source}:{args.parcels_layer} has no geometry')
    if not parcels.geometry.geom_type.str.contains('Polygon').all():
        raise ValueError(f'{source}:{args.parcels_layer} must hold polygons; got {sorted(set(parcels.geometry.geom_type))}')
    held = parcels['status'] == 'hold' if 'status' in parcels else parcels.geometry.isna()
    measured = parcels[~held] if args.exclude_hold else parcels
    if not len(measured):
        raise ValueError(f'{source}:{args.parcels_layer} has no parcels left to measure without the hold ones')
    minx, miny, maxx, maxy = (float(v) for v in measured.total_bounds)
    pad = OVERLAY_MARGIN*max(maxx-minx, maxy-miny, 1.0)
    view = (minx-pad, miny-pad, maxx+pad, maxy+pad)
    parcel_valid = measured.geometry.is_valid.to_numpy()
    parcel_geoms = measured.geometry.to_numpy()
    parcel_area = float(shapely.union_all(parcel_geoms[parcel_valid]).area)

    references, report, rows, labels = [], [], [], set()
    for spec in args.reference:
        path, layer = layer_spec(spec)
        names = [name for name, _ in pyogrio.list_layers(path)]
        given = layer is not None
        if not given and len(names) != 1:
            raise ValueError(f'{path} has {len(names)} layers ({", ".join(names)}); name one as {path}:LAYER')
        layer = layer or names[0]
        if layer not in names:
            raise ValueError(f'{path} has no layer {layer} ({", ".join(names)})')
        info = pyogrio.read_info(path, layer=layer)
        if info.get('crs') is None:
            raise ValueError(f'{path}:{layer} has no CRS; a reference must carry one')
        ref_crs = CRS.from_user_input(info['crs'])
        reprojected = not ref_crs.equals(crs)
        bbox = Transformer.from_crs(crs, ref_crs, always_xy=True).transform_bounds(*view, densify_pts=21) if reprojected else view
        check_vector(path, layer)
        frame = pyogrio.read_dataframe(path, layer=layer, bbox=tuple(bbox), fid_as_index=True)
        frame = frame[frame.geometry.notna() & ~frame.geometry.is_empty]
        dropped_z = bool(frame.geometry.has_z.any())
        if dropped_z:
            frame = frame.set_geometry(shapely.force_2d(frame.geometry.to_numpy()), crs=frame.crs)
        if reprojected:
            frame = frame.to_crs(crs)
        name = layer if given else path.stem
        base, n = name, 2
        while name in labels:
            name, n = f'{base}#{n}', n+1
        labels.add(name)
        kinds = frame.geometry.geom_type
        areal = kinds.str.contains('Polygon').to_numpy()
        valid = frame.geometry.is_valid.to_numpy()
        invalid = [int(fid) for fid in frame.index[areal & ~valid]]
        entry = {'label': name, 'path': str(path.absolute()), 'layer': layer, 'features': len(frame),
                 'geometry_types': sorted(set(kinds)), 'crs': ':'.join(ref_crs.to_authority()) if ref_crs.to_authority() else ref_crs.name,
                 'reprojected': reprojected, 'dropped_z': dropped_z, 'invalid_fids': invalid,
                 'parcels_touched': 0, 'intersection_area_m2': 0.0, 'share_of_parcels': 0.0}
        usable = frame[areal & valid]
        if len(usable):
            ref_geoms = usable.geometry.to_numpy()
            tree = shapely.STRtree(ref_geoms)
            left, right = tree.query(np.where(parcel_valid, parcel_geoms, None), predicate='intersects')
            pieces = shapely.intersection(parcel_geoms[left], ref_geoms[right])
            areas = shapely.area(pieces)
            keep = areas > 0
            left, right, pieces, areas = left[keep], right[keep], pieces[keep], areas[keep]
            parcel_areas, ref_areas = shapely.area(parcel_geoms[left]), shapely.area(ref_geoms[right])
            for i, j, piece, area, pa, ra in zip(left, right, pieces, areas, parcel_areas, ref_areas):
                row = measured.iloc[i]
                rows.append({'parcel_fid': int(measured.index[i]), **{f: row[f] for f in OVERLAY_FIELDS if f in parcels},
                             'reference': name, 'reference_fid': int(usable.index[j]), 'area_m2': round(float(area), 3),
                             'parcel_share': round(float(area/pa), 6), 'reference_share': round(float(area/ra), 6),
                             'geometry': polygon_parts(piece)})
            if len(pieces):
                union = float(shapely.union_all(pieces).area)
                entry.update(parcels_touched=len(set(left.tolist())), intersection_area_m2=round(union, 3),
                             share_of_parcels=round(union/parcel_area, 6) if parcel_area else None)
        references.append((name, frame))
        report.append(entry)

    if args.qa_layers is None:
        qa_names = [name for name in available if name != args.parcels_layer and (name.startswith('qa_') or name == 'label_conflicts')]
    else:
        qa_names = [name.strip() for name in args.qa_layers.split(',') if name.strip()]
        missing = [name for name in qa_names if name not in available]
        if missing:
            raise ValueError(f'{source} has no QA layer {", ".join(missing)} ({", ".join(available)})')
    qa = {}
    for name in qa_names:
        frame = pyogrio.read_dataframe(source, layer=name)
        if frame.crs is None:
            raise ValueError(f'{source}:{name} has no CRS; QA layers must share the parcels\' CRS')
        qa[name] = frame if CRS.from_user_input(frame.crs).equals(crs) else frame.to_crs(crs)
    hold = {'count': int(held.sum()), 'measured': not args.exclude_hold,
            'parcel_fids': [int(fid) for fid in parcels.index[held]],
            'parcel_ids': sorted(str(v) for v in parcels.loc[held, 'parcel_id'].dropna()) if 'parcel_id' in parcels else []}
    columns = ['parcel_fid', *[f for f in OVERLAY_FIELDS if f in parcels], 'reference', 'reference_fid', 'area_m2',
               'parcel_share', 'reference_share']
    intersections = gpd.GeoDataFrame(rows, columns=[*columns, 'geometry'], geometry='geometry', crs=crs)
    try:
        from _theme import get_chinese_font
        font = get_chinese_font()
    except ImportError:
        font = 'DejaVu Sans'
    code = ':'.join(crs.to_authority()) if crs.to_authority() else crs.name
    with bundle(args.output) as stage:
        target = stage / 'overlay.gpkg'
        with python_warnings.catch_warnings():
            python_warnings.filterwarnings('ignore', message='.*geometry column.*')
            pyogrio.write_dataframe(intersections, target, layer='intersections', driver='GPKG', geometry_type='MultiPolygon')
        back = int(pyogrio.read_info(target, layer='intersections', force_feature_count=True)['features'])
        if back != len(intersections):
            raise ValueError(f'Readback count {back} of intersections differs from {len(intersections)}; nothing was published')
        plot_overlay(parcels, references, qa, view, stage / 'check.png',
                     f'{source.name}:{args.parcels_layer} · {len(parcels)} parcels · {hold["count"]} hold · '
                     f'{len(references)} references', font)
        manifest = {'manifest': 'cad-overlay', 'schema_version': 1,
                    'input': {'name': source.name, 'layer': args.parcels_layer, 'sha256': digest(source) if source.is_file() else None,
                              'features': len(parcels), 'measured': len(measured), 'measured_area_m2': round(parcel_area, 3),
                              'invalid_fids': [int(fid) for fid in measured.index[~parcel_valid]]},
                    'crs': code, 'view': [round(v, 3) for v in view], 'margin': OVERLAY_MARGIN,
                    'references': report, 'hold': hold, 'qa': {name: len(frame) for name, frame in qa.items()},
                    'intersections': len(intersections),
                    'output': {'path': target.name, 'sha256': digest(target)}, 'plot': 'check.png'}
        write_json(stage / 'manifest.json', manifest)
        write_json(stage / 'record.json', {'schema_version': 1, 'status': 'complete', 'tool': 'cad overlay', 'version': VERSION,
                                            'parameters': {k: v for k, v in vars(args).items() if k not in ('output', 'command')},
                                            'source': {'path': str(source.absolute()), 'sha256': manifest['input']['sha256']},
                                            'files': {'manifest.json': digest(stage / 'manifest.json'), target.name: manifest['output']['sha256'],
                                                      'check.png': digest(stage / 'check.png')},
                                            'finished_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                                            'assumptions': ['References are read only within the view (parcels\' extent plus the margin) and '
                                                            'reprojected to the parcels\' CRS; Z is dropped',
                                                            'Invalid parcels and reference polygons are plotted and listed, not repaired or measured',
                                                            *(['Hold parcels are plotted and listed, not measured'] if args.exclude_hold else []),
                                                            'Areas are planar, in the parcels\' CRS']})
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest='command', required=True)
    sub = commands.add_parser('inspect', help='Inventory a DXF and list candidate CRS')
    sub.add_argument('source', help='DXF file')
    sub.add_argument('--output', required=True, help='New directory for cad.json, check.png and record.json')
    sub.add_argument('--near', help='Approximate LON,LAT of the area the drawing covers; candidates farther than --near-km are rejected')
    sub.add_argument('--near-km', type=float, default=50.0, help='Keep candidates within this distance of --near (default 50)')
    sub.add_argument('--encoding', help='Text encoding overriding $DWGCODEPAGE, e.g. CP936 for Chinese drawings without one')
    sub.add_argument('--reference', help='Layer in a known projected CRS (metres) covering the same area; extents are compared')
    sub.add_argument('--reference-layer', help='Layer name within --reference')
    pull = commands.add_parser('extract', help='Extract model-space entities of chosen layers to raw.gpkg, in metres')
    pull.add_argument('source', help='DXF file')
    pull.add_argument('--layers', help='Comma-separated CAD layers (default: every layer with model-space entities)')
    pull.add_argument('--units', required=True, choices=('m', 'mm', 'header'),
                      help='Drawing unit: m, mm, or header ($INSUNITS; refused when unitless or contradicted by the coordinates)')
    pull.add_argument('--arc-tolerance', type=float, default=0.05, help='Largest chord-to-curve distance in metres (default 0.05)')
    pull.add_argument('--include-blocks', action='store_true', help='Explode block references instead of keeping insertion points')
    pull.add_argument('--drop-z', action='store_true', help='Flatten non-zero Z; without it such features are refused')
    pull.add_argument('--encoding', help='Text encoding overriding $DWGCODEPAGE, e.g. CP936')
    pull.add_argument('--output', required=True, help='New directory for raw.gpkg, manifest.json and record.json')
    faces = commands.add_parser('polygonize', help='Snap, node and polygonize extracted lines into candidate faces')
    faces.add_argument('source', help='raw.gpkg from cad.py extract (or a line layer with a "layer" field, in metres)')
    faces.add_argument('--lines-layer', default='lines', help='Line layer within source (default lines)')
    faces.add_argument('--cad-layers', help='Comma-separated CAD layers to use (default: every layer in the line layer)')
    faces.add_argument('--boundary', help='CAD layer whose lines enclose the planning area; faces outside are dropped and counted')
    faces.add_argument('--snap', type=float, required=True, help='Merge endpoints, and move free ends onto lines, within this many metres')
    faces.add_argument('--max-gap', type=float, required=True, help='Report gaps up to this many metres (must exceed --snap)')
    faces.add_argument('--min-area', type=float, required=True, help='Faces smaller than this many m² are held as suspected slivers')
    faces.add_argument('--exclude', help='YAML rules leaving lines out: layers / blocks (name patterns), boxes ([minx, miny, maxx, maxy])')
    faces.add_argument('--exclude-area', help='Polygons (file or file:layer) whose lines are left out, e.g. title blocks and legends')
    faces.add_argument('--close-gaps', action='store_true', help='Close reported gaps with added lines, kept in added_lines and marked')
    faces.add_argument('--output', required=True, help='New directory for faces.gpkg, check.png, manifest.json and record.json')
    compare = commands.add_parser('compare-fills', help='Compare hatch areas from extract with polygonized or 湘源 faces')
    compare.add_argument('faces', help='Faces: faces.gpkg from polygonize, parcels.gpkg from label or xiangyuan')
    compare.add_argument('--faces-layer', default='faces', help='Polygon layer within faces (default faces; parcels for label / xiangyuan)')
    compare.add_argument('--fills', required=True, help='raw.gpkg from cad.py extract')
    compare.add_argument('--fills-layer', default='hatches', help='Hatch layer within --fills (default hatches)')
    compare.add_argument('--cad-layers', help='Comma-separated CAD layers of the hatches to use (default: all)')
    compare.add_argument('--exclude-hold', action='store_true', help='Leave out faces whose status is hold (counted)')
    compare.add_argument('--output', required=True, help='New directory for fill_compare.gpkg, check.png, manifest.json and record.json')
    tag = commands.add_parser('label', help='Attach parcel fields from CAD texts to polygonized faces by rules')
    tag.add_argument('faces', help='faces.gpkg from cad.py polygonize')
    tag.add_argument('--faces-layer', default='faces')
    tag.add_argument('--texts', required=True, help='raw.gpkg from cad.py extract (its texts and leaders layers are read)')
    tag.add_argument('--texts-layer', default='texts')
    tag.add_argument('--rules', required=True, help='YAML label rules: fields, combined, landuse_major, road')
    tag.add_argument('--output', required=True, help='New directory for parcels.gpkg, check.png, rules.yaml, manifest.json and record.json')
    xy = commands.add_parser('xiangyuan', help='Parcels of a 湘源控规 drawing from closed polylines and their XDATA')
    xy.add_argument('source', help='DXF of a 湘源控规 drawing (convert DWG first)')
    xy.add_argument('--units', required=True, choices=('m', 'mm', 'header'), help='Drawing unit, as for extract')
    xy.add_argument('--arc-tolerance', type=float, default=0.05, help='Largest chord-to-arc distance in metres (default 0.05)')
    xy.add_argument('--max-area-diff', type=float, help='Hold parcels whose area differs from 用地面积 by more than this ratio, e.g. 0.005')
    xy.add_argument('--boundary-layer', default='----规划边界', help='Layer of the planning boundary (default ----规划边界; empty to skip)')
    xy.add_argument('--unit-layer', default='DY范围', help='Layer of unit outlines (default DY范围; empty to skip)')
    xy.add_argument('--unit-label-layer', default='DY编号', help='Layer of the unit labels inside them (default DY编号)')
    xy.add_argument('--facility-layer', default='配套设施', help='Layer of facility blocks (default 配套设施; empty to skip)')
    xy.add_argument('--schema', help='Optional JSON naming more 湘源 fields (land_attribute_fields_order, gis_field_mapping) '
                                     'and numeric ones (numeric_like)')
    xy.add_argument('--encoding', help='Text encoding overriding the drawing')
    xy.add_argument('--output', required=True, help='New directory for parcels.gpkg, check.png, manifest.json and record.json')
    dwg = commands.add_parser('convert', help='Convert DWG to DXF with an installed ODA File Converter or LibreDWG')
    dwg.add_argument('source', help='DWG file')
    dwg.add_argument('--tool', choices=('auto', 'oda', 'libredwg'), default='auto', help='auto prefers ODA, then LibreDWG')
    dwg.add_argument('--version', choices=tuple(DXF_VERSIONS), default='r2018', help='DXF version to write (default r2018)')
    dwg.add_argument('--timeout', type=float, default=600, help='Seconds before the converter is stopped (default 600)')
    dwg.add_argument('--output', required=True, help='New directory for <source>.dxf, manifest.json and record.json')
    assign = commands.add_parser('assign-crs', help='Label a CRS-less layer from an inspected drawing with a confirmed CRS')
    assign.add_argument('source', help='Vector file converted from the drawing, without a CRS (e.g. convert.py output)')
    assign.add_argument('--layer', help='Only this layer of source (default: every layer, names kept)')
    assign.add_argument('--inspection', required=True, help='cad.json (or its directory) from cad.py inspect of the drawing')
    assign.add_argument('--crs', required=True, help='Projected CRS in metres, e.g. EPSG:4547')
    assign.add_argument('--confirmed-by', required=True, help='Who confirmed the CRS')
    assign.add_argument('--basis', required=True, help='What the confirmation rests on, e.g. provider statement, control-point check')
    assign.add_argument('--outside-candidates', action='store_true', help='Allow a CRS the inspection did not list; recorded')
    assign.add_argument('--check-reference', help='Layer (file or file:layer) sharing boundaries with the drawing, in a known CRS')
    assign.add_argument('--check-layer', help='Comma-separated source layers whose vertices are checked (default: line and polygon layers)')
    assign.add_argument('--tolerance', type=float, help='Metres: a vertex within this of a reference edge counts as on it')
    assign.add_argument('--min-share', type=float, help='Share of vertices (0-1) that must lie within --tolerance')
    assign.add_argument('--check-points', help='CSV of known points (id, src_x, src_y, dst_x, dst_y), dst in --crs')
    assign.add_argument('--max-check-error', type=float, help='Largest acceptable distance in metres between src and dst of a point')
    assign.add_argument('--output', required=True, help='New directory for <source>.gpkg, manifest.json and record.json')
    lay = commands.add_parser('overlay', help='Parcels over reference layers in one check plot, intersections measured')
    lay.add_argument('source', help='Vector file with parcels in a projected CRS, e.g. parcels.gpkg after assign-crs')
    lay.add_argument('--parcels-layer', default='parcels', help='Parcel layer within source (default parcels)')
    lay.add_argument('--reference', action='append', required=True,
                     help='PATH[:LAYER] of a reference layer with a CRS (repeatable); a single-layer file needs no layer')
    lay.add_argument('--qa-layers', help='Comma-separated layers of source to mark as QA (default: qa_* and label_conflicts; '
                                         'empty for none)')
    lay.add_argument('--exclude-hold', action='store_true',
                     help='Plot and list hold parcels but leave them out of the intersections and the view extent')
    lay.add_argument('--output', required=True, help='New directory for overlay.gpkg, check.png, manifest.json and record.json')
    args = parser.parse_args()
    try:
        if args.command == 'overlay':
            manifest = overlay(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'parcels': manifest['input']['features'],
                              'hold': manifest['hold']['count'], 'intersections': manifest['intersections'],
                              'references': {r['label']: {k: r[k] for k in ('features', 'parcels_touched', 'intersection_area_m2',
                                                                            'share_of_parcels', 'invalid_fids')}
                                             for r in manifest['references']},
                              'qa': manifest['qa']}, ensure_ascii=False))
            return
        if args.command == 'xiangyuan':
            manifest = xiangyuan(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'layers': manifest['layers'],
                              'summary': manifest['summary'], 'area': manifest['area'], 'warnings': manifest['warnings']},
                             ensure_ascii=False))
            return
        if args.command == 'convert':
            manifest = convert(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'tool': manifest['converter']['tool'],
                              'dwg': manifest['source']['release'] or manifest['source']['dwg_version'],
                              'readback': {k: v for k, v in manifest['readback'].items() if k != 'header'}}, ensure_ascii=False))
            return
        if args.command == 'extract':
            manifest = extract(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'layers': manifest['layers'],
                              'skipped': manifest['skipped'], 'warnings': manifest['warnings']}, ensure_ascii=False))
            return
        if args.command == 'polygonize':
            manifest = polygonize(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'layers': manifest['layers'], 'faces': manifest['faces'],
                              'area': manifest['area']}, ensure_ascii=False))
            return
        if args.command == 'compare-fills':
            manifest = compare_fills(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'summary': manifest['summary'], 'area': manifest['area'],
                              'warnings': manifest['warnings']}, ensure_ascii=False))
            return
        if args.command == 'label':
            manifest = label(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'parcels': manifest['parcels'],
                              'conflicts': manifest['conflicts'], 'coverage': manifest['coverage']}, ensure_ascii=False))
            return
        if args.command == 'assign-crs':
            manifest = assign_crs(args)
            print(json.dumps({'output': str(Path(args.output).absolute()), 'crs': manifest['crs']['id'],
                              'features': {k: v['features'] for k, v in manifest['checks'].items()},
                              'outside_candidates': manifest['outside_candidates'],
                              'evidence': {k: {f: v[f] for f in ('share_within_tolerance', 'median_distance_m', 'max_error_m', 'passed') if f in v}
                                           for k, v in manifest['evidence'].items() if isinstance(v, dict)}},
                             ensure_ascii=False))
            return
        if args.near_km <= 0:
            raise ValueError('--near-km must be positive')
        report = inspect(args)
        candidates = [row['crs'] for reading in report['crs']['readings'] for match in reading['matches'] for row in match['candidates']]
        print(json.dumps({'output': str(Path(args.output).absolute()), 'layers': len(report['layers']),
                          'entities': report['entities']['model_space'], 'crs_status': report['crs']['status'],
                          'candidates': candidates, 'warnings': report['warnings']}, ensure_ascii=False))
    except (ValueError, OSError, RuntimeError) as exc:  # pyogrio's DataSourceError is a RuntimeError
        parser.exit(1, f'ERROR: {exc}\n')


if __name__ == '__main__':
    main()
