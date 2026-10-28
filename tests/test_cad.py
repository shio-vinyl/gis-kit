"""cad.py inspect: header, layer inventory, outliers and candidate CRS on hand-written R12 DXF files."""
from __future__ import annotations

import json
import math
import os
import shutil
from pathlib import Path
import subprocess
import sys

import geopandas as gpd
import pyogrio
import pytest
import shapely
from shapely.geometry import LineString, Point, Polygon, box

SCRIPTS = Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0, str(SCRIPTS))


def pairs(*items):
    return ''.join(f'{code}\n{value}\n' for code, value in items)


def dxf(path, entities, units=6, codepage='ANSI_936', encoding='gbk', header=True):
    """Minimal R12 DXF: HEADER (variables optional), a block with one circle, then the given ENTITIES pairs."""
    variables = [(9, '$ACADVER'), (1, 'AC1009'), (9, '$INSUNITS'), (70, units), (9, '$DWGCODEPAGE'), (3, codepage)]
    text = pairs((0, 'SECTION'), (2, 'HEADER'), *(variables if header else []), (0, 'ENDSEC'))  # GDAL needs the section
    text += pairs((0, 'SECTION'), (2, 'BLOCKS'), (0, 'BLOCK'), (8, '0'), (2, 'TREE'), (70, 0), (10, 0), (20, 0), (30, 0),
                  (0, 'CIRCLE'), (8, '0'), (10, 0), (20, 0), (30, 0), (40, 1), (0, 'ENDBLK'), (8, '0'), (0, 'ENDSEC'))
    text += pairs((0, 'SECTION'), (2, 'ENTITIES'), *entities, (0, 'ENDSEC'), (0, 'EOF'))
    path.write_text(text, encoding=encoding)
    return path


def line(layer, x0, y0, x1, y1, paper=False):
    return [(0, 'LINE'), (8, layer), *([(67, 1)] if paper else []), (10, x0), (20, y0), (30, 0), (11, x1), (21, y1), (31, 0)]


def label(layer, x, y, text):
    return [(0, 'TEXT'), (8, layer), (10, x), (20, y), (30, 0), (40, 2.5), (1, text)]


def site(x0, y0, step=1.0):
    """Roads, a closed building outline, a street name and a block reference around (x0, y0)."""
    s = step
    return [*line('道路', x0, y0, x0+100*s, y0+50*s), *line('道路', x0, y0+80*s, x0+120*s, y0+80*s),
            (0, 'POLYLINE'), (8, '建筑'), (66, 1), (70, 1), (10, 0), (20, 0), (30, 0),
            *[item for x, y in ((10, 10), (30, 10), (30, 30)) for item in ((0, 'VERTEX'), (8, '建筑'), (10, x0+x*s), (20, y0+y*s), (30, 0))],
            (0, 'SEQEND'), (8, '建筑'), *label('注记', x0+20*s, y0+20*s, '人民路'),
            (0, 'INSERT'), (8, '绿化'), (2, 'TREE'), (10, x0+50*s), (20, y0+5*s), (30, 0)]


# Runs cad.py as __main__ with ezdxf made unimportable, to exercise the GDAL-only inventory.
WITHOUT_EZDXF = ("import runpy, sys; sys.modules['ezdxf'] = None; sys.path.insert(0, sys.argv[1]); "
                 "sys.argv = sys.argv[1:]; sys.argv[0] = sys.argv[0] + '/cad.py'; runpy.run_path(sys.argv[0], run_name='__main__')")


def inspect(source, out, *extra, success=True, ezdxf=True):
    command = [sys.executable, str(SCRIPTS/'cad.py')] if ezdxf else [sys.executable, '-c', WITHOUT_EZDXF, str(SCRIPTS)]
    result = subprocess.run([*command, 'inspect', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return (json.loads((out/'cad.json').read_text()) if success else result.stderr)


def codes(report):
    return [row['crs'] for reading in report['crs']['readings'] for match in reading['matches'] for row in match['candidates']]


def test_zone_prefixed_grid_inventory_and_outputs(tmp_path):
    pytest.importorskip('ezdxf')
    source = dxf(tmp_path/'site.dxf', site(39_500_000, 3_400_000))
    report = inspect(source, tmp_path/'out')
    assert report['header'] == {'readable': True, 'acadver': 'AC1009', 'release': 'R12', 'insunits': 6, 'units': 'meters',
                                'codepage': 'ANSI_936', 'extent': None, 'handles': False}
    layers = {row['layer']: row for row in report['layers']}
    assert set(layers) == {'道路', '建筑', '注记', '绿化'} and layers['道路']['entities'] == 2
    assert layers['注记']['text_samples'] == [{'raw': '人民路', 'clean': '人民路'}] and 'texts' not in layers['道路']
    assert layers['道路']['dxf_types'] == {'LINE': 2} and layers['建筑']['dxf_types'] == {'POLYLINE closed': 1}
    assert layers['绿化']['blocks'] == {'TREE': 1} and report['reader']['missing'] == []
    assert report['entities']['model_space'] == 5 and report['reader']['inventory'].startswith('ezdxf')
    assert report['crs']['status'] == 'candidate' and report['crs']['assigned'] is None
    assert codes(report) == ['EPSG:4527', 'EPSG:4775', 'EPSG:2363', 'EPSG:2415']  # CGCS2000 first, all four datums
    lon, lat = report['crs']['readings'][0]['matches'][0]['candidates'][0]['lonlat']
    assert lon == pytest.approx(117, abs=0.01) and lat == pytest.approx(30.72, abs=0.01)
    assert (tmp_path/'out'/'check.png').stat().st_size > 10_000
    record = json.loads((tmp_path/'out'/'record.json').read_text())
    assert record['source']['sha256'] == report['source']['sha256'] and set(record['files']) == {'cad.json', 'check.png'}
    assert 'new directory' in inspect(source, tmp_path/'out', success=False)
    assert not list(tmp_path.glob('.delivery-*'))


def test_unprefixed_grid_is_held_even_with_near(tmp_path):
    source = dxf(tmp_path/'hefei.dxf', site(518_900, 3_520_000))  # CGCS2000 CM 117E without prefix, or a local grid
    for name, extra in (('held', ()), ('near', ('--near', '117.2,31.8'))):
        report = inspect(source, tmp_path/name, *extra)
        assert report['crs']['status'] == 'hold' and codes(report) == [] and report['crs']['rejected'] == []
        assert any('local grid' in note and 'no EPSG code' in note for note in report['crs']['notes'])


def test_near_rejects_distant_candidates(tmp_path):
    source = dxf(tmp_path/'site.dxf', site(39_500_000, 3_400_000))
    kept = inspect(source, tmp_path/'kept', '--near', '117.2,31.0')
    assert codes(kept)[0] == 'EPSG:4527'
    assert all(row['distance_km'] <= 50 for match in kept['crs']['readings'][0]['matches'] for row in match['candidates'])
    far = inspect(source, tmp_path/'far', '--near', '87.6,43.8')  # Urumqi
    assert codes(far) == [] and far['crs']['status'] == 'hold'
    assert {row['crs'] for row in far['crs']['rejected']} == {'EPSG:4527', 'EPSG:4775', 'EPSG:2363', 'EPSG:2415'}


def test_millimetre_drawing_uses_header_units(tmp_path):
    source = dxf(tmp_path/'mm.dxf', site(39_500_000_000, 3_400_000_000, step=1000), units=4)
    report = inspect(source, tmp_path/'out')
    first, second = report['crs']['readings']
    assert first['matches'] == [] and second['scale_to_metre'] == 0.001 and second['basis'] == 'header $INSUNITS = millimeters'
    assert codes(report)[0] == 'EPSG:4527'


def test_outliers_paper_space_and_local_grid(tmp_path):
    entities = site(1200, 800) + line('图框', 0, 0, 0.5, 0.5) + line('图框', 9e5, 9e5, 9e5+1, 9e5+1) \
        + line('布局', 0, 0, 420, 297, paper=True)
    source = dxf(tmp_path/'local.dxf', entities, header=False)  # no $DWGCODEPAGE: GDAL falls back to Latin-1
    auto = inspect(source, tmp_path/'auto')  # GBK read as Latin-1 is noticed and read again
    assert auto['encoding']['source'] == 'auto' and auto['encoding']['used'] == 'CP936'
    assert auto['encoding']['tried'] == {'CP936': 0, 'GB18030': 0} and '道路' in {r['layer'] for r in auto['layers']}
    assert any('read again as CP936' in w for w in auto['warnings'])
    garbled = inspect(source, tmp_path/'garbled', '--encoding', 'CP1252')  # an explicit choice is kept, however wrong
    assert garbled['encoding']['source'] == 'option' and '道路' not in {r['layer'] for r in garbled['layers']}
    report = inspect(source, tmp_path/'out', '--encoding', 'CP936')
    assert report['header']['readable'] is False and report['entities']['paper_space'] == 1
    assert report['extent']['outliers'] == 2 and report['extent']['core'][2] < 2000
    assert report['crs']['status'] == 'hold'
    assert any('图框 2' in w for w in report['warnings']) and any('paper-space' in w for w in report['warnings'])
    assert not any('--encoding' in w for w in report['warnings'])


def test_geographic_coordinates(tmp_path):
    report = inspect(dxf(tmp_path/'lonlat.dxf', site(117.2, 31.8, step=0.0001)), tmp_path/'out')
    assert codes(report)[:2] == ['EPSG:4490', 'EPSG:4326']


def test_r2000_entity_types_and_hatch_is_not_text(tmp_path):
    pytest.importorskip('ezdxf')
    frame = gpd.GeoDataFrame({'Layer': ['A', 'B', 'C']}, geometry=[LineString([(0, 0), (10, 10)]), Point(5, 5),
                                                                   Polygon([(0, 0), (10, 0), (10, 10)])])
    source = tmp_path/'r2000.dxf'
    with pytest.warns(UserWarning):
        pyogrio.write_dataframe(frame, source, driver='DXF')
    report = inspect(source, tmp_path/'out', ezdxf=False)
    layers = {row['layer']: row for row in report['layers']}
    assert layers['C']['entity_types'] == {'AcDbHatch': 1} and 'texts' not in layers['C']
    assert report['header']['release'] == 'R2004'
    full = {row['layer']: row for row in inspect(source, tmp_path/'ezdxf')['layers']}
    assert full['C']['dxf_types'] == {'HATCH': 1} and 'texts' not in full['C']


def test_gdal_fallback_lists_what_is_missing(tmp_path):
    source = dxf(tmp_path/'site.dxf', site(39_500_000, 3_400_000))
    report = inspect(source, tmp_path/'out', ezdxf=False)
    assert report['reader']['inventory'] == 'gdal' and 'block names and attribute tags' in report['reader']['missing']
    layers = {row['layer']: row for row in report['layers']}
    assert layers['注记']['text_samples'] == [{'raw': '人民路'}] and 'dxf_types' not in layers['注记']
    assert report['layouts'] is None and report['unused_layers'] is None
    assert codes(report)[0] == 'EPSG:4527'  # geometry, extent and CRS do not depend on ezdxf


def test_ezdxf_inventory_of_an_r2018_drawing(tmp_path):
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018', setup=True)
    doc.header['$INSUNITS'] = 6
    for name in ('地块界线', '地块编号', '用地代码', '图框', '辅助线'):
        doc.layers.add(name)
    doc.layers.get('辅助线').off()
    doc.layers.get('图框').freeze()
    msp = doc.modelspace()
    x0, y0 = 38_500_000, 2_490_000
    msp.add_lwpolyline([(x0, y0), (x0+100, y0), (x0+100, y0+80), (x0, y0+80)], close=True, dxfattribs={'layer': '地块界线'})
    msp.add_lwpolyline([(x0+100, y0), (x0+180, y0), (x0+180, y0+80)], dxfattribs={'layer': '地块界线'})
    msp.add_mtext(r'{\fSimSun|b0;03-03-02}\P', dxfattribs={'layer': '地块编号', 'insert': (x0+50, y0+40)})
    msp.add_text('Ｒ２＋Ｃ１', dxfattribs={'layer': '用地代码', 'insert': (x0+50, y0+30)})
    msp.add_line((x0, y0), (x0+180, y0+80), dxfattribs={'layer': '辅助线'})
    block = doc.blocks.new('地块标签')
    block.add_attdef('BH', (0, 0))
    block.add_circle((0, 0), 1)
    for point, value in (((x0+20, y0+20), '03-01'), ((x0+140, y0+20), '03-02')):
        msp.add_blockref('地块标签', point, dxfattribs={'layer': '地块编号'}).add_attrib('BH', value, point)
    layout = doc.layouts.new('图则')
    layout.add_viewport((150, 100), (200, 150), (x0+90, y0+40), 100)
    source = tmp_path/'r2018.dxf'
    doc.saveas(source)

    report = inspect(source, tmp_path/'out')
    layers = {row['layer']: row for row in report['layers']}
    assert layers['地块界线']['dxf_types'] == {'LWPOLYLINE closed': 1, 'LWPOLYLINE open': 1}
    assert layers['地块编号']['dxf_types'] == {'INSERT': 2, 'MTEXT': 1}
    assert layers['地块编号']['blocks'] == {'地块标签': 2} and layers['地块编号']['attributes'] == {'BH': ['03-01', '03-02']}
    assert layers['地块编号']['text_samples'] == [{'raw': r'{\fSimSun|b0;03-03-02}\P', 'clean': '03-03-02'}]
    assert layers['用地代码']['text_samples'] == [{'raw': 'Ｒ２＋Ｃ１', 'clean': 'R2+C1'}]
    assert layers['辅助线']['state'] == ['off'] and 'state' not in layers['地块界线']
    assert report['unused_layers'] == ['图框']
    assert report['encoding'] == {'declared': 'ANSI_1252', 'used': 'UTF-8', 'source': 'dxf-version', 'undecodable_strings': 0}
    assert layers['地块界线']['bounds'] is not None  # GDAL names match ezdxf names only when GDAL reads UTF-8 too
    assert {row['name']: row['viewports'] for row in report['layouts']}['图则'] == 1
    assert report['header']['release'] == 'R2018' and codes(report)[0] == 'EPSG:4526'


def test_refuses_dwg_and_other_inputs(tmp_path):
    (tmp_path/'plan.dwg').write_bytes(b'AC1032')
    assert 'save it as DXF' in inspect(tmp_path/'plan.dwg', tmp_path/'a', success=False)
    assert '.dxf' in inspect(tmp_path/'missing.dxf', tmp_path/'b', success=False)
    assert not (tmp_path/'a').exists() and not (tmp_path/'b').exists()


def reference(path, crs, bounds):
    gpd.GeoDataFrame({'name': ['范围']}, geometry=[box(*bounds)], crs=crs).to_file(path)
    return path


def test_reference_overlap_lists_its_crs_for_unprefixed_coordinates(tmp_path):
    source = dxf(tmp_path/'qianhai.dxf', site(495_000, 2_490_000))  # CGCS2000 CM 114E without prefix
    ref = reference(tmp_path/'range.gpkg', 'EPSG:4547', (494_000, 2_489_000, 497_000, 2_492_000))
    report = inspect(source, tmp_path/'out', '--reference', ref)
    comparison = report['crs']['reference']['comparisons'][0]
    assert comparison['verdict'] == 'extents-overlap' and comparison['overlap'] == 1.0
    assert report['crs']['status'] == 'candidate' and codes(report) == ['EPSG:4547']
    match = report['crs']['readings'][0]['matches'][0]
    assert match['grid'] == 'reference layer (extents-overlap)' and match['candidates'][0]['basis'] == 'reference'
    record = json.loads((tmp_path/'out'/'record.json').read_text())
    assert record['reference']['sha256'] and record['reference']['layer'] == 'range'


def test_reference_zone_prefix_both_ways(tmp_path):
    bare = dxf(tmp_path/'bare.dxf', site(495_000, 2_490_000))
    prefixed_ref = reference(tmp_path/'zone38.gpkg', 'EPSG:4526', (38_494_000, 2_489_000, 38_497_000, 2_492_000))
    report = inspect(bare, tmp_path/'a', '--reference', prefixed_ref)
    comparison = report['crs']['reference']['comparisons'][0]
    assert comparison['verdict'] == 'zone-prefix' and comparison['zone'] == 38 and comparison['drawing_has_prefix'] is False
    assert codes(report) == ['EPSG:4547']  # the reference's grid without the prefix
    prefixed = dxf(tmp_path/'prefixed.dxf', site(38_495_000, 2_490_000))
    bare_ref = reference(tmp_path/'cm114.gpkg', 'EPSG:4547', (494_000, 2_489_000, 497_000, 2_492_000))
    report = inspect(prefixed, tmp_path/'b', '--reference', bare_ref)
    assert report['crs']['reference']['comparisons'][0]['twin']['crs'] == 'EPSG:4526'
    assert codes(report).count('EPSG:4526') == 1  # already a coordinate candidate; not listed twice


def test_reference_shift_and_magnitude_are_held(tmp_path):
    ref = reference(tmp_path/'range.gpkg', 'EPSG:4547', (494_000, 2_489_000, 497_000, 2_492_000))
    shifted = inspect(dxf(tmp_path/'shift.dxf', site(480_000, 2_470_000)), tmp_path/'shift', '--reference', ref)
    comparison = shifted['crs']['reference']['comparisons'][0]
    assert comparison['verdict'] == 'shift-or-rotation' and comparison['overlap'] == 0.0
    assert comparison['offset'][0] == pytest.approx(15_440, abs=1) and shifted['crs']['status'] == 'hold'
    assert any('control points' in note and 'offset' in note for note in shifted['crs']['notes'])
    local = inspect(dxf(tmp_path/'local.dxf', site(105_000, 32_000)), tmp_path/'local', '--reference', ref)
    assert local['crs']['reference']['comparisons'][0]['verdict'] == 'different-magnitude'
    assert local['crs']['status'] == 'hold' and codes(local) == []


def test_reference_must_be_projected_in_metres(tmp_path):
    source = dxf(tmp_path/'site.dxf', site(495_000, 2_490_000))
    ref = reference(tmp_path/'lonlat.gpkg', 'EPSG:4326', (113.8, 22.4, 114.0, 22.6))
    assert 'projected CRS in metres' in inspect(source, tmp_path/'out', '--reference', ref, success=False)
    assert not (tmp_path/'out').exists()


def assign(source, out, *extra, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'assign-crs', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def converted(tmp_path, entities, name='drawing', **options):
    """A drawing, its inspection and its CRS-less GPKG made by convert.py."""
    source = dxf(tmp_path/f'{name}.dxf', entities, **options)
    inspect(source, tmp_path/f'{name}-inspect')
    gpkg = tmp_path/f'{name}.gpkg'
    run = subprocess.run([sys.executable, str(SCRIPTS/'convert.py'), str(source), '--output', str(gpkg)], capture_output=True, text=True)
    assert run.returncode == 0, run.stdout+run.stderr
    return gpkg, tmp_path/f'{name}-inspect'


CONFIRM = ('--confirmed-by', '张三', '--basis', '数据提供方书面说明为 CGCS2000 3 度带 39 带')


def test_assign_crs_labels_without_touching_coordinates(tmp_path):
    gpkg, inspection = converted(tmp_path, site(39_500_000, 3_400_000))
    before = gpd.read_file(gpkg)
    assert before.crs is None
    manifest = assign(gpkg, tmp_path/'out', '--inspection', inspection, '--crs', 'EPSG:4527', *CONFIRM)
    after = gpd.read_file(tmp_path/'out'/'drawing.gpkg')
    assert after.crs.to_epsg() == 4527 and len(after) == len(before)
    assert list(shapely.to_wkb(after.geometry.to_numpy())) == list(shapely.to_wkb(before.geometry.to_numpy()))
    assert manifest['confirmed_by'] == '张三' and manifest['basis'].startswith('数据提供方')
    assert manifest['crs'] == {'id': 'EPSG:4527', 'name': 'CGCS2000 / 3-degree Gauss-Kruger zone 39', 'datum': 'China 2000'}
    assert manifest['outside_candidates'] is False and manifest['checks']['drawing']['geometry_wkb_identical'] is True
    assert manifest['inspection']['drawing'] == 'drawing.dxf' and manifest['input']['sha256']
    assert gpd.read_file(gpkg).crs is None  # the input is not modified


def test_assign_crs_refusals(tmp_path):
    gpkg, inspection = converted(tmp_path, site(39_500_000, 3_400_000))
    base = ('--inspection', inspection)
    assert 'must not be blank' in assign(gpkg, tmp_path/'a', *base, '--crs', 'EPSG:4527', '--confirmed-by', ' ', '--basis', 'x', success=False)
    assert 'not among the inspection candidates' in assign(gpkg, tmp_path/'b', *base, '--crs', 'EPSG:4547', *CONFIRM, success=False)
    assert 'projected CRS in metres' in assign(gpkg, tmp_path/'c', *base, '--crs', 'EPSG:4490', *CONFIRM, success=False)
    outside = assign(gpkg, tmp_path/'d', *base, '--crs', 'EPSG:4547', *CONFIRM, '--outside-candidates')
    assert outside['outside_candidates'] is True and outside['inspection']['candidates_as_drawn'][0] == 'EPSG:2363'
    labelled = tmp_path/'d'/'drawing.gpkg'
    assert 'already has a CRS' in assign(labelled, tmp_path/'e', *base, '--crs', 'EPSG:4527', *CONFIRM, success=False)
    moved = tmp_path/'moved.gpkg'
    frame = gpd.read_file(gpkg)
    frame.set_geometry(frame.translate(5000, 0)).to_file(moved)
    assert 'outside the inspected drawing' in assign(moved, tmp_path/'f', *base, '--crs', 'EPSG:4527', *CONFIRM, success=False)
    assert 'new directory' in assign(gpkg, tmp_path/'d', *base, '--crs', 'EPSG:4527', *CONFIRM, success=False)
    assert not any((tmp_path/name).exists() for name in 'abcef')


def test_assign_crs_refuses_a_candidate_that_needs_unit_scaling(tmp_path):
    gpkg, inspection = converted(tmp_path, site(39_500_000_000, 3_400_000_000, step=1000), name='mm', units=4)
    message = assign(gpkg, tmp_path/'out', '--inspection', inspection, '--crs', 'EPSG:4527', *CONFIRM, '--outside-candidates', success=False)
    assert 'only after unit scaling' in message and not (tmp_path/'out').exists()


def extract(source, out, *extra, success=True, ezdxf=True):
    command = [sys.executable, str(SCRIPTS/'cad.py')] if ezdxf else [sys.executable, '-c', WITHOUT_EZDXF, str(SCRIPTS)]
    result = subprocess.run([*command, 'extract', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def layer(out, name):
    return gpd.read_file(out/'raw.gpkg', layer=name)


def plan_drawing(path, units=6, scale=1.0, z=0.0):
    """R2018 parcels: a bulged boundary, an MTEXT number, a full-width code, an aligned TEXT and attributed blocks."""
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    doc.header['$INSUNITS'] = units
    msp = doc.modelspace()
    x0, y0, s = 38_500_000*scale, 2_490_000*scale, scale
    # Two straight sides and a semicircular side of radius 50: bulge 1 between (x0+100, y0) and (x0+100, y0+100).
    msp.add_lwpolyline([(x0, y0, 0, 0, 0), (x0+100*s, y0, 0, 0, 1), (x0+100*s, y0+100*s, 0, 0, 0), (x0, y0+100*s, 0, 0, 0)],
                       format='xyseb', close=True, dxfattribs={'layer': '地块界线', 'elevation': z})
    msp.add_line((x0, y0+100*s), (x0+300*s, y0+100*s), dxfattribs={'layer': '道路'})
    msp.add_mtext(r'{\fSimSun|b0;03-03-02}\P', dxfattribs={'layer': '地块编号', 'insert': (x0+50*s, y0+50*s), 'char_height': 2.5*s})
    msp.add_text('Ｒ２＋Ｃ１', dxfattribs={'layer': '用地代码', 'insert': (x0+50*s, y0+40*s), 'height': 2*s})
    aligned = msp.add_text('C1', dxfattribs={'layer': '用地代码', 'height': 2*s})
    aligned.set_placement((x0+60*s, y0+40*s), align=ezdxf.enums.TextEntityAlignment.MIDDLE_CENTER)
    block = doc.blocks.new('地块标签')
    block.add_attdef('BH', (0, 0))
    block.add_circle((0, 0), 1*s)
    for point, value in (((x0+20*s, y0+20*s), '03-01'), ((x0+80*s, y0+20*s), '03-02')):
        msp.add_blockref('地块标签', point, dxfattribs={'layer': '地块编号'}).add_attrib('BH', value, point)
    doc.saveas(path)
    return path


def test_extract_layers_fields_and_counts_match_inspect(tmp_path):
    source = plan_drawing(tmp_path/'plan.dxf')
    report = inspect(source, tmp_path/'inspect')
    manifest = extract(source, tmp_path/'out', '--units', 'header')
    counts = {row['layer']: row['entities'] for row in report['layers']}
    assert {k: sum(v.values()) for k, v in manifest['read'].items()} == counts  # same entities as the inspection
    assert manifest['layers'] == {'lines': 2, 'closed_polygons': 1, 'texts': 5, 'blocks': 2, 'leaders': 0, 'hatches': 0} and manifest['crs'] is None
    lines, polygons, texts, blocks = (layer(tmp_path/'out', n) for n in ('lines', 'closed_polygons', 'texts', 'blocks'))
    assert lines.crs is None and set(lines['layer']) == {'地块界线', '道路'} and lines['handle'].notna().all()
    assert polygons['valid'].all() and polygons['handle'].iloc[0] == lines.set_index('dxf_type').loc['LWPOLYLINE', 'handle']
    by_raw = texts.set_index('text_raw')
    assert by_raw.loc[r'{\fSimSun|b0;03-03-02}\P', 'text_clean'] == '03-03-02'
    assert by_raw.loc['Ｒ２＋Ｃ１', 'text_clean'] == 'R2+C1' and by_raw.loc['C1', 'anchor'] == 'align'
    assert by_raw.loc['C1'].geometry.x == pytest.approx(38_500_060)
    attribs = texts[texts['dxf_type'] == 'ATTRIB']
    assert sorted(attribs['text_clean']) == ['03-01', '03-02'] and set(attribs['tag']) == {'BH'}
    assert set(attribs['block']) == {'地块标签'} and attribs['source_handle'].notna().all()
    assert sorted(json.loads(a)['BH'] for a in blocks['attributes']) == ['03-01', '03-02']
    assert manifest['warnings'] == []  # R2018 carries real handles


def test_extract_flattens_arcs_within_tolerance(tmp_path):
    source = plan_drawing(tmp_path/'plan.dxf')
    for tolerance in (0.05, 0.5):
        extract(source, tmp_path/f'out{tolerance}', '--units', 'm', '--layers', '地块界线', '--arc-tolerance', tolerance)
        ring = layer(tmp_path/f'out{tolerance}', 'closed_polygons').geometry.iloc[0].exterior
        centre = (38_500_100, 2_490_050)
        arc = [p for p in ring.coords if p[0] > 38_500_100 + 1e-6]
        assert arc, 'the bulge side bows outwards'
        mids = [((a[0]+b[0])/2, (a[1]+b[1])/2) for a, b in zip(arc, arc[1:])]
        assert max(50 - math.hypot(x-centre[0], y-centre[1]) for x, y in mids) <= tolerance + 1e-9
    assert len(layer(tmp_path/'out0.05', 'lines').geometry.iloc[0].coords) > len(layer(tmp_path/'out0.5', 'lines').geometry.iloc[0].coords)


def test_extract_explodes_blocks_on_request(tmp_path):
    source = plan_drawing(tmp_path/'plan.dxf')
    manifest = extract(source, tmp_path/'out', '--units', 'm', '--layers', '地块编号', '--include-blocks')
    assert manifest['layers'] == {'lines': 2, 'closed_polygons': 0, 'texts': 3, 'blocks': 0, 'leaders': 0, 'hatches': 0}
    circles = layer(tmp_path/'out', 'lines')
    assert set(circles['dxf_type']) == {'CIRCLE'} and set(circles['layer']) == {'地块编号'}  # layer 0 takes the reference's
    assert circles['handle'].isna().all() and circles['source_handle'].notna().all() and set(circles['block']) == {'地块标签'}


def test_extract_units(tmp_path):
    mm = plan_drawing(tmp_path/'mm.dxf', units=4, scale=1000)
    manifest = extract(mm, tmp_path/'mm', '--units', 'header', '--layers', '道路')
    assert manifest['parameters']['scale_to_metre'] == 0.001
    assert layer(tmp_path/'mm', 'lines').total_bounds[0] == pytest.approx(38_500_000)
    lying = plan_drawing(tmp_path/'lying.dxf', units=4)  # says millimetres, drawn in metres
    assert 'pass --units m' in extract(lying, tmp_path/'lying', '--units', 'header', success=False)
    assert extract(lying, tmp_path/'forced', '--units', 'm')['parameters']['units_basis'] == '--units m'
    unitless = plan_drawing(tmp_path/'none.dxf', units=0)
    assert 'unitless' in extract(unitless, tmp_path/'none', '--units', 'header', success=False)


def test_extract_refuses_z_unless_dropped(tmp_path):
    source = plan_drawing(tmp_path/'z.dxf', z=12.5)
    assert 'non-zero Z' in extract(source, tmp_path/'a', '--units', 'm', success=False)
    manifest = extract(source, tmp_path/'b', '--units', 'm', '--drop-z')
    assert manifest['z']['features'] == 2 and manifest['z']['dropped'] and manifest['z']['max'] == 12.5
    assert not layer(tmp_path/'b', 'lines').has_z.any()


def test_extract_refusals_and_r12_handles(tmp_path):
    source = dxf(tmp_path/'site.dxf', site(39_500_000, 3_400_000))
    message = extract(source, tmp_path/'a', '--units', 'm', '--layers', '不存在', success=False)
    assert 'No model-space entities on layers 不存在' in message and '道路' in message
    assert 'needs ezdxf' in extract(source, tmp_path/'b', '--units', 'm', success=False, ezdxf=False)
    manifest = extract(source, tmp_path/'c', '--units', 'm')
    assert any('handles' in w for w in manifest['warnings'])
    assert 'new directory' in extract(source, tmp_path/'c', '--units', 'm', success=False)
    assert not (tmp_path/'a').exists() and not (tmp_path/'b').exists()


def test_extract_then_assign_crs(tmp_path):
    source = plan_drawing(tmp_path/'plan.dxf')
    inspect(source, tmp_path/'inspect')
    extract(source, tmp_path/'raw', '--units', 'header')
    manifest = assign(tmp_path/'raw'/'raw.gpkg', tmp_path/'crs', '--layer', 'closed_polygons', '--inspection', tmp_path/'inspect',
                      '--crs', 'EPSG:4526', *CONFIRM)
    assert manifest['crs']['id'] == 'EPSG:4526' and manifest['input']['layers'] == ['closed_polygons']
    every = assign(tmp_path/'raw'/'raw.gpkg', tmp_path/'all', '--inspection', tmp_path/'inspect', '--crs', 'EPSG:4526', *CONFIRM)
    assert every['input']['layers'] == ['lines', 'closed_polygons', 'texts', 'blocks', 'leaders', 'hatches']
    assert {k: v['features'] for k, v in every['checks'].items()} == {'lines': 2, 'closed_polygons': 1, 'texts': 5, 'blocks': 2, 'leaders': 0, 'hatches': 0}
    for name in every['input']['layers']:
        layer = gpd.read_file(tmp_path/'all'/'raw.gpkg', layer=name)
        assert layer.crs.to_epsg() == 4526  # names kept, empty layers included


def linework(path, lines):
    """A raw.gpkg-like `lines` layer: (CAD layer, handle, coordinates) per line, no CRS."""
    frame = gpd.GeoDataFrame({'layer': [l for l, _, _ in lines], 'handle': [h for _, h, _ in lines]},
                             geometry=[LineString(c) for _, _, c in lines])
    with pytest.warns(UserWarning):
        frame.to_file(path, layer='lines')
    return path


def rect(x, y, w, h, gap=0.0):
    """A rectangle as one open line; with `gap`, its last vertex stops `gap` short of the first along the bottom edge."""
    return [(x+gap, y), (x+w, y), (x+w, y+h), (x, y+h), (x, y)] if gap else [(x, y), (x+w, y), (x+w, y+h), (x, y+h), (x, y)]


def polygonize(source, out, *extra, success=True, snap=0.05, max_gap=0.5, min_area=10):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'polygonize', str(source), '--output', str(out),
                             '--snap', str(snap), '--max-gap', str(max_gap), '--min-area', str(min_area), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def faces_of(out, name='faces'):
    return gpd.read_file(out/'faces.gpkg', layer=name)


def test_polygonize_closed_rectangles_and_shared_edge(tmp_path):
    source = linework(tmp_path/'raw.gpkg', [('地块', 'A', rect(0, 0, 50, 40)), ('地块', 'B', rect(50, 0, 30, 40)),
                                            ('地块', 'C', rect(200, 0, 20, 20))])
    manifest = polygonize(source, tmp_path/'out')
    faces = faces_of(tmp_path/'out')
    assert manifest['faces'] == {'pass': 3, 'hold': 0, 'outside_boundary': None}
    assert sorted(faces['area_m2']) == [400, 1200, 2000] and faces['valid'].all()
    shared = faces[faces['area_m2'] == 2000].iloc[0]
    assert shared['source_handles'] == 'A,B' and shared['source_layers'] == '地块'  # the shared side belongs to both
    assert manifest['layers']['qa_gaps'] == 0 and manifest['layers']['qa_dangles'] == 0
    assert not any(row['check'] == 'overlaps' for _, row in faces_of(tmp_path/'out', 'qa_topology').iterrows())
    assert manifest['area']['boundary_area_m2'] is None and manifest['notes'][0].startswith('No --boundary')
    assert (tmp_path/'out'/'check.png').stat().st_size > 10_000 and faces.crs is None


def test_polygonize_small_gap_snaps_large_gap_is_reported(tmp_path):
    source = linework(tmp_path/'raw.gpkg', [('地块', 'S', rect(0, 0, 10, 10, gap=0.03)), ('地块', 'L', rect(100, 0, 10, 10, gap=0.3))])
    manifest = polygonize(source, tmp_path/'out')
    assert manifest['snap']['endpoints_merged'] == 2 and manifest['snap']['largest_move_m'] == pytest.approx(0.015)
    faces = faces_of(tmp_path/'out')
    assert len(faces) == 1 and faces['source_handles'].iloc[0] == 'S'
    gap = faces_of(tmp_path/'out', 'qa_gaps')
    assert len(gap) == 1 and gap['kind'].iloc[0] == 'endpoint' and gap['length_m'].iloc[0] == pytest.approx(0.3)
    assert gap['handles_a'].iloc[0] == 'L' and not gap['closed'].iloc[0]
    closed = polygonize(source, tmp_path/'closed', '--close-gaps')
    faces = faces_of(tmp_path/'closed')
    assert len(faces) == 2 and faces.set_index('source_handles').loc['L', 'closed_gap'] and closed['layers']['added_lines'] == 1


def test_polygonize_t_junction_dangle_sliver_and_stray(tmp_path):
    source = linework(tmp_path/'raw.gpkg', [
        ('地块', 'R', rect(0, 0, 40, 20)),
        ('地块', 'T', [(20, 0.02), (20, 20)]),        # stops 0.02 short of the bottom side
        ('地块', 'D', [(40, 10), (60, 10)]),          # sticks out: a dangle
        ('地块', 'P', rect(100, 0, 1, 2)),            # 2 m²: a suspected sliver
        ('地块', 'X', [(300, 0), (300.4, 0)]),        # short stray line, not a gap
    ])
    manifest = polygonize(source, tmp_path/'out')
    assert manifest['snap']['endpoints_moved_to_line'] == 1
    faces = faces_of(tmp_path/'out')
    assert sorted(faces['area_m2']) == [2, 400, 400]
    sliver = faces[faces['area_m2'] == 2].iloc[0]
    assert sliver['status'] == 'hold' and 'sliver' in sliver['reason']
    loose = faces_of(tmp_path/'out', 'qa_dangles')
    assert set(loose['handles']) == {'D', 'X'} and set(loose['kind']) == {'dangle'}
    assert manifest['layers']['qa_gaps'] == 0
    assert any(row['check'] == 'slivers' for _, row in faces_of(tmp_path/'out', 'qa_topology').iterrows())


def test_polygonize_boundary_limits_faces(tmp_path):
    source = linework(tmp_path/'raw.gpkg', [
        ('规划范围线', 'B', rect(0, 0, 100, 50)),
        ('地块', 'A', [(50, 0), (50, 50)]),
        ('地块', 'O', rect(200, 0, 20, 20)),
    ])
    manifest = polygonize(source, tmp_path/'out', '--cad-layers', '地块', '--boundary', '规划范围线')
    assert manifest['faces']['outside_boundary'] == 1 and manifest['parameters']['cad_layers'] == ['地块', '规划范围线']
    assert manifest['area'] == {'faces_area_m2': 5000.0, 'boundary_area_m2': 5000.0, 'share_of_boundary': 1.0}
    assert len(faces_of(tmp_path/'out')) == 2
    open_boundary = linework(tmp_path/'open.gpkg', [('规划范围线', 'B', rect(0, 0, 100, 50, gap=2))])
    assert 'do not close' in polygonize(open_boundary, tmp_path/'bad', '--boundary', '规划范围线', success=False)


def test_polygonize_refusals_and_crs(tmp_path):
    source = linework(tmp_path/'raw.gpkg', [('地块', 'A', rect(0, 0, 10, 10))])
    assert '--max-gap > --snap' in polygonize(source, tmp_path/'a', snap=0.5, max_gap=0.5, success=False)
    assert 'No lines on CAD layers 不存在' in polygonize(source, tmp_path/'b', '--cad-layers', '不存在', success=False)
    projected = tmp_path/'projected.gpkg'
    gpd.read_file(source, layer='lines').set_crs('EPSG:4547').to_file(projected, layer='lines')
    assert polygonize(projected, tmp_path/'c')['source']['crs'] == 'EPSG:4547'
    assert faces_of(tmp_path/'c').crs.to_epsg() == 4547
    lonlat = tmp_path/'lonlat.gpkg'
    gpd.read_file(source, layer='lines').set_crs('EPSG:4326').to_file(lonlat, layer='lines')
    assert 'projected in metres' in polygonize(lonlat, tmp_path/'d', success=False)


def test_extract_then_polygonize(tmp_path):
    source = plan_drawing(tmp_path/'plan.dxf')
    extract(source, tmp_path/'raw', '--units', 'header', '--layers', '地块界线,道路')
    manifest = polygonize(tmp_path/'raw'/'raw.gpkg', tmp_path/'faces')
    faces = faces_of(tmp_path/'faces')
    # 100 x 100 square plus a half disc of radius 50 on its right side, flattened at 0.05 m.
    assert len(faces) == 1 and faces['area_m2'].iloc[0] == pytest.approx(10_000 + math.pi*50**2/2, rel=1e-3)
    assert manifest['layers']['qa_dangles'] == 1  # the road line runs on past the parcel


RULES = r"""
fields:
  parcel_id:
    layers: [地块编号]
    pattern: '\d{2}-\d{2}'
  landuse:
    layers: [用地代码]
    pattern: '(R[1-4]|C\d|M\d|GIC\d|S\d)(\+[A-Z]+\d*)*'
combined:
  - layers: [地块标注]
    pattern: '(?P<parcel_id>\d{2}-\d{2})\s+(?P<landuse>[A-Z]+\d(\+[A-Z]+\d*)*)'
landuse_major: {R: 居住, C: 商业服务业, G: 绿地广场, GIC: 公共管理与服务}
road:
  layers: [道路红线]
  landuse: [S1]
"""


def label_drawing(path):
    """Seven faces and texts covering every labelling case; see test_label_end_to_end for what each one is."""
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    doc.header['$INSUNITS'] = 6
    msp = doc.modelspace()
    def box(x, y, w, h, layer):
        msp.add_lwpolyline([(x, y), (x+w, y), (x+w, y+h), (x, y+h)], close=True, dxfattribs={'layer': layer})
    def text(value, x, y, layer, mtext=False):
        if mtext:
            return msp.add_mtext(value, dxfattribs={'layer': layer, 'insert': (x, y), 'char_height': 2})
        return msp.add_text(value, dxfattribs={'layer': layer, 'insert': (x, y), 'height': 2})
    for x, y in ((0, 0), (50, 0), (0, 40), (50, 40), (0, 80), (50, 80)):   # A B / C D / E F
        box(x, y, 50, 40, '地块界线')
    box(0, 200, 100, 20, '道路红线')                                        # G: a road strip
    text(r'{\fSimSun;01-01}', 20, 30, '地块编号', mtext=True)               # A
    text('R2', 20, 25, '用地代码')
    text('编号待定', 25, 10, '地块编号')                                     # A: unparsed
    text('01-02 GIC1', 70, 30, '地块标注')                                  # B: combined
    text('01-03', 10, 70, '地块编号')                                       # C: two ids
    text('01-09', 30, 70, '地块编号')
    text('C1', 20, 60, '用地代码')
    note = text('01-04', 150, 60, '地块编号', mtext=True)                  # D: by leader, text outside
    leader = msp.add_leader([(75, 60), (120, 60), (150, 60)], dxfattribs={'layer': '标注'})
    leader.dxf.annotation_handle = note.dxf.handle
    text('R2+C1', 70, 50, '用地代码')
    text('01-01', 70, 110, '地块编号')                                      # F: duplicate of A
    text('M1', 70, 100, '用地代码')
    text('09-09', 50, 20, '地块编号')                                       # on the A|B edge
    text('01-05', 300, 300, '地块编号')                                     # outside every face
    text('人民路', 50, 210, '道路名称')                                     # not covered by the rules
    msp.add_leader([(25, 35), (40, 45)], dxfattribs={'layer': '标注'})      # points at nothing extracted
    doc.saveas(path)
    return path


def run_label(faces, texts, rules, out, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'label', str(faces), '--texts', str(texts), '--rules', str(rules),
                             '--output', str(out)], capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def test_label_end_to_end(tmp_path):
    source = label_drawing(tmp_path/'fatu.dxf')
    raw = extract(source, tmp_path/'raw', '--units', 'm')
    assert raw['layers']['leaders'] == 2
    leader_text = layer(tmp_path/'raw', 'texts').set_index('text_clean').loc['01-04']
    assert leader_text['anchor'] == 'leader' and leader_text.geometry.x == 75 and leader_text['arrows'] == 1
    polygonize(tmp_path/'raw'/'raw.gpkg', tmp_path/'faces', '--cad-layers', '地块界线,道路红线')
    rules = tmp_path/'rules.yaml'
    rules.write_text(RULES, encoding='utf-8')
    manifest = run_label(tmp_path/'faces'/'faces.gpkg', tmp_path/'raw'/'raw.gpkg', rules, tmp_path/'out')
    parcels = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='parcels')
    at = lambda x, y: parcels[parcels.contains(Point(x, y))].iloc[0]
    a, b, c, d, e, f, g = (at(*p) for p in ((25, 20), (75, 20), (25, 60), (75, 60), (25, 100), (75, 100), (50, 210)))
    assert (a['parcel_id'], a['landuse'], a['landuse_major']) == ('01-01', 'R2', '居住') and 'duplicate' in a['reason']
    assert (b['parcel_id'], b['landuse'], b['landuse_major'], b['status']) == ('01-02', 'GIC1', '公共管理与服务', 'pass')
    assert c['parcel_id'] is None and c['status'] == 'hold' and 'multi-value parcel_id' in c['reason'] and c['landuse_major'] == '商业服务业'
    assert (d['parcel_id'], d['landuse_major'], d['status']) == ('01-04', '居住', 'pass')
    assert e['status'] == 'hold' and e['reason'] == 'unknown use' and not e['road']
    assert f['parcel_id'] == '01-01' and f['landuse_major'] is None and 'duplicate parcel_id' in f['reason']
    assert g['road'] and g['status'] == 'pass'
    assert manifest['conflicts'] == {'multi-value': 2, 'duplicate-id': 2, 'on-boundary': 1, 'unassigned': 1, 'leader-unlinked': 1}
    assert manifest['layers'] == {'parcels': 7, 'label_conflicts': 7, 'unparsed_text': 1}
    assert manifest['ignored_text_layers'] == {'道路名称': 1}
    assert manifest['coverage'] == {'faces': 7, 'road_faces': 1, 'unknown_use': 1, 'non_road_known_use': 5,
                                    'with_parcel_id': 4, 'parcel_id_share': 0.8}
    conflicts = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='label_conflicts')
    assert set(conflicts[conflicts['kind'] == 'multi-value']['value']) == {'01-03', '01-09'}
    assert conflicts[conflicts['kind'] == 'on-boundary']['value'].tolist() == ['09-09']
    assert (tmp_path/'out'/'rules.yaml').read_text(encoding='utf-8') == RULES
    assert (tmp_path/'out'/'check.png').stat().st_size > 10_000


def test_label_rule_validation(tmp_path):
    import _cad_label as L
    for text, message in (("fields: {face_id: {layers: [a], pattern: 'x'}}", 'reserved'),
                          ("fields: {parcel_id: {pattern: 'x'}}", 'layers and/or tags'),
                          ("combined: [{layers: [a], pattern: 'x'}]", 'named groups'),
                          ("fields: {parcel_id: {layers: [a], pattern: '('}}", 'bad pattern'),
                          ("fields: {parcel_id: {layers: [a], pattern: 'x'}}\nlanduse_major: {R: 居住}", 'needs a landuse field'),
                          ("colours: {}", 'Unknown rule sections')):
        with pytest.raises(ValueError, match=message):
            L.Rules(__import__('yaml').safe_load(text))
    rules = L.Rules(__import__('yaml').safe_load(RULES))
    assert rules.major('GIC1') == '公共管理与服务' and rules.major('G1') == '绿地广场' and rules.major('R2+C1') == '居住'
    assert rules.classify('地块编号', None, '01-01') == (True, {'parcel_id': '01-01'})
    assert rules.classify('地块编号', None, '01-01x') == (True, {})          # full match only
    assert rules.classify('道路名称', None, '01-01') == (False, {})


def test_inspect_keeps_large_blocks_and_skips_polyface_meshes(tmp_path):
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2000')
    block = doc.blocks.new('区域位置图')
    for i in range(10_050):  # GDAL stops at 10000 features per block unless told otherwise
        block.add_line((i, 0), (i, 1))
    msp = doc.modelspace()
    msp.add_blockref('区域位置图', (39_500_000, 3_400_000), dxfattribs={'layer': '图框'})
    mesh = msp.add_polyface(dxfattribs={'layer': '面'})
    mesh.append_face([(39_500_000, 3_400_000, 0), (39_500_001, 3_400_000, 0), (39_500_001, 3_400_001, 1)])
    source = tmp_path/'big.dxf'
    doc.saveas(source)
    report = inspect(source, tmp_path/'out')
    assert report['entities']['geometries'] == 1  # GDAL merges one block reference into one geometry collection
    import cad
    frame, unreadable, _ = cad.read_gdal(source)
    parts = shapely.get_num_geometries(frame.geometry.to_numpy())
    assert parts.max() == 10_050 and unreadable == 1  # nothing cut at 10000
    assert any('1 polyface meshes' in w for w in report['warnings'])
    assert {row['layer']: row['dxf_types'] for row in report['layers']} == {'图框': {'INSERT': 1}, '面': {'POLYLINE polyface': 1}}


def test_inspect_reads_libredwg_xdata_in_xrecords(tmp_path):
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    doc.modelspace().add_line((0, 0), (10, 10), dxfattribs={'layer': '地块界线'})
    doc.objects.add_xrecord(doc.rootdict.dxf.handle)  # empty, like the 湘源 reactor records in the Mawan drawing
    source = tmp_path/'libredwg.dxf'
    doc.saveas(source)
    text = source.read_text(encoding='utf-8')
    marker = 'AcDbXrecord\n280\n1\n'
    assert marker in text
    # LibreDWG writes pointers (330) after an application name inside XRECORDs; ezdxf refuses them.
    source.write_text(text.replace(marker, marker + '1001\nHYGggl\n330\n2E\n', 1), encoding='utf-8')
    from ezdxf import recover
    with pytest.raises(Exception, match='XDATA'):
        recover.readfile(source)
    report = inspect(source, tmp_path/'out')
    assert report['reader']['inventory'].startswith('ezdxf') and report['layers'][0]['dxf_types'] == {'LINE': 1}
    assert any('1 malformed XDATA tags' in w for w in report['warnings'])


def fake_dwg2dxf(folder, body):
    """A dwg2dxf stand-in, alone on PATH: the converter adapter is tested without LibreDWG or ODA installed."""
    folder.mkdir()
    script = folder/'dwg2dxf'
    script.write_text(f'#!{sys.executable}\nimport shutil, sys\nargs = sys.argv\n' + body, encoding='utf-8')
    script.chmod(0o755)
    return folder


def run_convert(source, out, *extra, path, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'convert', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PATH': str(path), 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def test_convert_runs_the_converter_on_a_copy_and_reads_back(tmp_path):
    prepared = dxf(tmp_path/'prepared.dxf', site(39_500_000, 3_400_000))
    tool = fake_dwg2dxf(tmp_path/'bin', f"print('Warning: Unknown object, skipping eed/reactors/xdic', file=sys.stderr)\n"
                                        f"shutil.copy({str(prepared)!r}, args[args.index('-o') + 1])\n")
    source = tmp_path/'plan.dwg'
    source.write_bytes(b'AC1032' + bytes(200))
    before = source.read_bytes()
    manifest = run_convert(source, tmp_path/'out', '--version', 'r2013', path=tool)
    assert manifest['source']['release'] == 'R2018' and manifest['converter']['tool'] == 'libredwg'
    command = manifest['converter']['command']
    assert command[:3] == ['dwg2dxf', '-y', '--as'] and command[3] == 'r2013' and command[-1] != str(source)
    assert manifest['readback']['model_space_entities'] == 5 and manifest['readback']['layers_with_entities'] == 4
    assert manifest['readback']['header']['codepage'] == 'ANSI_936'
    assert manifest['warnings'] == ['Warning: Unknown object, skipping eed/reactors/xdic']
    assert (tmp_path/'out'/'plan.dxf').read_bytes() == prepared.read_bytes() and source.read_bytes() == before


def test_convert_refusals(tmp_path):
    source = tmp_path/'plan.dwg'
    source.write_bytes(b'AC1032' + bytes(200))
    empty = tmp_path/'empty'
    empty.mkdir()
    assert 'No ODA File Converter or LibreDWG' in run_convert(source, tmp_path/'a', path=empty, success=False)
    failing = fake_dwg2dxf(tmp_path/'bin', "print('Error: bad DWG section', file=sys.stderr)\nsys.exit(1)\n")
    message = run_convert(source, tmp_path/'b', path=failing, success=False)
    assert 'libredwg failed (exit 1)' in message and 'bad DWG section' in message
    assert 'No oda found' in run_convert(source, tmp_path/'c', '--tool', 'oda', path=failing, success=False)
    fake = tmp_path/'fake.dwg'
    fake.write_bytes(b'not a drawing')
    assert 'does not start like a DWG' in run_convert(fake, tmp_path/'d', path=failing, success=False)
    assert not any((tmp_path/name).exists() for name in 'abcd')


@pytest.mark.skipif(shutil.which('dwg2dxf') is None or shutil.which('dxf2dwg') is None, reason='LibreDWG not installed')
def test_convert_with_libredwg(tmp_path):
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2000')
    doc.modelspace().add_line((0, 0), (10, 10))
    doc.saveas(tmp_path/'seed.dxf')
    subprocess.run(['dxf2dwg', '-y', '-o', str(tmp_path/'seed.dwg'), str(tmp_path/'seed.dxf')], capture_output=True, check=True)
    manifest = run_convert(tmp_path/'seed.dwg', tmp_path/'out', path=os.environ['PATH'])
    # LibreDWG's dxf2dwg is experimental, so only the conversion itself is checked here, not its content.
    assert manifest['converter']['tool'] in ('libredwg', 'oda') and manifest['source']['release'] == 'R2000'
    assert manifest['readback']['header']['acadver'] and (tmp_path/'out'/'seed.dxf').stat().st_size > 0


def xiangyuan_drawing(path):
    """A small drawing laid out the way 湘源控规 writes one: parcel outlines on <code>_K with one XDATA application per
    attribute, number blocks repeating them, units on DY范围 named by DY编号, facilities and a planning boundary."""
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    doc.header['$INSUNITS'] = 1  # 湘源 drawings often say inches while drawn in metres
    msp = doc.modelspace()
    fields = ('地块编号', '用地代码', '用地性质', '用地面积', '容积率', '建筑限高', '单位规模', '删除标记', '设施小类名称', '地块', 'LHB_YJ')
    for name in fields:
        doc.appids.new(name)
    def parcel(layer, number, code, name, box_, declared=None, deleted='', bulge=False, height='24'):
        x0, y0, x1, y1 = box_
        points = [(x0, y0, 0, 0, 0), (x1, y0, 0, 0, 1 if bulge else 0), (x1, y1, 0, 0, 0), (x0, y1, 0, 0, 0)]
        outline = msp.add_lwpolyline(points, format='xyseb', close=True, dxfattribs={'layer': layer})
        area = (x1-x0)*(y1-y0) + (math.pi*((y1-y0)/2)**2/2 if bulge else 0)
        for key, value in (('地块编号', number), ('用地代码', code), ('用地性质', name),
                           ('用地面积', str(round(declared if declared is not None else area))), ('容积率', '2.5'),
                           ('建筑限高', height), ('单位规模', '120'), ('删除标记', deleted)):
            outline.set_xdata(key, [(1000, value)])
        return outline
    def fill(layer, box_):
        hatch = msp.add_hatch(color=1, dxfattribs={'layer': layer})
        x0, y0, x1, y1 = box_
        hatch.paths.add_polyline_path([(x0, y0), (x1, y0), (x1, y1), (x0, y1)], is_closed=True)
    parcel('C1_K', '01-01', 'C1', '商业用地', (0, 0, 100, 80), bulge=True)
    second = parcel('G1+C1', '01-02', 'G1+C1', '公园绿地', (100, 0, 200, 80))  # suffix dropped: still consistent
    second.set_xdata('LHB_YJ', [(1000, 'another plug-in')])                    # not a 湘源 field
    second.set_xdata('用地性质', [(1000, '旧值'), (1000, '公园绿地')])             # several values: the last wins
    parcel('R2_K', '01-03', 'R3', '三类居住用地', (0, 80, 100, 160))           # layer says R2, XDATA says R3
    parcel('C1_K', '01-01', 'C1', '商业用地', (100, 80, 200, 160))            # duplicate number
    parcel('C1_K', '01-05', 'C1', '商业用地', (200, 0, 300, 80), deleted='1')
    parcel('G1_K', '01-06', 'G1', '公园绿地', (200, 80, 300, 160), declared=5000, height='N/A')  # 8000 m² drawn
    fill('C1_T', (10, 10, 90, 70))       # 01-01
    fill('G1_T', (110, 10, 190, 70))     # 01-02 is G1+C1: wrong layer code
    fill('R2_T', (600, 600, 610, 610))   # under no parcel
    parcel('DY_K', '02-01', 'DY', '规划控制单元', (0, 200, 300, 300))
    block = doc.blocks.new('地块编号')
    for tag in ('地块编号', '用地代码'):
        block.add_attdef(tag, (0, 0))
    for point, number in (((150, 40), '01-02'), ((50, 120), '09-09'), ((500, 500), '01-07')):
        msp.add_blockref('地块编号', point, dxfattribs={'layer': '12 地块编号'}).add_attrib('地块编号', number, point)
    facility = msp.add_blockref('地块编号', (250, 250), dxfattribs={'layer': '配套设施'})
    facility.set_xdata('设施小类名称', [(1000, '通信机房')])
    facility.set_xdata('地块编号', [(1000, '02-01')])
    for point, number in (((250, 260), '99-99'), ((50, 40), '01-05')):   # unknown number; number of another parcel
        msp.add_blockref('地块编号', point, dxfattribs={'layer': '配套设施'}).set_xdata('地块编号', [(1000, number)])
    roads = doc.blocks.new('luwang')
    copy = roads.add_line((0, 0), (300, 0), dxfattribs={'layer': '0-道路红线'})
    copy.set_xdata('地块编号', [(1000, '01-01')])
    copy.set_xdata('用地代码', [(1000, 'C1')])
    msp.add_blockref('luwang', (0, 0), dxfattribs={'layer': '----道路中心线'})
    msp.add_lwpolyline([(-10, -10), (310, -10), (310, 310), (-10, 310)], close=True, dxfattribs={'layer': '----规划边界'})
    msp.add_lwpolyline([(0, 0), (300, 0), (300, 160), (0, 160)], close=True, dxfattribs={'layer': 'DY范围'})
    msp.add_lwpolyline([(0, 160), (300, 160), (300, 0), (0, 0)], close=True, dxfattribs={'layer': 'DY范围'})  # same, reversed
    msp.add_lwpolyline([(0, 200), (300, 200), (300, 300), (0, 300)], close=True, dxfattribs={'layer': 'DY范围'})
    msp.add_mtext('DY01', dxfattribs={'layer': 'DY编号', 'insert': (150, 100)})
    msp.add_lwpolyline([(400, 0), (500, 0), (500, 50)], dxfattribs={'layer': '地块'})  # no XDATA: ignored
    doc.saveas(path)
    return path


def xiangyuan(source, out, *extra, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'xiangyuan', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def test_xiangyuan_parcels_from_xdata(tmp_path):
    source = xiangyuan_drawing(tmp_path/'xy.dxf')
    manifest = xiangyuan(source, tmp_path/'out', '--units', 'm', '--max-area-diff', 0.005)
    parcels = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='parcels')
    by = {(row['parcel_id'], row['layer']): row for _, row in parcels.iterrows()}
    first, second = parcels[parcels['parcel_id'] == '01-01'].sort_values('area_m2', ascending=False).itertuples()
    assert first.landuse == 'C1' and first.landuse_name == '商业用地' and first.容积率 == 2.5 and first.fills == 1
    assert first.area_m2 == pytest.approx(8000 + math.pi*40**2/2, rel=1e-3) and abs(first.area_diff) < 1e-3
    assert 'duplicate 地块编号' in first.reason and 'duplicate 地块编号' in second.reason
    g = by[('01-02', 'G1+C1')]
    assert g['status'] == 'pass' and g['number_block'] and g['unit_id'] == 'DY01'
    assert 'layer R2_K does not match 用地代码 R3' in by[('01-03', 'R2_K')]['reason']
    assert by[('01-05', 'C1_K')]['reason'] == 'marked deleted'
    assert by[('01-06', 'G1_K')]['reason'].startswith('area differs from 用地面积 by +60.00%')
    unit = by[('02-01', 'DY_K')]
    assert unit['control_type'] == '规划控制单元' and unit['status'] == 'pass' and unit['unit_id'] is None
    assert {'地块编号', '用地代码', '用地性质', '用地面积', '容积率', '删除标记'} <= set(manifest['fields'])
    checks = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='qa_number_blocks')
    assert sorted(zip(checks['kind'], checks['parcel_id'], checks['found'].fillna(''))) == [
        ('number-block-mismatch', '09-09', '01-03'), ('number-block-outside', '01-07', '')]
    units = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='units').sort_values('handle')
    assert list(units['status']) == ['pass', 'hold', 'hold']
    assert list(units['reason'].fillna('')) == ['', 'duplicate outline', 'no unit label inside']
    assert manifest['area']['units_by_id_m2'] == {'DY01': 48000.0}
    facilities = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='facilities')
    telecom = facilities[facilities['设施小类名称'].notna()]
    assert telecom['设施小类名称'].tolist() == ['通信机房'] and telecom['地块编号'].tolist() == ['02-01']
    assert manifest['summary']['parcels'] == 6 and manifest['summary']['control_units'] == 1
    assert manifest['area']['boundary_m2'] == pytest.approx(320*320)
    assert manifest['summary']['number_block_problems'] == 2 and (tmp_path/'out'/'check.png').stat().st_size > 10_000



def test_xiangyuan_fields_fills_facilities_and_copies(tmp_path):
    source = xiangyuan_drawing(tmp_path/'xy.dxf')
    manifest = xiangyuan(source, tmp_path/'out', '--units', 'm')
    parcels = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='parcels')
    g = parcels.set_index('parcel_id').loc['01-02']
    assert json.loads(g['other_xdata']) == {'LHB_YJ': 'another plug-in'} and 'LHB_YJ' not in parcels.columns
    assert g['landuse_name'] == '公园绿地' and g['fills'] == 1  # the last of two values; its fill sits on G1_T
    assert parcels['单位规模'].tolist()[0] == '120' and parcels['建筑限高'].dtype == object
    assert {'容积率', '用地面积'} <= set(manifest['summary']['numeric_fields']) and '建筑限高' not in manifest['summary']['numeric_fields']
    warnings = manifest['warnings']
    assert any("建筑限高 kept as text: 1 values are not numbers ('N/A')" in w for w in warnings)
    assert any('1 XDATA applications held several values' in w for w in warnings)
    assert manifest['block_copies'] == [{'block': 'luwang', 'entities': 1, 'inserts_in_model_space': 1}]
    assert manifest['summary']['parcels'] == 6  # the copy on the road line is not a parcel
    fills = manifest['summary']['fills']
    assert fills['parcels_without_fill'] == 4 and fills['fills_without_parcel'] == 1
    assert fills['fill_code_mismatches'] == [{'parcel_id': '01-02', 'fill_layer': 'G1_T', 'landuse': 'G1+C1'}]
    assert manifest['summary']['facility_matches'] == {'inside': 1, 'unknown id': 1, 'elsewhere': 1}
    facilities = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='facilities').set_index('parcel_match')
    assert facilities.loc['elsewhere', 'parcel_distance_m'] == 150.0  # (50, 40) to the 01-05 box starting at x = 200
    assert manifest['parameters']['units_basis'] == '--units m (header $INSUNITS = inches)'
    assert manifest['custom_entities'] == {'lzx_land_class_registered': False, 'lzx_land_instances': 0, 'proxy_entities': 0}
    schema = tmp_path/'schema.json'
    schema.write_text(json.dumps({'land_attribute_fields_order': ['单位规模'], 'numeric_like': ['单位规模']}, ensure_ascii=False),
                      encoding='utf-8')
    typed = xiangyuan(source, tmp_path/'schema', '--units', 'm', '--schema', schema)
    assert gpd.read_file(tmp_path/'schema'/'parcels.gpkg', layer='parcels')['单位规模'].tolist()[0] == 120.0
    assert typed['schema']['numeric'] == 1 and '单位规模' in typed['summary']['numeric_fields']


def test_xiangyuan_refuses_unexploded_parcels(tmp_path):
    ezdxf = pytest.importorskip('ezdxf')
    source = xiangyuan_drawing(tmp_path/'xy.dxf')
    owner = ezdxf.readfile(source).modelspace().block_record_handle
    text = source.read_text(encoding='utf-8')
    classes = text.index('  2\nCLASSES\n') + len('  2\nCLASSES\n')
    entities = text.index('  2\nENTITIES\n') + len('  2\nENTITIES\n')
    # The DXF form of a 湘源 parcel that was never exploded: a class record and an entity with only its marker.
    record = '  0\nCLASS\n  1\nLZX_LAND\n  2\nAcdbLzxLand\n  3\nLZXOBJ\n 90\n0\n 91\n1\n280\n0\n281\n1\n'
    entity = f'  0\nLZX_LAND\n  5\nFFFF0\n330\n{owner}\n100\nAcDbEntity\n  8\nC1_K\n100\nAcdbLzxLand\n 70\n4\n'
    source.write_text(text[:classes] + record + text[classes:entities] + entity + text[entities:], encoding='utf-8')
    message = xiangyuan(source, tmp_path/'out', '--units', 'm', success=False)
    assert 'unexploded LZX_LAND' in message and 'explode or export them in 湘源控规' in message


def test_xiangyuan_refusals(tmp_path):
    source = xiangyuan_drawing(tmp_path/'xy.dxf')
    message = xiangyuan(source, tmp_path/'a', '--units', 'header', success=False)  # header says inches
    assert 'matches the drawing in metres' in message and 'pass --units m' in message
    plain = dxf(tmp_path/'plain.dxf', site(0, 0))
    assert 'does not look like a 湘源控规 drawing' in xiangyuan(plain, tmp_path/'b', '--units', 'm', success=False)
    assert 'convert DWG first' in xiangyuan(tmp_path/'plan.dwg', tmp_path/'c', '--units', 'm', success=False)


def compound(streams):
    """A minimal compound file (CFB v3, 512-byte sectors) holding {name: bytes} streams under its root."""
    import struct
    size, cutoff, end, free, fat_sector, none = 512, 4096, 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    sectors, fat = [], []

    def place(data):
        count = -(-len(data)//size)
        first = len(sectors)
        for i in range(count):
            sectors.append(data[i*size:(i+1)*size].ljust(size, b'\0'))
            fat.append(first+i+1 if i < count-1 else end)
        return first if count else end
    mini, minifat, starts = b'', [], {}
    for name, data in streams.items():
        if len(data) < cutoff:
            count = -(-len(data)//64)
            starts[name] = len(mini)//64
            minifat += [starts[name]+i+1 for i in range(count-1)] + [end]
            mini += data.ljust(count*64, b'\0')
    first_minifat = place(struct.pack(f'<{len(minifat)}I', *minifat)) if minifat else end
    mini_start = place(mini) if mini else end
    for name, data in streams.items():
        if len(data) >= cutoff:
            starts[name] = place(data)

    def entry(name, kind, start, length, child=none, right=none):
        encoded = name.encode('utf-16-le') + b'\0\0'
        return (encoded.ljust(64, b'\0') + struct.pack('<HBB3I', len(encoded), kind, 1, none, right, child) + b'\0'*36
                + struct.pack('<IQ', start, length))
    names = list(streams)
    directory = entry('Root Entry', 5, mini_start, len(mini), child=1 if names else none)
    directory += b''.join(entry(n, 2, starts[n], len(streams[n]), right=i+2 if i+1 < len(names) else none) for i, n in enumerate(names))
    first_directory = place(directory.ljust(-(-len(directory)//size)*size, b'\0'))
    count = -(-len(sectors)//127)  # FAT sectors needed, each also listed in the FAT
    fat_ids = list(range(len(sectors), len(sectors)+count))
    fat += [fat_sector]*count
    fat_bytes = struct.pack(f'<{len(fat)}I', *fat).ljust(count*size, b'\xff')
    header = (bytes.fromhex('D0CF11E0A1B11AE1') + b'\0'*16 + struct.pack('<HHHHH', 0x3E, 3, 0xFFFE, 9, 6) + b'\0'*6
              + struct.pack('<9I', 0, count, first_directory, 0, cutoff, first_minifat, len(minifat)*4//size + bool(len(minifat)*4 % size),
                            end, 0)
              + struct.pack('<109I', *(fat_ids + [free]*(109-count))))
    return header + b''.join(sectors) + fat_bytes


def ole_frame(owner, handle, layer, corners, streams):
    """OLE2FRAME DXF text as LibreDWG writes it: placeholder 10/11 points, the real corners in the binary header."""
    import struct
    data = b'\x80\x55' + struct.pack('<12d', *[v for x, y in corners for v in (x, y, 0.0)]) + b'\0'*30 + compound(streams)
    chunks = ''.join(f'310\n{data[i:i+127].hex().upper()}\n' for i in range(0, len(data), 127))
    return (f'  0\nOLE2FRAME\n  5\n{handle}\n330\n{owner}\n100\nAcDbEntity\n  8\n{layer}\n100\nAcDbOle2Frame\n 70\n2\n  3\nOLE\n'
            f' 10\n30.0\n 20\n-18.0\n 30\n0.0\n 11\n35.0\n 21\n-22.0\n 31\n0.0\n 71\n2\n 72\n0\n 73\n2\n 90\n{len(data)}\n{chunks}  1\nOLE\n')


def workbook(active, unit_volume):
    import io
    openpyxl = pytest.importorskip('openpyxl')
    book = openpyxl.Workbook()
    book.active.title = 'Sheet1'
    book.active.append(['地块编号', '用地面积(平方米)'])
    book.active.append(['01-01', 9257])
    units = book.create_sheet('Sheet2')
    units.append(['规划控制单元编号', '建筑容积总量（万平方米）'])
    units.append(['DY01', unit_volume])
    book.active = book.sheetnames.index(active)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def comp_obj_stream(user, prog):
    def string(text, encoding='ascii'):
        raw = text.encode(encoding) + b'\0'
        return len(raw).to_bytes(4, 'little') + raw
    return (bytes.fromhex('0100FEFF030A0000FFFFFFFF') + b'\0'*16 + string(user, 'gbk') + string('ExcelML12') + string(prog)
            + b'\0'*12)


def embedded_drawing(path):
    """The 湘源 drawing plus two embedded Excel workbooks that disagree on Sheet2, a table block whose ACAD_TABLE
    came out of LibreDWG empty, and a number block with fields its outline lacks or already has."""
    ezdxf = pytest.importorskip('ezdxf')
    pytest.importorskip('olefile')
    source = xiangyuan_drawing(path)
    doc = ezdxf.readfile(source)
    msp = doc.modelspace()
    number = next(e for e in msp.query('INSERT') if any(a.dxf.text == '01-02' for a in e.attribs))
    number.add_attrib('备注', '依据政府批件', (150, 40))
    number.add_attrib('容积率', '9.9', (150, 40))      # the outline says 2.5: not overwritten
    number.add_attrib('用地代码', 'C1', (150, 40))     # a key field: never filled
    table = doc.blocks.new_anonymous_block(type_char='T')
    for x, y, text in ((10, -5, '用地代码'), (40, -5, '用地面积(公顷)'), (10, -15, 'R2'), (40, -15, '2.77'), (25, -10, '合并')):
        table.add_mtext(text, dxfattribs={'layer': '表格图层', 'insert': (x, y), 'char_height': 2.5})
    doc.saveas(source)
    owner = msp.block_record_handle
    frames = ''.join(ole_frame(owner, handle, '图框_图框', [(400, 100), (500, 100), (500, 50), (400, 50)],
                               {'\x01CompObj': comp_obj_stream('Microsoft Excel 工作表', 'Excel.Sheet.12'),
                                'Package': workbook(active, volume)})
                     for handle, active, volume in (('FFF01', 'Sheet1', 59.59), ('FFF02', 'Sheet2', 61.89)))
    # The OLE user type is in the embedding system's code page; ezdxf writes ANSI_1252 into every R2018 file.
    text = source.read_text(encoding='utf-8').replace('$DWGCODEPAGE\n  3\nANSI_1252', '$DWGCODEPAGE\n  3\nANSI_936')
    entities = text.index('  2\nENTITIES\n') + len('  2\nENTITIES\n')
    source.write_text(text[:entities] + frames + '  0\nACAD_TABLE\n' + text[entities:], encoding='utf-8')
    return source, table.name


def digest_of(path):
    import hashlib
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_xiangyuan_reads_embedded_workbooks_tables_and_block_fields(tmp_path):
    source, block = embedded_drawing(tmp_path/'xy.dxf')
    manifest = xiangyuan(source, tmp_path/'out', '--units', 'm')
    out = tmp_path/'out'
    frames = gpd.read_file(out/'parcels.gpkg', layer='ole_frames').set_index('handle')
    assert list(frames.index) == ['FFF01', 'FFF02'] and set(frames['corners_from']) == {'binary header'}
    assert frames.loc['FFF01'].geometry.bounds == (400, 50, 500, 100)
    assert frames.loc['FFF01', 'prog_id'] == 'Excel.Sheet.12' and frames.loc['FFF01', 'user_type'] == 'Microsoft Excel 工作表'
    assert dict(frames['shown_sheet']) == {'FFF01': 'Sheet1', 'FFF02': 'Sheet2'}
    import openpyxl
    saved = openpyxl.load_workbook(out/frames.loc['FFF02', 'file'])
    assert saved['Sheet2']['B2'].value == 61.89 and frames.loc['FFF02', 'sha256'] == digest_of(out/frames.loc['FFF02', 'file'])
    assert manifest['embedded']['sheet_versions'] == [{'sheet': 'Sheet2', 'versions': [{'frames': ['FFF01'], 'shown_in': []},
                                                                                         {'frames': ['FFF02'], 'shown_in': ['FFF02']}]}]
    assert any('different versions of 1 sheets (Sheet2)' in w for w in manifest['warnings'])
    assert any('1 ACAD_TABLE entities in the file carry no data' in w for w in manifest['warnings'])
    [table] = manifest['embedded']['tables']
    assert table == {'block': block, 'file': f'tables/{block[1:]}.csv', 'cells': 5, 'rows': 3, 'cols': 3,
                     'first_row': '用地代码 | 用地面积(公顷)', 'table_entity': None, 'table_owner': None, 'block_inserts': 0}
    cells = (out/table['file']).read_text(encoding='utf-8-sig').splitlines()
    assert cells[0] == 'row,col,text,x,y,layer,handle' and cells[3].startswith('2,2,"合并",25.0,-10.0,表格图层,')
    parcels = gpd.read_file(out/'parcels.gpkg', layer='parcels').set_index('handle')
    filled = parcels[parcels['filled_from_block'].notna()]
    assert list(filled['filled_from_block']) == ['备注'] and filled['备注'].tolist() == ['依据政府批件']
    assert filled['容积率'].tolist() == [2.5] and filled['landuse'].tolist() == ['G1+C1']
    assert manifest['summary']['filled_from_block'] == {'备注': 1}
    record = json.loads((out/'record.json').read_text())
    assert {'ole/FFF01.xlsx', 'ole/FFF02.xlsx', table['file']} <= set(record['files'])


def test_xiangyuan_lists_ole_frames_without_olefile(tmp_path):
    source, _ = embedded_drawing(tmp_path/'xy.dxf')
    hide = ("import runpy, sys; sys.modules['olefile'] = None; sys.argv = sys.argv[1:]; "
            "sys.path.insert(0, sys.argv[0].rsplit('/', 1)[0]); runpy.run_path(sys.argv[0], run_name='__main__')")
    result = subprocess.run([sys.executable, '-c', hide, str(SCRIPTS/'cad.py'), 'xiangyuan', str(source), '--units', 'm',
                             '--output', str(tmp_path/'out')], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    manifest = json.loads((tmp_path/'out'/'manifest.json').read_text())
    assert any('2 OLE objects are listed without their contents: olefile is not installed' in w for w in manifest['warnings'])
    frames = gpd.read_file(tmp_path/'out'/'parcels.gpkg', layer='ole_frames')
    assert len(frames) == 2 and frames['file'].isna().all() and frames.geometry.notna().all()
    assert not (tmp_path/'out'/'ole').exists()


def overlay(source, out, *extra, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'overlay', str(source), '--output', str(out), *map(str, extra)],
                            capture_output=True, text=True, env={**os.environ, 'PYTHONDONTWRITEBYTECODE': '1'})
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


X, Y = 500_000, 2_500_000


def overlay_inputs(tmp_path):
    """Parcels in EPSG:4547 (one hold, one self-crossing) with QA layers, and four references."""
    parcels = tmp_path/'parcels.gpkg'
    bowtie = Polygon([(X+200, Y), (X+300, Y+100), (X+300, Y), (X+200, Y+100)])
    gpd.GeoDataFrame({'parcel_id': ['01-01', '01-02', '01-03'], 'landuse': ['R2', 'C1', 'G1'], 'status': ['pass', 'hold', 'pass']},
                     geometry=[box(X, Y, X+100, Y+100), box(X+100, Y, X+200, Y+100), bowtie], crs='EPSG:4547').to_file(parcels, layer='parcels')
    gpd.GeoDataFrame({'kind': ['gap']}, geometry=[LineString([(X+10, Y+10), (X+12, Y+10)])], crs='EPSG:4547').to_file(parcels, layer='qa_gaps')
    gpd.GeoDataFrame({'kind': ['number-block-outside', 'number-block-outside']}, geometry=[Point(X+50, Y+50), Point(X+150, Y+50)],
                     crs='EPSG:4547').to_file(parcels, layer='qa_number_blocks')
    gpd.GeoDataFrame({'name': ['设施']}, geometry=[Point(X+5, Y+5)], crs='EPSG:4547').to_file(parcels, layer='facilities')
    supplied = tmp_path/'supplied.gpkg'
    gpd.GeoDataFrame({'宗地号': ['T1']}, geometry=[box(X+50, Y, X+150, Y+100)], crs='EPSG:4547').to_crs('EPSG:4490').to_file(supplied, layer='已供地')
    renewal = tmp_path/'renewal.shp'
    gpd.GeoDataFrame({'name': ['更新']}, geometry=[Polygon([(X, Y, 10), (X+50, Y, 10), (X+50, Y+50, 10), (X, Y+50, 10)])],
                     crs='EPSG:4547').to_file(renewal)
    broken = tmp_path/'broken.gpkg'
    gpd.GeoDataFrame({'n': [1, 2, 3]}, geometry=[Polygon([(X, Y), (X+100, Y+100), (X+100, Y), (X, Y+100)]), box(X+60, Y+60, X+80, Y+80),
                                                  box(X+5000, Y+5000, X+5100, Y+5100)], crs='EPSG:4547').to_file(broken, layer='plan')
    roads = tmp_path/'roads.gpkg'
    gpd.GeoDataFrame({'n': [1]}, geometry=[LineString([(X-20, Y+100), (X+320, Y+100)])], crs='EPSG:4547').to_file(roads, layer='roads')
    gpd.GeoDataFrame({'n': [1]}, geometry=[Point(X, Y)], crs='EPSG:4547').to_file(roads, layer='points')
    return parcels, ['--reference', f'{supplied}:已供地', '--reference', renewal, '--reference', f'{broken}:plan',
                     '--reference', f'{roads}:roads']


def test_overlay_measures_intersections_and_plots(tmp_path):
    parcels, references = overlay_inputs(tmp_path)
    manifest = overlay(parcels, tmp_path/'out', *references)
    assert manifest['crs'] == 'EPSG:4547' and manifest['input']['invalid_fids'] == [3]
    assert manifest['input']['measured_area_m2'] == pytest.approx(20000)
    assert manifest['hold'] == {'count': 1, 'measured': True, 'parcel_fids': [2], 'parcel_ids': ['01-02']}
    assert manifest['qa'] == {'qa_gaps': 1, 'qa_number_blocks': 2}
    refs = {r['label']: r for r in manifest['references']}
    assert list(refs) == ['已供地', 'renewal', 'plan', 'roads']
    supplied = refs['已供地']
    assert supplied['reprojected'] and supplied['crs'] == 'EPSG:4490' and not supplied['dropped_z']
    assert supplied['parcels_touched'] == 2 and supplied['intersection_area_m2'] == pytest.approx(10000, rel=1e-6)
    assert supplied['share_of_parcels'] == pytest.approx(0.5, rel=1e-6)
    assert refs['renewal']['dropped_z'] and not refs['renewal']['reprojected'] and refs['renewal']['intersection_area_m2'] == pytest.approx(2500)
    plan = refs['plan']
    assert plan['features'] == 2 and plan['invalid_fids'] == [1] and plan['intersection_area_m2'] == pytest.approx(400)  # third is outside the view
    assert refs['roads']['geometry_types'] == ['LineString'] and refs['roads']['parcels_touched'] == 0
    pieces = gpd.read_file(tmp_path/'out'/'overlay.gpkg', layer='intersections')
    assert len(pieces) == manifest['intersections'] == 4 and pieces.crs.to_epsg() == 4547
    row = pieces[(pieces['reference'] == '已供地') & (pieces['parcel_id'] == '01-01')].iloc[0]
    assert row['parcel_fid'] == 1 and row['reference_fid'] == 1 and row['landuse'] == 'R2' and row['status'] == 'pass'
    assert row['area_m2'] == pytest.approx(5000, rel=1e-6) and row['parcel_share'] == pytest.approx(0.5, rel=1e-6)
    assert row['reference_share'] == pytest.approx(0.5, rel=1e-6) and row.geometry.geom_type == 'MultiPolygon'
    assert (tmp_path/'out'/'check.png').stat().st_size > 0
    record = json.loads((tmp_path/'out'/'record.json').read_text())
    assert record['tool'] == 'cad overlay' and set(record['files']) == {'manifest.json', 'overlay.gpkg', 'check.png'}


def test_overlay_exclude_hold_and_qa_layers(tmp_path):
    parcels, references = overlay_inputs(tmp_path)
    manifest = overlay(parcels, tmp_path/'out', *references[:2], '--exclude-hold', '--qa-layers', 'qa_gaps,facilities')
    assert manifest['input']['measured'] == 2 and manifest['input']['measured_area_m2'] == pytest.approx(10000)
    assert manifest['hold']['measured'] is False and manifest['qa'] == {'qa_gaps': 1, 'facilities': 1}
    assert manifest['references'][0]['parcels_touched'] == 1 and manifest['references'][0]['share_of_parcels'] == pytest.approx(0.5, rel=1e-6)
    assert manifest['view'] == pytest.approx([X-30, Y-30, X+330, Y+130])  # the hold parcel still sits inside the measured extent
    assert overlay(parcels, tmp_path/'none', *references[:2], '--qa-layers', '')['qa'] == {}
    assert 'no QA layer missing' in overlay(parcels, tmp_path/'x', *references[:2], '--qa-layers', 'missing', success=False)


def test_overlay_refusals(tmp_path):
    parcels, references = overlay_inputs(tmp_path)
    loose = tmp_path/'loose.gpkg'
    gpd.GeoDataFrame({'n': [1]}, geometry=[box(X, Y, X+10, Y+10)]).to_file(loose, layer='parcels')
    assert 'has no CRS; run cad.py assign-crs' in overlay(loose, tmp_path/'a', *references, success=False)
    geographic = tmp_path/'geographic.gpkg'
    gpd.read_file(parcels, layer='parcels').to_crs('EPSG:4490').to_file(geographic, layer='parcels')
    assert 'projected CRS in metres' in overlay(geographic, tmp_path/'b', *references, success=False)
    assert 'has no layer faces' in overlay(parcels, tmp_path/'c', '--parcels-layer', 'faces', *references, success=False)
    assert 'name one as' in overlay(parcels, tmp_path/'d', '--reference', tmp_path/'roads.gpkg', success=False)
    assert 'a reference must carry one' in overlay(parcels, tmp_path/'e', '--reference', f'{loose}:parcels', success=False)
    (tmp_path/'exists').mkdir()
    assert 'Output must be a new directory' in overlay(parcels, tmp_path/'exists', *references, success=False)


def square(x, y, size):
    return [(x, y), (x+size, y), (x+size, y+size), (x, y+size)]


def hatch_drawing(path):
    """Hatches for each style around nested islands, a coloured pattern, an arc edge, a bow tie and a loop drawn twice."""
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    msp = doc.modelspace()

    def add(layer, loops, style=0, **attributes):
        hatch = msp.add_hatch(dxfattribs={'layer': layer, 'hatch_style': style, **attributes})
        for loop in loops:
            hatch.paths.add_polyline_path(loop, is_closed=True)
        return hatch
    nested = [square(0, 0, 100), square(20, 20, 20), square(25, 25, 10)]  # 10000 - 400 + 100 under even-odd
    add('R2_T', nested, 0)
    add('R2_T', [[(x+200, y) for x, y in loop] for loop in nested], 1)
    add('R2_T', [[(x+400, y) for x, y in loop] for loop in nested], 2)
    pattern = msp.add_hatch(dxfattribs={'layer': 'C1_T'})
    pattern.set_pattern_fill('ANSI31', color=3, scale=2)
    pattern.rgb = (255, 128, 0)
    pattern.paths.add_polyline_path([(0, 200, 0), (100, 200, 1), (100, 300, 0), (0, 300, 0)], is_closed=True)  # half disc east
    add('G1_T', [[(0, 400), (100, 500), (100, 400), (0, 500)]])   # bow tie
    add('G1_T', [square(200, 400, 50), square(200, 400, 50)])     # the same loop twice
    doc.saveas(path)
    return path


def test_extract_hatches_fill_by_style_and_keep_pattern(tmp_path):
    source = hatch_drawing(tmp_path/'fills.dxf')
    manifest = extract(source, tmp_path/'out', '--units', 'm', '--arc-tolerance', 0.01)
    assert manifest['layers']['hatches'] == 6 and 'HATCH' not in manifest['skipped']
    hatches = layer(tmp_path/'out', 'hatches').sort_values('handle').reset_index(drop=True)
    normal, outer, ignore, pattern, bow, twice = (hatches.iloc[i] for i in range(6))
    assert (normal.area_m2, outer.area_m2, ignore.area_m2) == (9700.0, 9600.0, 10000.0)
    assert [h.hatch_style for h in (normal, outer, ignore)] == [0, 1, 2] and normal.loops == 3 and normal.solid
    assert len(list(normal.geometry.geoms)) == 2  # the frame with a hole, and the island inside the hole
    assert pattern.pattern == 'ANSI31' and not pattern.solid and pattern.color == 3 and pattern.true_color == '#ff8000'
    assert pattern.area_m2 == pytest.approx(10000 + math.pi*50**2/2, rel=1e-3)
    assert not bow.valid and math.isnan(bow.area_m2) and twice.valid and twice.area_m2 == 0
    assert any('1 hatches have boundary loops that are not simple polygons' in w for w in manifest['warnings'])
    assert any('1 hatches fill no area' in w for w in manifest['warnings'])


def compare(faces, fills, out, *extra, success=True):
    result = subprocess.run([sys.executable, str(SCRIPTS/'cad.py'), 'compare-fills', str(faces), '--fills', str(fills),
                             '--output', str(out), *map(str, extra)], capture_output=True, text=True)
    assert (result.returncode == 0) == success, result.stdout+result.stderr
    return json.loads((out/'manifest.json').read_text()) if success else result.stderr


def fill_compare_inputs(tmp_path):
    """Faces: A matches its fill, B is half covered by a fill that runs past it, C has none, D and E share one fill,
    and a hold frame encloses everything. A sixth fill lies outside every face."""
    ezdxf = pytest.importorskip('ezdxf')
    doc = ezdxf.new('R2018')
    msp = doc.modelspace()
    for layer, loop in (('R2_T', square(0, 0, 100)), ('C1_T', [(150, 0), (300, 0), (300, 100), (150, 100)]),
                        ('DY控制地块', [(0, 200), (200, 200), (200, 300), (0, 300)]), ('图例', square(900, 900, 10))):
        msp.add_hatch(dxfattribs={'layer': layer}).paths.add_polyline_path(loop, is_closed=True)
    doc.saveas(tmp_path/'fills.dxf')
    extract(tmp_path/'fills.dxf', tmp_path/'raw', '--units', 'm')
    faces = gpd.GeoDataFrame({'face_id': ['A', 'B', 'C', 'D', 'E', 'frame'], 'status': ['pass']*5 + ['hold']},
                             geometry=[box(0, 0, 100, 100), box(100, 0, 200, 100), box(400, 0, 500, 100),
                                       box(0, 200, 100, 300), box(100, 200, 200, 300), box(-50, -50, 1000, 1000)])
    faces.to_file(tmp_path/'faces.gpkg', layer='faces')  # no CRS, like extract output
    return tmp_path/'faces.gpkg', tmp_path/'raw'/'raw.gpkg'


def test_compare_fills_matches_by_iou_and_measures_differences(tmp_path):
    faces, fills = fill_compare_inputs(tmp_path)
    manifest = compare(faces, fills, tmp_path/'out', '--exclude-hold')
    out = tmp_path/'out'/'fill_compare.gpkg'
    rows = gpd.read_file(out, layer='fills').set_index('layer')
    assert rows.loc['R2_T', 'match'] == 'one face' and rows.loc['R2_T', 'face_face_id'] == 'A' and rows.loc['R2_T', 'iou'] == 1.0
    c1 = rows.loc['C1_T']   # 150-300 over B (100-200): 5000 shared, 10000 past B
    assert c1['match'] == 'one face' and c1['face_face_id'] == 'B' and c1['fill_only_m2'] == 10000.0
    assert c1['iou'] == pytest.approx(5000/20000)
    unit = rows.loc['DY控制地块']
    assert unit['match'] == 'several faces' and unit['faces'] == 2 and unit['fill_only_m2'] == 0.0
    assert rows.loc['图例', 'match'] == 'no face'
    faces_out = gpd.read_file(out, layer='faces').set_index('face_id')
    assert faces_out.loc['A', 'fills'] == 1 and faces_out.loc['C', 'fills'] == 0 and faces_out.loc['C', 'face_only_m2'] == 10000.0
    assert faces_out.loc['B', 'face_only_m2'] == 5000.0 and faces_out.loc['E', 'fills_overlapping'] == 1
    assert manifest['summary']['faces_held_out'] == 1 and manifest['summary']['matches'] == {'one face': 2, 'several faces': 1, 'no face': 1}
    assert manifest['area']['fill_only_m2'] == 10100.0 and manifest['area']['face_only_m2'] == 15000.0
    assert (tmp_path/'out'/'check.png').stat().st_size > 0
    # With the hold frame kept, IoU still picks the faces the fills fit, and the legend swatch now falls in the frame.
    compare(faces, fills, tmp_path/'kept')
    kept = gpd.read_file(tmp_path/'kept'/'fill_compare.gpkg', layer='fills').set_index('layer')
    assert kept.loc['R2_T', 'face_face_id'] == 'A' and kept.loc['图例', 'face_face_id'] == 'frame'
    only = compare(faces, fills, tmp_path/'r2', '--exclude-hold', '--cad-layers', 'R2_T')
    assert only['summary']['fills'] == 1 and only['area']['fill_only_m2'] == 0.0


def test_compare_fills_refusals(tmp_path):
    faces, fills = fill_compare_inputs(tmp_path)
    assert 'not a hatches layer' in compare(faces, faces, tmp_path/'a', '--fills-layer', 'faces', success=False)
    assert 'No hatches on CAD layers X_T' in compare(faces, fills, tmp_path/'b', '--cad-layers', 'X_T', success=False)
    located = tmp_path/'located.gpkg'
    gpd.read_file(faces).set_crs('EPSG:4547').to_file(located, layer='faces')
    assert 'must share a CRS' in compare(located, fills, tmp_path/'c', success=False)
